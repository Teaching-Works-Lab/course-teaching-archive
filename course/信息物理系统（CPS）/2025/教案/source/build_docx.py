from __future__ import annotations
from pathlib import Path
from copy import deepcopy
import re, json, shutil, zipfile, hashlib, subprocess
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

BASE=Path('/mnt/data')
TEMPLATE=BASE/'机械测试技术教案.docx'
ROOT=BASE/'25JD32417-信息物理系统（CPS）-教案_源码包'
IMGDIR=ROOT/'images'
MD_NAME='25JD32417-信息物理系统（CPS）-教案.md'
DOCX_NAME='25JD32417-信息物理系统（CPS）-教案.docx'
FINAL_MD=BASE/MD_NAME
FINAL_DOCX=BASE/DOCX_NAME
SOURCE_ZIP=BASE/'25JD32417-信息物理系统（CPS）-教案_源码包.zip'
COMPLETE_DIR=BASE/'25JD32417-信息物理系统（CPS）-教案_完整交付包'
COMPLETE_ZIP=BASE/'25JD32417-信息物理系统（CPS）-教案_完整交付包.zip'

COURSE={
 'name':'信息物理系统（CPS）','english':'Cyber-Physical Systems (CPS)','code':'25JD32417','category':'专业类','nature':'选修','language':'中文','semester':'第7学期','credits':'2',
 'hours_total':32,'hours_theory':32,'major':'智能制造工程','college':'机械与动力工程学院',
 'prereq':'工程建模与科学计算可视化基础（Python）、微机原理及接口技术、机械控制工程基础',
 'followup':'智能制造综合实践、毕业设计（论文）',
 'textbook':'张利国, 詹璟原, 邓恒, 石睿. 信息物理系统原理与设计[M]. 北京：机械工业出版社，2025. ISBN 9787111776550',
 'assessment':'总评成绩 = 平时成绩×20% + 课程作业×20% + 期末考核×60%',
 'assessment_detail':'期末考核采用综合CPS课程设计（综合报告/设计成果 + 答辩与个人核验），不增加闭卷笔试或机考。',
}
REFS='''[1] 张利国, 詹璟原, 邓恒, 石睿. 信息物理系统原理与设计[M]. 北京：机械工业出版社，2025. ISBN 9787111776550\n[2] 任志贵, 张明德. 数字孪生技术[M]. 北京：机械工业出版社，2025. ISBN 9787111770077\n[3] 李双寿, 高君. 数字孪生在智能制造中的工程实践[M]. 北京：机械工业出版社，2024. ISBN 9787111758341\n[4] 黄源, 张婧慧, 唐京瑞. 工业互联网导论[M]. 北京：机械工业出版社，2024. ISBN 9787111748069'''

lessons=[]
def add(no,unit,topic,branches,focus,diff,intro,teach,product,evidence,checks,homework,ideology,innovation,k,a,v,ass):
    lessons.append(dict(no=no,unit=unit,topic=topic,branches=branches,focus=focus,diff=diff,intro=intro,teach=teach,product=product,evidence=evidence,checks=checks,homework=homework,ideology=ideology,innovation=innovation,obj={'k':k,'a':a,'v':v},assessment=ass))

# Unit 1: 4h -> 2 lessons
add(1,'第一单元 CPS概念、架构与智能制造场景','CPS概念、特征、边界与相关技术辨析',
 {'CPS核心':['计算-通信-控制','物理对象/反馈'],'系统特征':['实时耦合','闭环/网络化'],'相关系统':['嵌入式/IoT','工业互联网/数字孪生'],'系统边界':['对象/人员','输入输出/环境']},
 'CPS定义、核心组成及与相邻概念的联系和区别。','从“有传感器和网络”进一步识别真正的计算—物理闭环，并明确系统边界。',
 '给出“联网温度传感器”“自动温控设备”“数字孪生监控平台”三个案例，要求学生判断哪个构成完整CPS、依据是什么。',
 ['从计算、通信、控制、物理对象和反馈五个要素建立CPS基本概念。','比较CPS与嵌入式系统、物联网、工业互联网、数字孪生的侧重点与交叉关系。','学生用“对象—信号—状态—动作—反馈”模板分析一个3D打印设备或机器人工作站，并明确系统内外边界。'],
 '一张CPS概念辨析表 + 一个智能装备CPS边界草图。','概念表、边界图、1条错误判断修正记录。',
 ['CPS与普通联网设备的关键差别是什么？','数字孪生是否天然等同于CPS？','确定系统边界时至少要考虑哪些对象？'],
 '选择一个智能制造装备，列出物理对象、计算节点、通信链路、控制动作和反馈，并说明两项不属于系统内部的对象。',
 '结合智能装备自主化和工业软件与实体装备融合，强调系统工程责任与自主创新意识。',
 '把CPS边界图作为后续数字孪生、智能产线和装备系统概念设计的第一张需求图。',
 ['理解CPS概念、发展背景、基本特征和跨域闭环。','理解CPS与嵌入式、IoT、工业互联网、数字孪生的关系。'],
 ['能够识别CPS主要组成并划定基本系统边界。'],['形成跨机械、电气、控制、计算机和网络的系统思维。'],'平时成绩；目标1.1、2.1、3.1、3.4。')

add(2,'第一单元 CPS概念、架构与智能制造场景','计算—通信—控制—物理闭环与CPS架构分析',
 {'架构层次':['物理对象/感知','计算/决策'],'通信控制':['网络/接口','控制/执行'],'反馈闭环':['状态测量','决策-动作-反馈'],'制造映射':['智能装备','智能产线']},
 'CPS闭环架构、信号/接口和制造场景映射。','区分“结构连接图”和“因果反馈闭环”，说明每条信息/控制流的方向和意义。',
 '展示一张只有模块方框、没有信号方向的“智能产线架构图”，要求学生判断它为什么还不足以说明CPS闭环。',
 ['从感知、计算、通信、控制、执行和反馈角度拆解CPS架构。','用信号方向和反馈关系说明闭环中的状态测量、控制决策和物理响应。','学生以AGV物流或机器人工作站为例绘制CPS架构图，至少标注3类信号和2条反馈关系。'],
 '智能制造场景CPS架构图 + 信号/接口表。','架构图、接口表、闭环路径口头/文字解释。',
 ['CPS反馈闭环至少包括哪几个环节？','传感信号与控制命令的方向是否相同？','架构图为什么要标明系统边界和接口？'],
 '完善课堂架构图，增加网络时延、传感误差或执行故障中的任意两项风险点。',
 '通过跨域接口协作说明局部模块正确并不等于系统整体可靠，培养系统责任意识。',
 '将架构图、接口表视为智能制造CPS概念方案的核心交付物，连接后续综合设计。',
 ['掌握CPS计算、通信、控制和物理过程闭环的基本架构。'],['能够绘制CPS系统架构图并说明信号、接口和反馈关系。'],['形成跨域协同、接口规范和系统边界意识。'],'平时成绩；课程作业过程证据；目标1.1、2.1、3.1、3.4。')

# Unit 2: 6h -> 3 lessons
add(3,'第二单元 离散计算模型与状态机','状态、事件、迁移与有限状态机基础',
 {'基本元素':['状态','事件'],'状态迁移':['迁移','初始/终止状态'],'条件逻辑':['守卫条件','动作'],'行为表达':['状态图','状态转移表']},
 '状态、事件、迁移、守卫与有限状态机。','从自然语言操作流程提取有限、互斥且可验证的状态，而不是把动作名直接当状态。',
 '给出“设备启动、运行、暂停、故障、复位”自然语言流程，让学生先判断“启动”是状态还是事件。',
 ['定义状态、事件、迁移、初始/终止状态和守卫条件。','比较状态图与状态转移表的表达优缺点。','学生为“自动门/输送机/打印机”建立5—7个状态的有限状态机，并加入一个异常状态。'],
 '设备有限状态机 + 状态转移表。','状态图、转移表、至少1条修正前后差异。',
 ['状态和事件的区别是什么？','守卫条件解决什么问题？','一个状态机为什么必须考虑异常状态？'],
 '把“设备启停与故障复位”状态机补全为不少于8条迁移，并标注每条迁移触发条件。',
 '通过设备互锁和状态遗漏案例强化规则、异常和可追溯设计意识。',
 '将状态机作为设备软件、PLC逻辑、机器人任务管理的统一行为描述工具。',
 ['掌握状态、事件、迁移、有限状态机和自动机基本概念。'],['能够建立设备有限状态机并定义异常状态。'],['形成严谨逻辑和异常意识。'],'平时成绩；课程作业过程证据；目标1.2、2.2、3.2、3.3。')

add(4,'第二单元 离散计算模型与状态机','输入输出、同步交互与执行轨迹',
 {'输入输出':['输入事件','输出动作'],'交互':['同步','接口契约'],'执行轨迹':['状态序列','事件序列'],'完整性':['可达状态','缺失迁移']},
 '输入输出自动机、同步交互与执行轨迹。','由状态图生成执行轨迹，并据轨迹反查缺失迁移和接口不一致。',
 '给出“传感器→控制器→执行器”三个离散模块的事件序列，要求学生判断哪些事件必须同步匹配。',
 ['解释输入事件、输出动作和接口契约。','从状态机生成执行轨迹，区分合法轨迹与不可能轨迹。','学生分析一组正常/异常执行轨迹，定位缺失迁移、错误事件或不同步接口。'],
 '输入/输出接口表 + 3条执行轨迹判定表。','接口表、轨迹判定、错误定位说明。',
 ['执行轨迹由哪些信息组成？','两个模块同步交互需要哪些接口条件一致？','不可达状态一定是错误吗？为什么？'],
 '根据给定设备状态机写出2条正常轨迹和2条异常/非法轨迹，并说明原因。',
 '强调系统交互必须遵循接口契约，状态判断必须基于可追踪事件证据。',
 '把执行轨迹用于测试用例设计和后续模型验证，实现从模型到测试的连接。',
 ['理解输入输出、同步交互和执行轨迹。'],['能够根据状态模型生成并检查执行轨迹。'],['形成接口一致性和证据追踪意识。'],'平时成绩；课程作业过程证据；目标1.2、2.2、3.2、3.3。')

add(5,'第二单元 离散计算模型与状态机','设备互锁、任务调度与故障恢复状态模型',
 {'互锁逻辑':['前置条件','禁止条件'],'任务状态':['等待/执行','完成/取消'],'故障处理':['故障检测','安全状态/复位'],'模型检查':['死循环','遗漏/冲突迁移']},
 '互锁、任务状态、故障恢复和状态机完整性检查。','在正常流程之外设计“失效后如何进入安全状态、何时允许恢复”。',
 '给出机械臂“门未关闭却允许自动运行”的状态逻辑错误，让学生先指出缺失的互锁条件。',
 ['从前置条件、禁止条件和安全状态设计设备互锁。','将任务调度中的等待、执行、完成、取消、超时和故障纳入状态模型。','学生对课堂状态机实施“故障注入”，检查能否进入安全状态并可控复位。'],
 '带互锁和故障恢复的状态机 + 故障注入检查表。','状态机、故障用例、恢复轨迹、修改记录。',
 ['互锁条件应放在状态还是迁移条件中？','故障发生后为什么不应直接回到运行状态？','如何判断一个状态机是否存在迁移遗漏？'],
 '完成课程作业中的离散状态模型部分：至少包含正常、边界、故障与恢复轨迹。',
 '以设备安全互锁说明软件逻辑直接影响人员与设备安全，培养对异常路径负责的意识。',
 '将故障恢复状态模型作为智能装备可靠运行与维护服务的核心软件资产。',
 ['掌握设备互锁、任务调度和故障恢复状态模型。'],['能够检查状态机迁移完整性并通过故障轨迹修正模型。'],['形成安全优先、主动找错和可恢复设计意识。'],'平时成绩；课程作业重点证据；目标1.2、2.2、3.2、3.3。')

# Unit 3: 4h -> 2 lessons
add(6,'第三单元 连续物理过程建模','状态变量、输入输出与常微分方程建模',
 {'连续系统':['连续时间','状态变量'],'输入输出':['输入u(t)','输出y(t)'],'动态模型':['微分方程','参数'],'建模过程':['假设/边界','单位/初值']},
 '连续时间系统、状态变量、输入输出和常微分方程。','从物理规律选取状态变量和参数，并显式写出假设、单位和初始条件。',
 '给出“加热器—温度传感器—环境散热”案例，要求学生先指出哪些量是状态、输入、输出和参数。',
 ['解释连续时间状态变量、输入输出和参数的工程含义。','以一阶温度过程或直流电机简化模型说明由物理规律形成常微分方程。','学生完成一个一阶过程建模模板：对象、假设、变量、参数、方程、初值、输出。'],
 '一阶连续过程建模卡 + 方程和单位表。','变量/参数表、方程、假设、单位与初值检查。',
 ['状态变量与输出变量一定相同吗？','为什么模型必须明确参数单位？','建模假设改变时方程是否可能改变？'],
 '选择电机、移动平台或温度过程之一，完成简化ODE模型并说明至少2项假设。',
 '通过参数、单位和简化假设对结论的影响，培养尊重客观规律和模型边界的工程作风。',
 '把连续模型作为数字孪生预测、控制器设计和状态估计的基础模型资产。',
 ['理解连续时间系统、状态变量、输入输出和常微分方程。'],['能够建立简化连续模型并说明变量、参数和假设。'],['形成基于假设、单位和证据建模的习惯。'],'平时成绩；课程作业过程证据；目标1.3、2.3、3.2。')

add(7,'第三单元 连续物理过程建模','平衡点、动态响应、参数敏感性与模型边界',
 {'动态响应':['初始条件','瞬态/稳态'],'平衡点':['平衡条件','稳定概念'],'参数影响':['时间常数','增益/阻尼概念'],'模型边界':['有效范围','误差/局限']},
 '动态响应、平衡点、参数影响和模型适用边界。','避免把演示曲线等同于真实设备，能够说明参数变化和未建模因素。',
 '展示同一一阶温度模型在两个时间常数下的响应曲线，要求学生先判断哪一个系统“反应更快”并说明依据。',
 ['从初始条件、输入和参数解释瞬态/稳态响应。','说明平衡点与基本稳定性概念，不展开控制理论重复推导。','通过教师提供Python/仿真曲线比较参数变化，学生完成“参数—响应—工程含义—模型局限”表。'],
 '连续模型参数敏感性与边界分析表。','响应曲线/教师数据、参数对比表、模型局限说明。',
 ['平衡点的基本含义是什么？','时间常数变大通常意味着响应更快还是更慢？','为什么仿真曲线不能直接证明真实设备一定如此？'],
 '完善连续模型，给出一个模型有效范围和一个可能导致模型失效的现实因素。',
 '强调模型结果必须与假设和证据绑定，不把数学模型当作真实系统本身。',
 '为后续混杂模型中的每个离散模式准备可解释的连续动力学。',
 ['理解平衡点、动态响应和参数影响。'],['能够解释连续模型的工程边界和参数敏感性。'],['形成模型结果不等同于物理事实的证据意识。'],'平时成绩；课程作业重点证据；目标1.3、2.3、3.2。')

# Unit 4: 6h -> 3 lessons
add(8,'第四单元 混杂系统与组合建模','混杂系统、离散模式与连续状态',
 {'混杂系统':['离散模式','连续状态'],'模式动力学':['各模式ODE','状态变量'],'切换':['事件/条件','模式迁移'],'典型场景':['温控','运动/启停']},
 '混杂系统基本思想：离散模式 + 连续动态 + 模式切换。','同时跟踪离散模式和连续状态，理解“切换后动力学规则可能改变”。',
 '以“加热器OFF/ON但温度连续变化”为例，要求学生判断系统中哪些变量离散、哪些连续。',
 ['定义混杂系统和混杂状态。','用温控、运动控制或设备启停案例说明不同离散模式下采用不同连续动力学。','学生把上一单元连续模型嵌入两个离散模式，画出初步混杂状态图。'],
 '两模式混杂系统草图 + 离散/连续变量表。','模式表、每模式连续方程/规则、切换条件草案。',
 ['混杂状态由哪些部分组成？','模式切换后连续状态是否必须清零？','为什么纯状态机无法完整描述温控动态？'],
 '将课堂两模式模型扩展为至少3个模式，并说明每个模式的连续动态。',
 '通过软件模式切换影响真实物理状态的案例，强化软件—物理双向责任。',
 '混杂模型为智能装备逻辑+动力学联合仿真与数字孪生提供统一表达。',
 ['掌握混杂系统、离散模式和连续状态基本概念。'],['能够把状态机与连续模型组合为简化混杂系统。'],['形成软硬件协同和全局分析意识。'],'平时成绩；课程作业过程证据；目标1.4、2.3、3.1、3.3。')

add(9,'第四单元 混杂系统与组合建模','混杂自动机：守卫、重置、不变量与模式切换',
 {'混杂自动机':['模式/状态','连续流'],'守卫条件':['触发边界','切换条件'],'重置':['状态跳变','重置规则'],'不变量':['模式约束','越界禁止']},
 '混杂自动机中的守卫、重置和不变量。','区分守卫条件与模式不变量，并正确处理切换时连续状态的重置/保持。',
 '给出“温度达到80℃切换到降温模式；在加热模式温度不得超过85℃”两条规则，要求学生判断哪条是守卫、哪条是不变量。',
 ['从模式、连续流、守卫、重置和不变量五部分定义混杂自动机。','比较守卫与不变量在“允许切换”和“允许停留”上的不同职责。','学生完善温控/移动平台混杂自动机，加入1个重置和2个不变量并检查边界。'],
 '完整混杂自动机图 + 守卫/重置/不变量表。','自动机图、条件表、边界轨迹检查。',
 ['守卫和不变量的主要差别是什么？','什么情况下需要重置连续状态？','不变量被违反意味着什么？'],
 '完成一个含3模式、至少4守卫、1重置、2不变量的混杂自动机。',
 '以越界和模式切换风险说明形式化约束对物理安全的重要价值。',
 '把安全阈值和模式切换显式化，为后续可达性与反例验证提供模型基础。',
 ['掌握混杂自动机、守卫、重置和不变量。'],['能够建立带安全边界的简化混杂自动机。'],['形成边界前置和安全约束显式化意识。'],'平时成绩；课程作业重点证据；目标1.4、2.3、3.3。')

add(10,'第四单元 混杂系统与组合建模','接口、并行/层次组合与组合冲突分析',
 {'子系统接口':['输入输出','状态/事件接口'],'并行组合':['同步事件','共享变量'],'层次组合':['子系统','上层协调'],'组合风险':['接口冲突','局部正确≠整体安全']},
 '子系统接口、并行/层次组合和组合冲突。','在多个局部模型都“正确”时识别接口语义、时序或共享资源造成的系统冲突。',
 '给出“机器人允许进入工位”和“安全门允许打开”两个各自合理的状态机，组合后出现同一时刻冲突，要求学生找系统层原因。',
 ['解释CPS子系统的输入输出、事件和共享状态接口。','用并行组合与层次化组合描述多设备/多控制器协同。','学生组合两个课堂状态机，建立接口表并制造一个冲突，随后通过约束或协调状态修正。'],
 '双子系统组合模型 + 接口合同 + 冲突/修正记录。','接口表、组合状态图、冲突轨迹、修正后复验。',
 ['并行组合为什么可能产生状态空间增长？','局部安全为什么不一定推出整体安全？','接口合同至少应明确哪些信息？'],
 '完成课程作业中的混杂/组合模型部分，必须包含一个组合冲突及其修正。',
 '通过“局部正确不等于整体安全”强化跨专业协同、接口责任和系统风险意识。',
 '组合建模直接对应智能产线中机器人、输送、视觉、门禁等多子系统集成。',
 ['理解子系统接口、并行组合和层次化组合。'],['能够组合简化模型并识别接口/状态冲突。'],['形成全局安全、接口契约和跨域协同意识。'],'平时成绩；课程作业重点证据；目标1.4、2.3、3.1、3.3。')

# Unit 5: 6h -> 3 lessons
add(11,'第五单元 CPS需求、模型验证与安全','从工程需求到安全性、活性与时序约束',
 {'需求分类':['功能需求','性能/约束'],'安全性':['坏事不发生','边界/禁止状态'],'活性':['好事最终发生','任务完成'],'时序约束':['截止时间','顺序/响应时间']},
 '安全性、活性和时序需求的形式化表达思想。','把“系统要安全/要快/要完成任务”等模糊语言改写为可检查的状态/轨迹约束。',
 '给出“机械臂必须安全”“AGV要及时到达”两句需求，让学生指出为什么无法直接验证，并改写为可检查条件。',
 ['区分功能、性能和安全等工程需求。','用“坏事不发生”解释安全性，用“好事最终发生”解释活性。','学生将5条自然语言需求改写为状态禁止、最终到达或时间约束，并互审可检查性。'],
 'CPS需求表：原始需求、类型、可检查约束、验证证据。','需求表、修改前后表述、互审意见。',
 ['安全性与活性的核心区别是什么？','“系统可靠”为什么不是一个足够具体的验证需求？','时序约束至少应明确哪两个时间点？'],
 '为自己的智能制造场景编写至少2条安全性、2条活性和1条时序约束。',
 '强调安全需求必须前置并可检查，而不是事故后补救，培养质量责任。',
 '把需求表作为期末综合CPS设计的验收合同，连接模型与测试。',
 ['理解需求与模型关系、安全性、活性和时序约束。'],['能够把工程需求改写为基本可验证约束。'],['形成需求可验证、边界清晰和安全前置意识。'],'平时成绩；课程作业过程证据；目标1.4、2.4、3.3、3.4。')

add(12,'第五单元 CPS需求、模型验证与安全','可达性、不变量、执行轨迹与反例',
 {'可达性':['初始状态','可达集合'],'不变量验证':['始终成立','安全边界'],'轨迹检查':['正常轨迹','异常轨迹'],'反例':['失败路径','根因定位']},
 '可达性、不变量、轨迹检查和反例。','理解验证不是“多跑几个正常案例”，而是主动寻找违反需求的可达状态或轨迹。',
 '给出一个模型经过3步后到达“门开且机器人运动”状态，要求学生判断这是一条普通测试失败还是安全性反例。',
 ['用初始状态和迁移关系说明可达状态集合。','用不变量检查安全边界，用执行轨迹检查需求是否被违反。','学生对简化状态/混杂模型进行手工可达分析，找到或证明未找到一条反例。'],
 '可达状态表 + 不变量检查 + 反例轨迹。','状态集合、反例路径、违反需求说明。',
 ['什么是可达状态？','不变量被一个可达状态违反意味着什么？','反例的价值是什么？'],
 '为上一课5条需求至少选择2条做轨迹/可达性检查，并记录反例或验证证据。',
 '通过主动寻找失败条件培养“不回避错误、用证据改进系统”的工程质量文化。',
 '将反例转化为测试用例和设计修正输入，建立模型驱动验证闭环。',
 ['掌握可达性、不变量、轨迹和反例基本思想。'],['能够对简化模型开展基本可达/轨迹检查并识别反例。'],['形成主动找错和证据验证意识。'],'平时成绩；课程作业重点证据；目标1.4、2.4、3.3、3.4。')

add(13,'第五单元 CPS需求、模型验证与安全','死锁、不可达、越界、故障注入与修正复验',
 {'典型问题':['死锁','不可达'],'边界失败':['越界','时序违约'],'故障注入':['通信延迟','传感/执行故障'],'验证闭环':['发现','修正/复验']},
 '死锁、不可达、越界、故障注入及修正复验。','区分“模型中永远到不了的状态”和“系统卡死无后继”的不同问题，并设计修正后复验。',
 '给出一个既包含不可达“完成”状态又包含某状态无出边的模型，让学生分别标出不可达与死锁。',
 ['解释死锁、不可达、越界和时序违约的表现。','用通信延迟、传感器错误、执行器失效等故障注入检查CPS鲁棒性。','学生对本单元模型执行“发现—修正—复验”，保留修改前后模型和反例。'],
 'CPS故障注入与修正复验报告页。','故障条件、原反例、修正内容、复验结果。',
 ['死锁和不可达状态有什么不同？','故障注入为什么不能只测试最容易成功的情况？','修正后为什么必须复验原反例？'],
 '提交课程作业：CPS模型分析、需求与验证证据，必须包含至少1个故障注入与修正复验。',
 '通过故障传播和修正复验强化对人员、设备安全和设计质量负责的意识。',
 '将故障注入和回归复验迁移到智能装备软件测试和数字孪生验证流程。',
 ['理解死锁、不可达、越界和故障注入。'],['能够开展基本故障注入、修正和复验。'],['形成持续改进、可恢复和质量闭环意识。'],'课程作业重点证据；目标1.4、2.4、3.3、3.4。')

# Unit 6: 6h -> 3 lessons
add(14,'第六单元 智能制造CPS设计与数字孪生','智能制造CPS设计流程：感知、边缘、通信、控制与执行',
 {'设计流程':['场景/需求','边界/架构'],'技术链':['感知采集','边缘计算/通信'],'控制执行':['决策','执行/反馈'],'工程约束':['时延','可靠/安全']},
 '智能制造CPS设计流程与感知—计算—通信—控制—执行闭环。','从场景需求推导技术链，而不是先堆砌传感器、平台和算法名词。',
 '给出一张堆满“AI、云、5G、数字孪生”的方案图，要求学生找出缺失的物理对象、控制动作和验收指标。',
 ['按场景→需求→边界→架构→模型→接口→验证组织CPS概念设计流程。','梳理感知、边缘计算、网络通信、控制执行和反馈的职责与时延/可靠性约束。','学生选择智能产线、机器人工作站、AGV、数控或增材制造设备，完成综合方案骨架图。'],
 '期末综合CPS设计骨架：场景、边界、架构、接口和约束。','方案图、接口表、关键时延/安全约束列表。',
 ['CPS设计为什么应从场景和需求开始？','边缘计算在CPS中主要解决哪类问题？','控制闭环中网络时延为什么是工程约束？'],
 '确定期末综合设计场景，完成需求与系统边界初稿。',
 '结合智能工厂自主决策与人员安全，强调技术方案必须明确人工责任与安全边界。',
 '形成期末综合CPS课程设计的最小可评审架构，并训练系统方案表达能力。',
 ['理解智能制造CPS设计流程和各技术环节角色。'],['能够形成智能装备/产线CPS概念设计骨架。'],['形成需求驱动、跨域沟通和工程边界意识。'],'期末考核过程准备；目标1.1、2.1、2.4、3.1、3.4。')

add(15,'第六单元 智能制造CPS设计与数字孪生','数字孪生、工业互联网、边缘计算与数据/时序约束',
 {'数字孪生':['模型-数据-实体','同步/映射'],'工业互联':['设备连接','数据交换'],'边缘计算':['本地处理','云边协同'],'关键约束':['时间同步','数据质量/网络时延']},
 '数字孪生、工业互联网、边缘计算在CPS中的角色与约束。','明确这些技术是CPS的支撑手段而非CPS本身，并说明数据质量、时延、同步对闭环的影响。',
 '比较“实时控制数字孪生”和“每小时更新一次的离线可视化模型”，要求学生判断两者在CPS闭环中的作用是否相同。',
 ['解释数字孪生中的模型—数据—实体映射与同步。','说明工业互联网负责连接/数据交换、边缘计算负责近源计算等角色。','学生为自己的综合设计写出“技术组件—作用—输入输出—时延/数据质量约束—不适用边界”表。'],
 '数字孪生/工业互联网/边缘技术角色与约束表。','角色表、约束表、一个不当技术堆砌案例修正。',
 ['数字孪生为什么不等于CPS？','边缘计算相对云计算的典型优势是什么？','时间同步误差可能怎样影响物理控制判断？'],
 '把数字孪生、工业互联网或边缘计算中至少两项合理融入期末设计，并说明不采用其他技术的理由。',
 '强调数据可信、来源可追溯和模型边界，避免技术名词替代工程事实。',
 '训练技术选型和系统集成判断，为智能制造方案咨询、数字孪生应用和工业互联网项目打基础。',
 ['理解数字孪生、工业互联网和边缘计算在CPS中的角色。'],['能够分析时延、同步和数据质量等工程约束。'],['形成技术边界、数据可信和持续学习意识。'],'期末考核过程准备；目标1.1、2.4、3.4。')

add(16,'第六单元 智能制造CPS设计与数字孪生','综合CPS课程设计评审：模型、需求、验证与工程取舍',
 {'系统方案':['场景/边界','架构/接口'],'核心模型':['状态机','连续/混杂/组合'],'验证证据':['需求','反例/复验'],'工程交付':['技术约束','报告/答辩']},
 '综合CPS方案的架构、模型、验证和工程取舍。','把前五单元产物整合成相互一致、可解释、可验证的系统设计，而不是把图表简单拼接。',
 '给出一份“架构图很好看但状态模型与接口表使用不同信号名”的综合方案，要求学生找出一致性问题。',
 ['以期末考核评分项检查场景、系统边界、架构与接口是否完整。','检查状态模型、连续/混杂模型是否能支撑主要行为，需求是否可验证。','学生进行同伴方案评审：找1个一致性问题、1个安全/验证缺口、1个技术边界问题并提出修正。'],
 '综合CPS课程设计评审包：架构、接口、核心模型、需求/验证表和修正清单。','评审清单、修改前后证据、个人解释记录。',
 ['综合设计至少应包含哪些核心模型/证据？','为什么接口表与状态模型必须使用一致命名？','答辩中如何证明设计结论来自模型与验证而非主观描述？'],
 '完成期末综合CPS课程设计：综合报告/设计成果 + 答辩与个人核验准备。',
 '以综合方案评审强化工程诚信、独立解释、人工复核和对安全结果负责。',
 '形成可展示、可评审的智能制造CPS概念设计成果，连接综合实践与毕业设计。',
 ['综合理解CPS架构、离散/连续/混杂模型、需求验证与数字孪生等技术关系。'],['能够完成智能装备/产线CPS概念设计并通过评审修正。'],['形成可验证设计、工程表达、持续改进和自主学习意识。'],'期末考核核心准备；目标1.1—1.4、2.1—2.4、3.1—3.4。')

assert len(lessons)==16

# ---------- map drawing ----------
FONT_REG='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FONT_BOLD='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
if not Path(FONT_BOLD).exists(): FONT_BOLD=FONT_REG
W,H=1800,620

def font(path,size): return ImageFont.truetype(path,size)
def fit(draw,text,maxw,start=50,minsize=24,bold=False):
    fp=FONT_BOLD if bold else FONT_REG
    for s in range(start,minsize-1,-1):
        f=font(fp,s); b=draw.textbbox((0,0),text,font=f)
        if b[2]-b[0]<=maxw: return f
    return font(fp,minsize)
def box(draw,xy,fill,outline='#8AA1B5',radius=20,width=3): draw.rounded_rectangle(xy,radius=radius,fill=fill,outline=outline,width=width)
def ctext(draw,xy,text,f,fill='#26394C'):
    x0,y0,x1,y1=xy; b=draw.textbbox((0,0),text,font=f); tw,th=b[2]-b[0],b[3]-b[1]
    draw.text(((x0+x1-tw)/2,(y0+y1-th)/2-b[1]),text,font=f,fill=fill)
def draw_map(path,title,branches):
    im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im); line='#819AB5'
    rb=(300,18,1500,112); box(d,rb,'#DCECF7',outline='#4D7894',radius=28,width=4); ctext(d,rb,title,fit(d,title,1120,56,32,True),'#234C68')
    n=len(branches); margin=42; gap=28; cw=(W-2*margin-gap*(n-1))//n; bus=138
    centers=[margin+i*(cw+gap)+cw//2 for i in range(n)]
    d.line((W//2,112,W//2,bus),fill=line,width=4); d.line((centers[0],bus,centers[-1],bus),fill=line,width=4)
    for i,(bl,leaves) in enumerate(branches.items()):
        left=margin+i*(cw+gap); cx=left+cw//2; d.line((cx,bus,cx,160),fill=line,width=4)
        bb=(left+8,160,left+cw-8,236); box(d,bb,'#EDF4F8'); ctext(d,bb,bl,fit(d,bl,cw-50,40,26,True),'#294D68')
        trunk=left+28; d.line((cx,236,cx,258),fill=line,width=3); d.line((cx,258,trunk,258),fill=line,width=3)
        ys=[300,404,508]; d.line((trunk,258,trunk,ys[min(len(leaves),3)-1]+35),fill=line,width=3)
        for j,leaf in enumerate(leaves[:3]):
            y=ys[j]; lx=left+60; rx=left+cw-5; d.line((trunk,y+35,lx,y+35),fill=line,width=3)
            lb=(lx,y,rx,y+70); box(d,lb,'#FBFCFD'); ctext(d,lb,leaf,fit(d,leaf,rx-lx-24,34,23,False),'#343B43')
    im.save(path,dpi=(300,300),optimize=True)

# ---------- markdown ----------
def obj_md(o):
    out=[]
    for label,key in [('知识目标','k'),('能力目标','a'),('价值目标','v')]:
        out.append(f'**{label}：**'); out += [f'- （{i+1}）{x}' for i,x in enumerate(o[key])]
    return '\n'.join(out)
def process_md(l):
    items='\n'.join(f'【教师】{x}\n【学生】完成对应建模/分析并记录判断依据。' for x in l['teach'])
    return f'''【互动导入 15 min】\n【教师】{l['intro']}\n【学生】先独立判断/建模，再交换依据。\n\n【传授新知与课堂训练 70 min】\n{items}\n\n【课堂产物】{l['product']}\n【学习证据】{l['evidence']}\n【评价关联】{l['assessment']}'''
def reflect(l): return f'''【课后填写，不预填事实】\n- 1. 教学流程：记录100 min各环节实际用时及需调整位置；\n- 2. 教学内容：重点记录“{l['diff']}”的真实掌握证据与常见错误；\n- 3. 学生参与：仅依据实际课堂记录填写独立建模、讨论、互审情况，不补写推测；\n- 4. 评价证据：检查本课架构图/状态图/方程/轨迹/验证表是否完整，记录缺项；\n- 5. 后续改进：依据真实课堂证据确定下一轮调整。'''

# ---------- docx helpers ----------
def set_cell_text(cell,text,size=9,bold=False,align=None):
    cell.text=''; lines=str(text).split('\n')
    for i,line in enumerate(lines):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
        if align is not None: p.alignment=align
        r=p.add_run(line); r.bold=bold; r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(size)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def set_margins(cell,top=50,start=65,bottom=50,end=65):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=tcMar.find(qn('w:'+k))
        if el is None: el=OxmlElement('w:'+k); tcMar.append(el)
        el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa')
def clean_md(s):
    s=s.replace('**','').replace('`',''); s=re.sub(r'^- ','',s,flags=re.M); s=re.sub(r'!\[[^\]]*\]\([^\)]*\)','',s); return s.strip()

def parse_md(md_path):
    txt=md_path.read_text(encoding='utf-8'); parts=re.split(r'(?=^## 第\d{2}次课 )',txt,flags=re.M); parts=[p for p in parts if p.startswith('## 第')]
    out=[]
    for p in parts:
        h=re.match(r'^## 第(\d{2})次课 (.+)$',p,re.M); no=int(h.group(1)); topic=h.group(2).strip()
        def sect(name,next_names):
            m=re.search(rf'^### {re.escape(name)}\s*$',p,re.M)
            if not m:return ''
            start=m.end(); ends=[]
            for n in next_names:
                m2=re.search(rf'^### {re.escape(n)}\s*$',p[start:],re.M)
                if m2: ends.append(start+m2.start())
            return p[start:min(ends) if ends else len(p)].strip()
        names=['知识脉络图','教学目标','课程思政','专创融合 / 双创融合','教学重难点','教学方法','教学用具','教学设计','教学过程','过关检测','课堂小结','作业布置','考勤','课后教学反思','本章节参考文献']
        fields={n:sect(n,names[i+1:]) for i,n in enumerate(names)}
        info={}
        for k in ['章节/单元','周次','课时安排']:
            mm=re.search(rf'- \*\*{re.escape(k)}\*\*：(.+)',p)
            if mm: info[k]=mm.group(1).strip()
        img=re.search(r'!\[[^\]]*\]\((images/lesson-\d{2}-knowledge-map\.png)\)',fields['知识脉络图']); fields['image']=img.group(1) if img else ''
        out.append(dict(no=no,topic=topic,info=info,fields=fields))
    return txt,out

# ---------- build ----------
def main():
    for d in [ROOT,COMPLETE_DIR]:
        if d.exists(): shutil.rmtree(d)
    IMGDIR.mkdir(parents=True)
    for z in [SOURCE_ZIP,COMPLETE_ZIP]:
        if z.exists(): z.unlink()
    # semantic images first
    for l in lessons: draw_map(IMGDIR/f'lesson-{l["no"]:02d}-knowledge-map.png',l['topic'],l['branches'])

    # canonical Markdown authored before Word
    md=['# 信息物理系统（CPS）教案（16次课）','', '> **Canonical Markdown / 内容权威源**：本文件与 `images/`、`manifest.json` 共同构成课程教案内容主源；Word 仅由此编译。','', '## 教案封面信息',f'- **学院**：{COURSE["college"]}','- **专业（教研室）**：智能制造工程',f'- **课程名称**：{COURSE["name"]}','- **主讲教师**：（按实际填写）','- **职称**：（按实际填写）',f'- **授课学期**：{COURSE["semester"]}',f'- **授课专业**：{COURSE["major"]}', '- **编制日期**：（按实际填写）','','## 课程基本信息']
    items=[('课程代码',COURSE['code']),('课程名称（中文）',COURSE['name']),('课程名称（英文）',COURSE['english']),('课程类别',f'{COURSE["category"]}；课程性质：{COURSE["nature"]}；授课语言：{COURSE["language"]}'),('授课学期',f'{COURSE["semester"]}；学分：{COURSE["credits"]}'),('课程学时及分配','总学时32；理论32；实验0；上机0；项目式0'),('适用专业',COURSE['major']),('教材',COURSE['textbook']),('授课学院',COURSE['college']),('先修课程',COURSE['prereq']),('后续课程',COURSE['followup']),('考核类型','考查课'),('考核形式',COURSE['assessment_detail']),('考核方式','平时成绩、课程作业、期末综合CPS课程设计'),('总评成绩比例',COURSE['assessment'])]
    md += [f'- **{k}**：{v}' for k,v in items]
    md += ['', '### 课程简介','', '课程基本定位：本课程面向智能装备、智能产线和数字化工厂中的计算过程与物理过程深度融合问题，建立“CPS架构—离散状态模型—连续物理模型—混杂与组合模型—需求/安全/模型验证—智能制造CPS与数字孪生设计”的知识主线。','', '核心学习结果：学生能够解释计算、通信、控制与物理过程的耦合关系；建立状态机、常微分方程、简化混杂自动机和组合模型；把工程需求转化为安全性、活性或时序约束，使用可达性、不变量、轨迹和反例开展基本验证；完成智能装备或产线CPS概念设计。','', '学情分析：学生已学习Python科学计算、微机原理和机械控制工程基础，具备基本计算、嵌入式和控制建模基础。课程32学时全部为理论教学，Python或仿真结果仅用于教师演示辅助理解，不设置独立上机学时。','', '主要教学方法：采用“概念框架—数学/状态模型—执行轨迹—需求约束—案例验证—系统设计”的理论课闭环，以架构图、状态图、时序图、方程、可达状态、反例和设计评审作为主要课堂产物。','', '## 16次课教学设计','']
    for l in lessons:
        md += [f'## 第{l["no"]:02d}次课 {l["topic"]}','',f'- **章节/单元**：{l["unit"]}','- **周次**：按实际课表填写','- **课时安排**：2课时（100 min）','', '### 知识脉络图','',f'![第{l["no"]:02d}次课知识脉络图](images/lesson-{l["no"]:02d}-knowledge-map.png)','', '### 教学目标','',obj_md(l['obj']),'','### 课程思政','',l['ideology'],'','### 专创融合 / 双创融合','',l['innovation'],'','### 教学重难点','',f'教学重点：{l["focus"]}\n教学难点：{l["diff"]}','','### 教学方法','','讲授法、问题引导法、架构/状态建模法、方程分析法、案例研讨法、教师演示法、同伴互评、反例与证据复盘','','### 教学用具','','电脑、投影仪、多媒体课件、教材、状态图/时序图模板、需求与验证表、教师提供Python/仿真曲线或案例数据','','### 教学设计','','课前任务 → 互动导入（15 min）→ 传授新知与课堂训练（70 min）→ 过关检测（10 min）→ 课堂小结（3 min）→ 作业布置（1 min）→ 考勤（1 min）','','### 教学过程','', '#### 课前任务','',f'【教师】发布“{l["topic"]}”预习材料和一个最小建模/判断问题，不提前给出完整答案。\n【学生】阅读对应教材内容，记录3个核心术语、至少1个疑问，并准备课堂建模表。','', '**设计意图：** 通过课前阅读和一个最小建模问题建立先备认知，让学生带着可验证的问题进入课堂。','', '#### 互动导入与传授新知 / 课堂训练','',process_md(l),'', '**设计意图：** 采用“先判断/建模 → 最小讲解 → 模型/轨迹推演 → 独立产出 → 反例/证据复核”的理论课闭环，把抽象CPS原理落实到可检查的图、式、表和验证证据。','', '### 过关检测','', '【教师】组织当堂检测：\n'+'\n'.join(f'- {i+1}. {q}' for i,q in enumerate(l['checks']))+'\n【学生】独立作答；【教师】要求说明模型、轨迹或证据依据。','', '### 课堂小结','',f'【教师】回到知识脉络图，用“概念—模型—关系—边界—证据”总结“{l["topic"]}”。\n【学生】写下本课最重要的一条规则和一个最容易误判的模型边界。','', '### 作业布置','',f'【教师】布置课后任务：{l["homework"]}\n【学生】按课程文件命名规范整理架构图、状态图、方程、轨迹或验证说明。','', '### 考勤','', '【教师】利用学校/课程实际使用的平台进行签到。\n【学生】按课程要求完成签到。','', '### 课后教学反思','',reflect(l),'', '### 本章节参考文献','',REFS,'']
    md_path=ROOT/MD_NAME; md_path.write_text('\n'.join(md),encoding='utf-8'); shutil.copy2(md_path,FINAL_MD)

    # manifest/readme and content gate before Word
    def sha(p):
        h=hashlib.sha256(); h.update(Path(p).read_bytes()); return h.hexdigest()
    manifest={'course':COURSE['name'],'course_code':COURSE['code'],'canonical_markdown':MD_NAME,'lesson_count':16,'hours':{'total':32,'theory':32,'experiment':0,'computer':0,'project':0},'assessment':{'平时成绩':20,'课程作业':20,'期末考核':60},'lessons':[{'lesson':l['no'],'title':l['topic'],'unit':l['unit'],'image':f'images/lesson-{l["no"]:02d}-knowledge-map.png'} for l in lessons],'syllabus_source':'Teaching-Works-Lab/course-teaching-archive: course/信息物理系统（CPS）/2025/大纲/25JD32417-信息物理系统（CPS）-课程教学大纲.md'}
    (ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'README.md').write_text(f'# {COURSE["name"]}教案源码包\n\n本包采用 Markdown-first 工作流。`{MD_NAME}` 是唯一内容主源，`images/` 含16张课次知识脉络图，Word由Markdown内容和学校教案模板编译。课程32学时全部为理论教学，Python/仿真仅作课堂演示辅助，不包含实验、上机或项目式课时。\n',encoding='utf-8')
    txt,parsed=parse_md(md_path)
    assert len(parsed)==16
    assert len(re.findall(r'images/lesson-\d{2}-knowledge-map\.png',txt))==16
    assert txt.count('【课后填写，不预填事实】')==16
    assert all((ROOT/p['fields']['image']).exists() for p in parsed)
    assert COURSE['hours_total']==16*2
    assert '总学时32；理论32；实验0；上机0；项目式0' in txt
    assert '总评成绩 = 平时成绩×20% + 课程作业×20% + 期末考核×60%' in txt

    # compile Word from canonical Markdown
    src=Document(str(TEMPLATE)); proto=deepcopy(src.tables[2]._tbl); body=src._element.body
    # Remove inherited exemplar date; source explicitly uses an actual-fill placeholder.
    for par in src.paragraphs:
        if par.text.strip()=='2024年9月':
            par.text='（按实际填写）'
            for r in par.runs:
                r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(10.5)
            par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for ch in list(body)[26:]:
        if not ch.tag.endswith('}sectPr'): body.remove(ch)
    cover=src.tables[0]
    vals=[COURSE['college'],'智能制造工程',COURSE['name'],'（按实际填写）','（按实际填写）',COURSE['semester'],COURSE['major']]
    for i,v in enumerate(vals): set_cell_text(cover.cell(i,1),v,12,align=WD_ALIGN_PARAGRAPH.CENTER)
    ci=src.tables[1]
    set_cell_text(ci.cell(0,1),COURSE['name'],9); set_cell_text(ci.cell(1,1),COURSE['english'],9)
    set_cell_text(ci.cell(2,1),COURSE['category'],9,align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(ci.cell(2,3),COURSE['nature'],9,align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(ci.cell(2,5),COURSE['language'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(ci.cell(3,1),COURSE['semester'],9,align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(ci.cell(3,5),COURSE['credits'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
    for col,val in enumerate(['课程学时及分配','总学时','理论','实验','上机','其他']): set_cell_text(ci.cell(4,col),val,8.5,bold=(col==0),align=WD_ALIGN_PARAGRAPH.CENTER)
    for col,val in enumerate(['课程学时及分配','32','32','0','0','0']): set_cell_text(ci.cell(5,col),val,9,bold=(col==0),align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(ci.cell(6,1),COURSE['major'],9); set_cell_text(ci.cell(7,1),COURSE['textbook'],8.4); set_cell_text(ci.cell(8,1),COURSE['college'],9); set_cell_text(ci.cell(9,1),COURSE['prereq'],8.3); set_cell_text(ci.cell(10,1),COURSE['followup'],8.5)
    set_cell_text(ci.cell(11,1),'考试课（ ）；考查课（√）',9); set_cell_text(ci.cell(12,1),'闭卷（ ）；开卷（ ）；笔试（ ）；机试（ ）；口试（√，答辩与个人核验）；其它（√，综合CPS课程设计）',8.0)
    set_cell_text(ci.cell(13,1),'作业（√）；报告（√）；随堂分析（√）；线上考试（ ）；线下考试（ ）；其它（综合设计与个人核验）',8.0); set_cell_text(ci.cell(14,1),COURSE['assessment'],8.6)
    intro='课程基本定位：本课程面向智能装备、智能产线和数字化工厂中的计算过程与物理过程深度融合，构建CPS架构—离散状态模型—连续物理模型—混杂与组合模型—需求验证—智能制造CPS设计的知识主线。\n核心学习结果：学生能够建立状态机、常微分方程、混杂/组合模型，使用可达性、不变量、轨迹与反例开展基本验证，并完成智能装备或产线CPS概念设计。\n学情分析：学生已学习Python、微机原理和机械控制工程基础；32学时全部为理论教学，Python或仿真结果仅用于教师演示辅助。\n主要教学方法：概念框架、状态/数学模型、执行轨迹、需求约束、案例验证和系统设计相结合，以架构图、状态图、方程、需求表、反例和评审记录作为课堂学习证据。'
    set_cell_text(ci.cell(15,1),intro,8.0)

    for p in parsed:
        br=src.add_paragraph(); br.add_run().add_break(WD_BREAK.PAGE); new=deepcopy(proto); br._p.addnext(new); t=Table(new,src)
        info=p['info']; f=p['fields']
        set_cell_text(t.cell(0,0),'章节',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(0,1),info['章节/单元'],8.5,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(0,2),'授课题目',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(0,3),p['topic'],8.5,True,WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(t.cell(1,0),'周次',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(1,1),info['周次'],9,False,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(1,2),'课时安排',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(1,3),info['课时安排'],9,False,WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(t.cell(2,0),'知识脉络图',9,True,WD_ALIGN_PARAGRAPH.CENTER); mc=t.cell(2,1); tc=mc._tc; tcPr=tc.tcPr
        for child in list(tc):
            if child is not tcPr: tc.remove(child)
        pe=OxmlElement('w:p'); tc.append(pe); pp=mc.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.add_run().add_picture(str(ROOT/f['image']),width=Inches(4.75)); t.rows[2].height=Inches(1.75); t.rows[2].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        set_cell_text(t.cell(3,0),'教学目标',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(3,1),clean_md(f['教学目标']),8.5)
        set_cell_text(t.cell(4,0),'课程思政',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(4,1),clean_md(f['课程思政']),8.3)
        set_cell_text(t.cell(5,0),'专创融合 /\n双创融合',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(5,1),clean_md(f['专创融合 / 双创融合']),8.3)
        set_cell_text(t.cell(6,0),'教学重难点',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(6,1),clean_md(f['教学重难点']),8.4)
        set_cell_text(t.cell(7,0),'教学方法',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(7,1),clean_md(f['教学方法']),8.3)
        set_cell_text(t.cell(8,0),'教学用具',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(8,1),clean_md(f['教学用具']),8.1)
        set_cell_text(t.cell(9,0),'教学设计',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(9,1),clean_md(f['教学设计']),8.2)
        set_cell_text(t.cell(10,0),'教学过程',11,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(10,1),'教学步骤及主要教学内容',10,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(10,4),'设计意图',9,True,WD_ALIGN_PARAGRAPH.CENTER)
        proc=f['教学过程']
        def sub(name,nextname=None):
            m=re.search(rf'^#### {re.escape(name)}\s*$',proc,re.M)
            if not m:return ''
            end=len(proc)
            if nextname:
                m2=re.search(rf'^#### {re.escape(nextname)}\s*$',proc[m.end():],re.M)
                if m2:end=m.end()+m2.start()
            return proc[m.end():end].strip()
        pre=sub('课前任务','互动导入与传授新知 / 课堂训练'); main=sub('互动导入与传授新知 / 课堂训练')
        pre=re.sub(r'\*\*设计意图：\*\*.*','',pre,flags=re.S).strip(); main_clean=re.sub(r'\*\*设计意图：\*\*.*','',main,flags=re.S).strip()
        set_cell_text(t.cell(11,0),'课前任务',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(11,1),clean_md(pre),8.2); set_cell_text(t.cell(11,4),'通过课前阅读和最小建模问题建立先备认知，让学生带着可验证问题进入课堂。',8.0)
        set_cell_text(t.cell(12,0),'互动导入\n（15 min）\n\n传授新知与课堂训练\n（70 min）',8.5,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(12,1),clean_md(main_clean),7.9); set_cell_text(t.cell(12,4),'采用“先判断/建模→最小讲解→模型/轨迹推演→独立产出→反例/证据复核”的理论课闭环。',7.8)
        set_cell_text(t.cell(13,0),'过关检测\n（10 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(13,1),clean_md(f['过关检测']),8.0); set_cell_text(t.cell(13,4),'即时检验关键概念、模型与边界，要求说明依据。',7.8)
        set_cell_text(t.cell(14,0),'课堂小结\n（3 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(14,1),clean_md(f['课堂小结']),8.0); set_cell_text(t.cell(14,4),'回到知识脉络图压缩信息，形成概念—模型—证据层级结构。',7.8)
        set_cell_text(t.cell(15,0),'作业布置\n（1 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(15,1),clean_md(f['作业布置']),8.0); set_cell_text(t.cell(15,4),'固化课堂建模成果，形成可复查的过程证据。',7.8)
        set_cell_text(t.cell(16,0),'考勤\n（1 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(16,1),clean_md(f['考勤']),8.0); set_cell_text(t.cell(16,4),'保留模板考勤环节，不预填出勤事实。',7.8)
        set_cell_text(t.cell(17,0),'课后\n教学反思',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(17,1),clean_md(f['课后教学反思']),7.9); set_cell_text(t.cell(17,4),'反思必须依据真实课堂证据，课前只提供填写框架。',7.8)
        set_cell_text(t.cell(18,0),'本章节\n参考文献',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(18,1),clean_md(f['本章节参考文献']),7.6)
        for row in t.rows:
            for c in row.cells: set_margins(c)

    src.core_properties.title=f'{COURSE["name"]}教案'; src.core_properties.author=''; src.core_properties.last_modified_by=''; src.save(FINAL_DOCX)
    clean=BASE/'_cps_clean.docx'; subprocess.run(['python','/home/oai/skills/docx/scripts/privacy_scrub.py',str(FINAL_DOCX),'--out',str(clean)],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True); clean.replace(FINAL_DOCX)

    shutil.copy2(__file__,ROOT/'build_docx.py')
    manifest['sha256']={MD_NAME:sha(md_path)}
    for l in lessons: manifest['sha256'][f'images/lesson-{l["no"]:02d}-knowledge-map.png']=sha(IMGDIR/f'lesson-{l["no"]:02d}-knowledge-map.png')
    (ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    with zipfile.ZipFile(SOURCE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for p in ROOT.rglob('*'):
            if p.is_file(): z.write(p,p.relative_to(ROOT.parent))
    COMPLETE_DIR.mkdir(); shutil.copytree(ROOT,COMPLETE_DIR/'source'); shutil.copy2(FINAL_DOCX,COMPLETE_DIR/DOCX_NAME); shutil.copy2(FINAL_MD,COMPLETE_DIR/MD_NAME)
    (COMPLETE_DIR/'README.md').write_text('本包包含 canonical Markdown 源包和由其编译的学校格式 Word。内容修改请先改 source 中 Markdown，再重新编译 Word。\n',encoding='utf-8')
    with zipfile.ZipFile(COMPLETE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for p in COMPLETE_DIR.rglob('*'):
            if p.is_file(): z.write(p,p.relative_to(COMPLETE_DIR.parent))
    print(FINAL_MD); print(FINAL_DOCX); print(SOURCE_ZIP); print(COMPLETE_ZIP)

if __name__=='__main__': main()

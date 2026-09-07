from __future__ import annotations
import os, re, json, shutil, zipfile, hashlib, copy, subprocess
from pathlib import Path
from typing import Dict, List
from graphviz import Digraph
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path('/mnt/data')
TEMPLATE=BASE/'程序设计基础_C_C++_I_教案_24次课_最终版.docx'
OUT_ROOT=BASE/'工程建模与科学计算可视化基础_Python_教案_源码包'
IMG_DIR=OUT_ROOT/'images'
MD_NAME='工程建模与科学计算可视化基础_Python_教案_16次课_最终版.md'
DOCX_NAME='工程建模与科学计算可视化基础_Python_教案_16次课_最终版.docx'
SOURCE_ZIP=BASE/'工程建模与科学计算可视化基础_Python_教案_16次课_源码包.zip'
COMPLETE_ZIP=BASE/'工程建模与科学计算可视化基础_Python_教案_完整交付包.zip'
FINAL_MD=BASE/MD_NAME
FINAL_DOCX=BASE/DOCX_NAME

COURSE={
 'name':'工程建模与科学计算可视化基础（Python）',
 'english':'Fundamentals of Engineering Modeling and Scientific Computing Visualization (Python)',
 'code':'25JD21402','category':'学科基础类','nature':'必修课','language':'中文','semester':'第2学期','credits':'2',
 'hours_total':32,'hours_theory':16,'hours_lab':16,
 'major':'智能制造工程、机械工程、能源与动力工程、车辆工程','college':'机械与动力工程学院','prereq':'无',
 'followup':'人工智能技术与应用、数据结构（C）、机器人技术及应用（A）',
 'textbook':'王小银 等.《Python数据分析与科学计算》[M]. 北京：机械工业出版社，2024. ISBN 9787111742586',
 'assessment':'总评成绩=平时表现20%＋课程作业20%＋期末综合项目60%',
 'syllabus_path':'course/工程建模与科学计算可视化基础/2025/大纲/25JD21402-工程建模与科学计算可视化基础-课程教学大纲.md',
 'syllabus_sha':'6ec19a287e6c6be9f570aeeea36dd744ced765c5','syllabus_date':'2026年9月','version':'2025版目录 v2 定稿'
}
REFS={
 'main':'[1] 王小银 等.《Python数据分析与科学计算》[M]. 机械工业出版社，2024。',
 'sci':'[2] 李志远, 黄化人, 姚明菊, 等.《Python编程与科学计算（微课视频版）》[M]. 清华大学出版社，2023。',
 'python':'[3] 董付国.《Python程序设计（第4版·微课版·在线学习软件版）》[M]. 清华大学出版社，2024。'
}

def L(no, section, topic, kind, split, goals, map_, focus, diff, intro, blocks, product, evidence, assessment, checks, homework, ideology, innovation, tools, refs):
    return dict(no=no,section=section,topic=topic,kind=kind,split=split,goals=goals,map=map_,focus=focus,diff=diff,intro=intro,blocks=blocks,product=product,evidence=evidence,assessment=assessment,checks=checks,homework=homework,ideology=ideology,innovation=innovation,tools=tools,refs=refs)

lessons:List[Dict]=[
L(1,'第一单元 课程导论与工程建模思维','工程建模范式：从问题定义到模型验证','理论','理论2课时',
 {'k':['理解工程问题定义、对象边界、输入输出、变量、约束和评价指标。','理解假设简化、数学模型/计算模型、量纲、边界情况和对照验证的关系。'],
  'a':['能够把一个简单工程问题拆解为对象、变量、约束、模型、求解与验证清单。','能够识别“不合理假设、无验证结论、数据伦理风险”等典型问题。'],
  'v':['形成尊重原始数据、明确假设、用证据验证模型并诚实说明结论边界的工程意识。']},
 {'问题定义':['对象边界 / 输入输出','变量 / 约束 / 指标'],'模型建立':['假设简化','数学模型 / 计算模型'],'模型求解':['程序求解','参数 / 量纲'],'模型验证':['边界情况','对照结果 / 结论边界']},
 '工程建模完整链路；假设、求解与验证之间的关系。','把开放工程描述转化为可计算问题，并避免用“程序能运行”替代模型正确性验证。',
 '给出“预测一辆小车10 s后的位姿”问题：信息只给初始位置与速度，不给采样周期、转向模型和单位，要求学生指出为什么不能立即写代码。',
 [('建模链路','用“工程问题—假设—方程/规则—程序—验证—解释”链路拆解机械臂、小车和温度数据案例，明确模型不是现实本身。'),
  ('建模卡片','学生完成一个简化工程问题的对象、输入、输出、变量、约束、假设和验收指标卡片。'),
  ('验证与伦理','比较量纲检查、边界情况、已知解/对照实验；辨析选择性展示、删改原始数据和未经核验AI建议等风险。')],
 '一页“工程建模问题拆解卡”＋验证清单。','问题定义、假设清单、变量/单位表、至少3条验证方法、修订记录。','平时表现＋课程作业（第一单元10%模块）＋期末综合项目；目标1.1、2.1、3.1、3.3。',
 ['“模型假设”与“计算参数”有什么区别？','为什么量纲正确仍不能证明模型一定正确？','程序运行成功为什么不是工程结论成立的充分条件？'],
 '选择一个日常工程对象（温度、机构、运动平台等），补全问题定义—假设—模型—验证四栏表。',
 '结合工程计算中的模型失效与数据责任，强调结论必须建立在真实数据、合理模型和充分验证上。',
 '把“问题拆解卡＋验收指标”视为技术产品需求说明的最小原型，训练从工程问题到可验证计算任务的转化。',
 '电脑、投影仪、白板/关系图工具、课程示例数据。',REFS['main']+'\n'+REFS['sci']),

L(2,'第二单元 Python编程基础与计算逻辑','Python环境、数据类型与控制结构','理论','理论2课时',
 {'k':['掌握Python程序运行、基本数据类型、列表/元组/字典等组合数据、表达式、分支与循环。','理解变量命名、缩进、可读性和基础调试的意义。'],
  'a':['能够阅读并手工推演基础Python程序，预测输出和循环次数。','能够把简单计算规则实现为可运行脚本并定位常见语法/逻辑错误。'],
  'v':['形成规范命名、主动验证和对开源工具使用负责的编程习惯。']},
 {'运行环境':['Python解释器','脚本 / 交互环境'],'数据表达':['数值 / 字符串','列表 / 元组 / 字典'],'控制逻辑':['if','for / while'],'程序质量':['命名 / 缩进','调试 / 边界输入']},
 'Python基础语法与控制结构；程序阅读和预测。','区分语法正确与逻辑正确；处理边界输入、循环终止和数据类型变化。',
 '展示一个“计算0~n总和”的脚本，分别给n=10、0、-1，让学生预测输出和可能的问题。',
 [('环境与表达','演示解释器、脚本执行、变量、数值/字符串/列表/字典；强调单位和数据含义不能只藏在变量名里。'),
  ('控制结构','用分段函数、阈值报警和累加计算讲if/for/while；学生先手工推演再运行。'),
  ('错误复盘','分析缩进、类型不匹配、循环不终止、边界条件遗漏等错误，并用最小输入验证修正。')],
 '三个可运行Python小程序＋一张输入—预期—实际结果表。','源代码、运行结果、错误信息、边界测试和修订记录。','平时表现＋课程作业（第二单元20%模块）；目标1.2、2.2、3.4、3.5。',
 ['列表与元组在本课程常见用途上有何区别？','for与while分别适合哪类重复过程？','怎样判断一个循环不是“碰巧对一个输入有效”？'],
 '完成3个基础计算题，并为每个程序增加正常、边界和异常输入测试。',
 '通过开源生态、许可证与代码责任案例说明“可以复用”不等于“可以不理解”，最终采用的代码需要人工核验。',
 '把脚本视为工程计算原型，要求输入、输出、异常和运行说明清晰，为后续模块化和自动化计算打底。',
 '电脑、Python 3.x、VS Code/Jupyter或学校统一环境、投影仪。',REFS['main']+'\n'+REFS['python']),

L(3,'第二单元 Python编程基础与计算逻辑','函数、模块、文件与异常处理','理论','理论2课时',
 {'k':['掌握函数参数/返回值、模块导入、文本/CSV基础读写和try/except异常处理。','理解程序分解、接口与复用的基本思想。'],
  'a':['能够把重复计算重构为函数，并用最小测试验证参数和返回值。','能够读取简单数据文件、处理文件/数据异常，并组织为多函数脚本。'],
  'v':['形成接口清晰、错误显式、代码可读和可复现的工程习惯。']},
 {'函数':['参数','返回值 / 作用域'],'模块':['import','标准库 / 第三方库'],'文件':['读取 / 写入','路径 / 编码'],'异常':['try / except','失败信息 / 恢复策略']},
 '函数化与模块化；文件I/O；异常处理。','为函数划清输入输出和失败条件，避免用宽泛except掩盖真正错误。',
 '展示一段把读取文件、计算平均值、绘图全部写在main中的脚本，要求学生指出复用、测试和错误定位上的问题。',
 [('函数与接口','以“计算均方根/单位换算”为例讲函数参数、返回值、默认值和最小测试。'),
  ('文件与模块','读取简化CSV/文本，讲路径、编码和上下文管理；演示random/math等模块与第三方库导入。'),
  ('异常与重构','处理文件不存在、非法数值和空数据；把程序拆为load_data、compute、report等职责函数。')],
 '一个模块化计算脚本＋文件读写样例＋异常测试。','函数接口说明、测试输入输出、文件样例、异常信息与处理理由。','平时表现＋课程作业（第二单元20%模块）；目标1.2、2.2、3.4、3.5。',
 ['函数的输入/输出为什么是接口契约？','为什么不建议使用except:直接吞掉所有错误？','如何让数据文件路径和程序运行方式更可复现？'],
 '把上一课一个脚本重构为不少于3个函数，并添加一条文件异常测试和README运行说明。',
 '强调错误必须被看见、被记录、被解释；不能通过隐藏异常制造“程序正常”的假象。',
 '将函数/模块组织作为小型计算工具产品化第一步，训练接口、复用和运行说明。',
 '电脑、Python、VS Code/Jupyter、示例CSV/文本文件。',REFS['main']+'\n'+REFS['python']),

L(4,'上机一 Python基础与蒙特卡罗圆周率计算','蒙特卡罗圆周率：环境、随机采样与基础实现','上机','上机2课时（上机一前2/3）',
 {'k':['理解蒙特卡罗随机采样估计圆周率的几何思想及采样点判定条件。','掌握random/NumPy随机数的基本生成方式和循环/计数实现。'],
  'a':['能够完成可运行的圆周率估计程序，并记录采样数、估计值、误差与运行时间。','能够构造不同采样规模进行对照并解释随机波动。'],
  'v':['形成不选择性只报告“最好一次”结果、完整保留实验参数与结果的实证意识。']},
 {'几何模型':['单位正方形','四分之一圆'],'随机采样':['x,y随机点','命中判定'],'估计计算':['命中率','π估计'],'实验记录':['采样规模','误差 / 时间 / 随机波动']},
 '蒙特卡罗模型与随机采样程序；实验记录。','理解随机结果波动与采样规模关系；避免把单次偶然结果解释成稳定规律。',
 '先给出n=100的一次“非常接近π”的结果，再给出n=10000的稍差结果，要求学生判断能否据此说“小样本更准”。',
 [('任务建模','写出几何概率模型、输入输出、命中条件和误差定义；确认环境可运行。'),
  ('独立实现','学生用循环实现随机点生成、命中计数和π估计；记录n=100、1000、10000的结果。'),
  ('对照与排错','重复运行同一n，观察随机波动；检查坐标范围、平方和条件、计数和除法边界。')],
 'monte_carlo_pi_v1.py＋采样结果记录表。','源代码、至少3种n×多次运行结果、绝对误差、运行时间、1个失败/修正案例。','课程作业（第二单元20%模块，上机一前段）＋期末综合项目方法基础；目标1.2、1.3、2.2、3.1、3.5。',
 ['为什么随机采样结果每次不同？','怎样定义本实验的误差？','一次很准的结果是否代表算法在该n下稳定？'],
 '补做同一n至少5次重复实验，计算平均误差并准备下一课比较实现方案。',
 '以随机实验的波动说明工程结论要依赖重复验证和完整记录，不能选择性展示最有利结果。',
 '把蒙特卡罗脚本视为可配置计算工具，开始引入参数、结果表和可复现实验记录。',
 '机房、Python、VS Code/Jupyter、计时工具。',REFS['main']+'\n'+REFS['sci']),

L(5,'上机一后1课时＋第三单元理论起始','蒙特卡罗误差复盘与NumPy数组入门','理实一体','上机1课时＋理论1课时',
 {'k':['理解采样规模、误差和运行时间之间的基本关系。','掌握NumPy ndarray、shape、dtype、数组创建和基础索引概念。'],
  'a':['能够整理蒙特卡罗多次实验并用数据而非单次结果下结论。','能够把Python列表数据转换为NumPy数组并检查形状/类型。'],
  'v':['形成从实验复盘过渡到工具升级时仍保留验证证据和边界说明的习惯。']},
 {'实验收口':['重复实验','误差 / 时间'],'结果解释':['随机波动','结论边界'],'NumPy数组':['ndarray / shape','dtype / 创建'],'数组访问':['索引','切片初步']},
 '蒙特卡罗实验结论的证据链；NumPy数组的结构。','在一节混合课中完成上机项目收口并建立从Python容器到数值数组的认知迁移。',
 '把学生上机一结果汇总成表，提问：若平均误差下降但运行时间上升，应如何描述“更好”？随后展示同一批数据转成ndarray后的shape/dtype。',
 [('上机一验收','学生整理多次实验，计算平均/最大误差和运行时间，补全README；教师检查是否选择性删结果。'),
  ('NumPy概念迁移','从list到ndarray，讲shape、dtype、zeros/arange/linspace及索引切片，解释数值数组为何适合科学计算。'),
  ('最小练习','把实验结果表转为数组，完成列选取、误差向量和基础统计，为向量化做准备。')],
 '上机一完整小报告＋NumPy数组练习脚本。','多次实验结果表、结论文字、README、数组shape/dtype/索引输出。','课程作业（第二单元20%模块完成）＋第三单元30%模块起始；目标1.2、1.3、2.2、3.1、3.5。',
 ['如何用重复实验避免“只看最好一次”？','ndarray的shape和dtype分别描述什么？','为什么科学计算中通常要关注数组形状？'],
 '提交上机一完整包；完成NumPy数组创建/索引练习并记录3个常见shape错误。',
 '强调实验报告不得删去不利结果；结论必须和记录一致，工具升级不能替代证据意识。',
 '把实验结果数组化，为后续参数扫描、性能比较和自动化报告建立可扩展数据结构。',
 '机房、Python、NumPy、VS Code/Jupyter。',REFS['main']+'\n'+REFS['sci']),

L(6,'第三单元 科学计算与工程数据处理','NumPy向量化、广播、矩阵运算与SciPy数值方法','理论','理论2课时',
 {'k':['掌握数组索引切片、广播、向量化、矩阵乘法与常见线性代数操作。','了解SciPy数值积分、插值、方程求解和优化的基本用途及“工具有适用条件”的原则。'],
  'a':['能够比较循环与向量化计算结构并判断数组shape是否兼容。','能够为一个简单数值问题选择合适的SciPy工具并设计结果校核。'],
  'v':['形成关注浮点误差、计算效率和工具适用边界的科学计算意识。']},
 {'数组计算':['索引 / 切片','广播 / 向量化'],'矩阵运算':['点积 / 矩阵乘法','转置 / 线性代数'],'数值方法':['积分 / 插值','求根 / 优化'],'数值质量':['浮点误差','性能 / 结果校核']},
 '广播与向量化；矩阵运算；SciPy方法选择。','理解广播shape规则和数值函数输入条件；避免“库函数返回数值就认为正确”。',
 '比较同一个100万点运算的for循环与向量化版本，让学生先预测哪个更快、结果是否应完全相同。',
 [('向量化与广播','用批量单位换算、距离计算展示广播；学生画shape变化图并判断可否广播。'),
  ('矩阵与误差','讲矩阵乘法、转置、线性代数基本接口；用浮点比较和条件数/误差现象强调数值结果需要校核。'),
  ('SciPy用途地图','用积分、插值、求根、优化四个最小案例说明函数接口、初值/区间/边界等条件，要求学生写验证方案。')],
 '“循环 vs 向量化”对比卡＋SciPy方法选择表。','shape推演、运行时间对照、结果差值、SciPy输入条件和验证思路。','平时表现＋课程作业（第三单元30%模块）；目标1.3、2.2、3.1。',
 ['广播的核心判断是什么？','为什么浮点数不宜总用==判断理论相等？','数值求根为什么要关心初值/区间？'],
 '完成3组循环—向量化改写；为积分/求根各写一条独立校核方式。',
 '通过数值误差和性能证据强调“快”和“对”都需要测量，不能凭经验或工具输出做结论。',
 '把向量化、矩阵计算和SciPy视为工程计算服务的核心计算层，训练选择工具而非堆叠工具。',
 '电脑、Python、NumPy、SciPy、计时工具。',REFS['main']+'\n'+REFS['sci']),

L(7,'第三单元理论后1课时＋上机二前1课时','Pandas数据流程与NumPy向量化性能实验','理实一体','理论1课时＋上机1课时',
 {'k':['掌握Pandas DataFrame/Series、读取、类型检查、缺失值/异常值、筛选、分组和连接的基本流程。','理解循环与向量化性能对比实验的测量方法。'],
  'a':['能够检查工程数据的字段、类型、缺失和异常，提出清洗策略。','能够完成一个循环/向量化对照实验并记录时间与结果一致性。'],
  'v':['形成原始数据留存、清洗可追溯和性能比较需控制变量的意识。']},
 {'数据对象':['Series','DataFrame'],'数据检查':['字段 / dtype','缺失 / 异常'],'数据变换':['筛选 / 转换','groupby / merge'],'性能实验':['循环','向量化 / 时间 / 一致性']},
 'Pandas工程数据流程；性能实验设计。','同时处理“数据质量”和“计算性能”两个维度，避免为了效率牺牲语义或为了清洗破坏原始数据。',
 '给出一份包含空值、异常温度和重复设备编号的CSV，以及一段慢循环代码，要求学生分别指出“数据问题”和“实现问题”。',
 [('Pandas流程','演示read_csv、info、isna、筛选、groupby、merge；强调原始数据只读保存和清洗日志。'),
  ('性能实验启动','学生完成一个数组计算的循环与向量化两版，统一输入、重复测时，并比较结果误差。'),
  ('数据质量练习','学生对示例工程数据生成质量报告：缺失、异常、重复、单位/类型问题及处理理由。')],
 '性能对照脚本＋数据质量检查表。','原始数据副本、清洗日志、循环/向量化运行时间、结果一致性检查。','课程作业（第三单元30%模块，上机二前段）；目标1.3、1.4、2.2、2.3、3.1、3.3。',
 ['为什么清洗前要保留原始数据？','循环和向量化比较怎样做到“同一问题同一输入”？','异常值应删除、修正还是保留由什么决定？'],
 '完善工程数据质量报告；为下节课准备分组统计、表连接和结果导出任务。',
 '强调数据处理过程透明，禁止为了得到预期结论随意删除异常数据。',
 '将数据质量报告视为分析产品的“数据验收层”，将性能对照作为计算模块选型证据。',
 '机房、Python、NumPy、Pandas、CSV工程数据。',REFS['main']+'\n'+REFS['sci']),

L(8,'上机二 NumPy矩阵运算与Pandas工程数据处理','工程数据清洗、分组统计、表连接与结果导出','上机','上机2课时（上机二中段）',
 {'k':['掌握Pandas缺失/异常处理、筛选、派生列、groupby聚合、merge连接和结果导出。','巩固NumPy数组与矩阵运算在批量计算中的应用。'],
  'a':['能够完成“读取—检查—清洗—转换—汇总—连接—导出”工程数据流程。','能够用小样本或独立查询核对关键统计结果。'],
  'v':['形成数据来源、处理步骤、单位和结果可复核的工程数据治理习惯。']},
 {'数据清洗':['缺失值','异常 / 重复'],'数据变换':['筛选 / 派生列','单位统一'],'统计汇总':['groupby / agg','排序 / Top-N'],'表连接':['merge','键 / 基数 / 导出']},
 'Pandas清洗、聚合与连接；结果核验。','merge键/基数导致的重复扩展；清洗策略对结论的影响。',
 '提供“设备运行表＋设备台账表”，其中一张表设备ID重复，要求学生先预测merge后行数，再运行验证。',
 [('任务拆解','明确原始数据只读、清洗副本、字段单位和输出要求；先写处理流程图。'),
  ('独立实现','完成缺失/异常处理、派生指标、groupby统计、两表merge和CSV结果导出。'),
  ('结果验收','核对行数、连接基数、单位和3个关键统计；保留处理前后对照和不能确认的数据。')],
 'data_process.py＋clean_data.csv＋summary.csv＋清洗日志。','原始/清洗数据、代码、行数/字段检查、groupby结果、merge前后基数、独立复核记录。','课程作业（第三单元30%模块，上机二中段）；目标1.3、1.4、2.2、2.3、3.1、3.3。',
 ['merge后行数突然增加可能有哪些原因？','为什么“填充缺失值”必须说明依据？','怎样证明一个groupby结果没有分错组？'],
 '补充一份数据字典，说明字段单位、缺失处理、异常规则和连接键。',
 '通过字段口径和清洗日志强化数据真实性与可追溯性，无法确认的数据不强行“修正”。',
 '将清洗后的数据与统计结果作为后续可视化/模型输入的数据产品，训练可复用数据管道。',
 '机房、Python、NumPy、Pandas、工程CSV数据。',REFS['main']+'\n'+REFS['sci']),

L(9,'上机二后1课时＋第四单元理论起始','数据处理验收与Matplotlib工程可视化入门','理实一体','上机1课时＋理论1课时',
 {'k':['理解数据处理验收需要检查字段、单位、缺失、连接行数和统计口径。','掌握Matplotlib figure/axes、折线/散点/柱状图和坐标轴/单位/图例的基本结构。'],
  'a':['能够收口上机二并给出可复核的数据处理结论。','能够根据时间序列、关系比较、类别比较选择基础图表并规范标注。'],
  'v':['形成图表服务于工程问题、不得用视觉技巧歪曲数据的表达责任。']},
 {'上机验收':['清洗日志','统计 / merge复核'],'绘图对象':['Figure','Axes'],'图表选择':['折线 / 散点','柱状图'],'规范表达':['坐标 / 单位','图例 / 标题 / 来源']},
 '数据处理闭环；Matplotlib对象模型与图表选择。','把“好看”与“表达正确”分开，理解图表选择必须服从数据类型和分析问题。',
 '展示同一组运行数据的折线图、柱状图、散点图，提问哪个更适合“随时间变化”和“两个变量关系”，并指出缺单位图表的问题。',
 [('上机二验收','学生按数据字典复查清洗、统计和merge结果，形成上机二README与结论边界。'),
  ('Matplotlib入门','讲Figure/Axes、plot/scatter/bar、xlabel/ylabel/legend/title和保存图像。'),
  ('图表选择练习','给出3类工程问题，让学生选择图表并说明理由；辨析截断坐标、面积错觉和过度装饰。')],
 '上机二完整交付＋3张规范基础图及图表选择说明。','README、数据字典、统计复核、绘图脚本、PNG图、单位/图例检查表。','课程作业（第三单元30%模块完成）＋第四单元20%模块起始；目标1.4、1.5、2.3、3.1、3.3、3.5。',
 ['折线图与散点图分别强调什么关系？','工程图为什么必须标单位？','坐标轴截断何时可能误导？'],
 '选择上一课summary.csv中的3个指标，分别制作适当图表并写一句“图表支持的结论/不支持的结论”。',
 '通过误导性图表反例强调客观表达与受众责任，不用视觉设计隐藏不利数据。',
 '将数据处理结果转为可读“工程信息产品”，训练图表作为决策沟通界面。',
 '机房、Python、Pandas、Matplotlib、课程工程数据。',REFS['main']+'\n'+REFS['sci']),

L(10,'第四单元 数据可视化设计与工程表达','Matplotlib规范绘图、子图、等值图与基础三维表达','理论','理论2课时',
 {'k':['掌握坐标轴范围、刻度、单位、图例、标注、颜色、子图布局和图像保存。','了解等值图、热力表达和基础三维线/散点/曲面图的适用场景。'],
  'a':['能够把同一数据按工程目的设计成信息层次清晰的图表组。','能够识别比例失真、颜色误导、信息过载和3D滥用。'],
  'v':['形成准确、简洁、可复核的工程可视化规范。']},
 {'图表结构':['轴 / 刻度 / 单位','图例 / 标注'],'布局':['subplot','共享轴 / 版式'],'场图表达':['等值图 / 热力图','基础3D'],'质量边界':['颜色 / 比例','信息密度 / 可复核']},
 '规范标注与布局；等值/三维图适用性。','在信息完整和图面简洁之间权衡；避免为了“高级感”使用不必要3D。',
 '给出一个3D柱状图和一个2D条形图展示相同类别数据，让学生比较哪一个更容易精确判断差异。',
 [('规范元素','系统讲轴、单位、图例、标注、网格、保存分辨率和中文/符号显示。'),
  ('多图布局','用位置—速度—误差三子图展示共享时间轴；讨论布局、对齐与比较任务。'),
  ('场与三维','介绍contour/contourf和mplot3d基础；比较何时3D提供空间结构、何时只是视觉负担。')],
 '一张“工程图表规范检查表”＋一组图表重构草案。','错误图找错记录、重构前后对比、图表用途/单位/比例说明。','平时表现＋课程作业（第四单元20%模块）；目标1.5、2.3、3.3、3.5。',
 ['什么时候子图比叠加多条曲线更清晰？','为什么3D图不一定比2D图信息更多？','图例、单位、数据来源中哪些属于可复核性信息？'],
 '找一张公开工程图（不涉及个人/敏感数据），按规范检查表分析3个优点和3个可改进点。',
 '强调图表既是技术表达，也是责任文件；坐标、单位和视觉编码不能误导受众。',
 '把可视化规范转化为报告/看板设计标准，为后续机构运动学和轨迹仿真提供统一输出规范。',
 '电脑、Python、Matplotlib、示例二维/三维图。',REFS['main']+'\n'+REFS['sci']),

L(11,'第四单元理论后1课时＋上机三前1课时','坐标变换与两连杆机械臂正运动学建模','理实一体','理论1课时＋上机1课时',
 {'k':['理解平面坐标系、关节角、连杆长度和两连杆正运动学关系。','理解末端位置、构型、轨迹和工作空间的区别。'],
  'a':['能够从机械臂几何关系推导并实现末端坐标计算函数。','能够用特殊姿态和长度边界测试正运动学模型。'],
  'v':['形成尊重物理约束、先推导/验算再扩大数值采样的建模习惯。']},
 {'几何对象':['基坐标系','两连杆 / 关节角'],'正运动学':['x(θ1,θ2)','y(θ1,θ2)'],'模型实现':['函数接口','向量化输入'],'验证':['零位 / 共线姿态','长度边界 / 量纲']},
 '两连杆正运动学模型；特殊姿态验证。','坐标系和角度定义必须一致；公式、程序和图形三者要互相校验。',
 '画出两连杆机械臂并给出θ1=θ2=0，先让学生不用代码判断末端应在哪里，再运行程序检验。',
 [('理论推导','建立基坐标系，推导x=l1 cosθ1+l2 cos(θ1+θ2)、y=l1 sinθ1+l2 sin(θ1+θ2)等关系，明确角度与单位。'),
  ('函数实现','学生实现forward_kinematics(theta1,theta2,l1,l2)，支持标量输入，输出末端位置。'),
  ('特殊姿态测试','测试0°、90°、折叠/伸直等姿态，比较手工几何判断和程序结果。')],
 '两连杆正运动学推导页＋forward_kinematics.py＋特殊姿态测试表。','公式、变量/单位定义、代码、手算/程序对照、失败修正记录。','课程作业（第四单元20%模块，上机三起始）＋期末综合项目；目标1.1、1.3、1.5、2.1、2.4、3.1、3.5。',
 ['θ2是相对第一杆还是绝对角度会造成什么差别？','什么特殊姿态最适合做模型单元测试？','如果图画对了但数值不对，应先检查哪些定义？'],
 '补全至少5组特殊姿态测试；准备下一课绘制机械臂构型和末端轨迹。',
 '通过机械臂物理边界和安全姿态强调模型必须尊重真实约束，不得用图形“遮住”错误数值。',
 '把正运动学函数设计成可复用计算模块，作为机器人仿真/数字样机的最小计算服务。',
 '机房、Python、NumPy、Matplotlib、白板/公式推导。',REFS['main']+'\n'+REFS['sci']),

L(12,'上机三 平面连杆机械臂运动学与工作空间可视化','机械臂构型、末端轨迹与参数化仿真','上机','上机2课时（上机三中段）',
 {'k':['巩固正运动学、角度序列、末端轨迹与机械臂构型绘制。','掌握用NumPy生成关节角序列和用Matplotlib同步表达构型/轨迹。'],
  'a':['能够完成关节轨迹→末端轨迹计算与可视化。','能够比较不同连杆长度/角度范围对轨迹的影响并说明模型边界。'],
  'v':['形成参数变化必须保留配置、图表和解释依据的仿真实验习惯。']},
 {'关节输入':['θ1序列','θ2序列'],'批量计算':['向量化','末端x,y'],'构型绘制':['基座 / 关节 / 连杆','等比例坐标'],'轨迹分析':['路径','参数变化 / 边界']},
 '批量正运动学与轨迹可视化。','保证坐标等比例、单位一致和角度弧度转换正确；区分机械臂瞬时构型与末端轨迹。',
 '展示一张因坐标轴比例不一致而“看起来连杆长度变化”的机械臂图，让学生判断是模型错还是画图错。',
 [('任务设置','定义l1/l2和θ1(t)、θ2(t)序列，明确参数单位、采样数量和验收姿态。'),
  ('独立仿真','学生用NumPy批量计算末端位置，绘制若干构型、末端轨迹和关键点标注。'),
  ('参数对照','改变连杆长度或角度范围，比较轨迹；使用axis equal和特殊姿态验证图形/数值一致。')],
 'robot_arm_trajectory.py＋构型图＋末端轨迹图＋参数对照表。','代码、参数配置、关键姿态数值、等比例图、参数变化前后图和解释。','课程作业（第四单元20%模块，上机三中段）＋期末综合项目；目标1.3、1.5、2.4、3.1、3.5。',
 ['为什么机械臂图通常要使用等比例坐标？','末端轨迹和工作空间有什么本质差别？','怎样证明轨迹图中的关键点与数值一致？'],
 '增加一组边界关节角参数并记录末端轨迹变化；准备工作空间采样。',
 '通过图形比例失真案例强调工程图不能制造错误直觉，数值与图形必须交叉验证。',
 '把参数化机械臂仿真视为数字样机雏形，训练“配置—计算—图形—验收”可复用流程。',
 '机房、Python、NumPy、Matplotlib。',REFS['main']+'\n'+REFS['sci']),

L(13,'上机三 平面连杆机械臂运动学与工作空间可视化','机械臂工作空间采样、边界与模型验收','上机','上机2课时（上机三后段）',
 {'k':['理解工作空间的数值采样表示、可达边界和姿态/采样密度对结果的影响。','理解“数值点云近似工作空间”与理论连续集合之间的区别。'],
  'a':['能够在给定关节范围内进行二维网格/随机采样并绘制工作空间。','能够用理论半径边界、特殊姿态和采样密度对结果进行合理性验证。'],
  'v':['形成说明采样近似、模型假设和结论边界的工程表达意识。']},
 {'工作空间':['可达集合','角度约束'],'数值采样':['网格 / 随机','采样密度'],'边界检查':['最大/最小半径','特殊姿态'],'结果表达':['散点 / 密度','近似性 / 模型假设']},
 '工作空间数值采样与边界验证。','避免把采样点云当作精确连续边界；识别关节限位、连杆长度和采样密度的影响。',
 '先展示低密度采样得到“有很多洞”的工作空间，让学生判断洞是机构不可达还是采样不足。',
 [('采样设计','确定θ1/θ2范围与采样策略；先估计理论最大/最小可达半径。'),
  ('独立实现','批量计算工作空间点，绘制散点/密度图，标出基座和理论参考圆。'),
  ('模型验收','比较不同采样密度和角度限位；检查最大半径、折叠半径、边界姿态并总结模型假设。')],
 'workspace.py＋工作空间图＋采样/边界验收报告。','采样参数、工作空间点数、理论/数值边界对照、密度对比图、假设与局限说明。','课程作业（第四单元20%模块完成）＋期末综合项目；目标1.1、1.3、1.5、2.1、2.4、3.1、3.3、3.5。',
 ['采样密度增加会改变真实工作空间吗？','理论最大可达半径如何用于校验？','关节限位应放在模型的哪一部分表达？'],
 '整理上机三完整包：公式、代码、轨迹、工作空间、验证和模型边界；写出两条“模型未覆盖的真实因素”。',
 '强调计算模型的边界说明与安全约束，不能把数值采样近似表述为绝对真实。',
 '将工作空间图作为机器人选型/布置的早期分析产品，训练从数值模型到工程决策支持的表达。',
 '机房、Python、NumPy、Matplotlib。',REFS['main']+'\n'+REFS['sci']),

L(14,'第五单元 综合建模流程与前沿工具边界','综合工程计算项目、版本控制与AI辅助编程边界','理论','理论2课时',
 {'k':['理解需求分解、模型/数据/代码/参数组织、测试、版本记录和可复现交付的项目流程。','理解开源库文档、AI辅助编程的作用、常见错误类型和人工验证责任。'],
  'a':['能够为双轮小车综合任务设计目录、接口、测试用例、数据记录和验收清单。','能够区分“AI建议可参考”与“工程结论必须由人验证”。'],
  'v':['形成团队协作、版本可追溯、负责任使用AI和高质量工程交付意识。']},
 {'项目组织':['需求 / 验收','目录 / 接口'],'版本与复现':['依赖 / 参数','版本 / README'],'AI辅助':['代码建议','解释 / 调试 / 风险'],'验证责任':['测试','人工审查 / 结论边界']},
 '综合项目组织与可复现性；AI辅助工具边界。','把工具产出转化为可验证工程成果，而不是把AI/开源库当作正确性来源。',
 '给出一段AI生成的双轮小车更新公式，其中左右轮速度符号被写反，但代码可运行，要求学生设计最少测试发现问题。',
 [('项目结构','用双轮小车任务拆解requirements、model、simulation、data、plot、tests、README，明确接口和验收。'),
  ('版本与复现','说明依赖、随机种子/参数、文件路径和版本记录；示范“别人从空目录能否运行”的验收思路。'),
  ('AI与开源工具边界','总结幻觉API、单位错误、符号错误、边界遗漏、过拟合解释等风险，要求AI建议必须经过阅读、测试和人工修改记录。')],
 '双轮小车综合项目设计书＋测试/验收清单＋AI使用记录模板。','项目目录图、接口表、依赖/参数清单、测试用例、AI建议采纳/拒绝理由。','平时表现＋课程作业（第五单元20%模块）＋期末综合项目60%；目标1.6、2.1、2.5、3.2、3.4、3.5。',
 ['怎样定义“别人可以复现我的程序”？','AI生成代码最少需要哪些人工核验？','为什么版本记录是工程证据而不只是管理习惯？'],
 '搭建双轮小车项目目录和README骨架；写出直行、原地转动、零速度等基础验收用例。',
 '强调能力越强的自动化工具越需要清晰责任边界，最终工程判断和交付责任不能外包给AI。',
 '将项目目录、README、测试和版本记录作为可交付软件产品的组成部分，训练从“代码作业”向“工程项目”升级。',
 '电脑、Python、Git/等价版本工具、课程示例项目。',REFS['main']+'\n'+REFS['sci']+'\n'+REFS['python']),

L(15,'上机四 双轮小车轨迹仿真与综合工程表达','双轮小车离散运动学、状态更新与数据记录','上机','上机2课时（上机四前半）',
 {'k':['理解双轮小车状态(x,y,θ)、左右轮速度、轮距、离散时间步长与运动学更新。','掌握状态日志、控制量和误差数据的Pandas记录结构。'],
  'a':['能够实现直行、原地转动、圆弧运动等基本状态更新并通过解析/几何预期验证。','能够记录每步状态、输入和时间并输出可复核数据表。'],
  'v':['形成状态模型、单位、时间步长和测试用例必须显式记录的仿真质量意识。']},
 {'运动学状态':['x / y / θ','左右轮速度 / 轮距'],'离散更新':['v / ω','Δt / 状态方程'],'基础用例':['直行','原地转动 / 圆弧'],'数据记录':['时间 / 状态','输入 / Pandas日志']},
 '双轮小车离散运动学与基础测试；状态记录。','符号约定、角度单位和Δt会直接影响轨迹；必须用可预期基础运动验证。',
 '给出vL=vR、vL=-vR、vL=0/vR>0三组输入，让学生不运行代码先画预期运动趋势。',
 [('模型确认','定义坐标系、θ正方向、轮距L、速度单位和Δt；写出v=(vR+vL)/2、ω=(vR-vL)/L及离散更新。'),
  ('独立实现','实现step(state,control,dt)和simulate函数，完成直行/原地转动/圆弧用例。'),
  ('数据记录','用Pandas记录time,x,y,theta,vL,vR等列，检查步数、时间和最终状态；保留失败修正。')],
 'differential_drive.py＋test_basic_motion.py＋trajectory_log.csv。','公式/坐标约定、代码、3类基础用例预期与实际、状态日志、错误修订记录。','课程作业（第五单元20%模块，上机四前半）＋期末综合项目60%；目标1.4、1.6、2.3、2.5、3.1、3.5。',
 ['vL=vR时角速度应为什么？','Δt减半且总时间不变，仿真点数会怎样变化？','如何用原地转动用例发现左右轮符号错误？'],
 '补充一个时间步长敏感性对照；为下一课准备目标轨迹或给定速度序列。',
 '通过基础运动验收强调自动化仿真必须先通过可解释测试，不能直接相信动画效果。',
 '把状态更新函数和日志设计为可替换模块，为移动机器人轨迹服务/数字孪生原型提供基础。',
 '机房、Python、NumPy、Pandas、Matplotlib、版本管理工具。',REFS['main']+'\n'+REFS['sci']),

L(16,'上机四 双轮小车轨迹仿真与综合工程表达','轨迹跟踪、误差可视化与综合项目验收','上机','上机2课时（上机四后半）',
 {'k':['理解目标轨迹/给定速度、状态误差、轨迹与误差可视化和综合项目验收逻辑。','理解期末综合项目对模型、代码、数据、图表、版本和个人解释的整体要求。'],
  'a':['能够完成简单轨迹跟踪或给定速度仿真，记录误差并制作规范轨迹/误差图。','能够按README从项目源文件复现结果，并现场解释一个模型决策、一个测试和一个AI/工具核验记录。'],
  'v':['形成可复现、可解释、客观表达、团队贡献可追溯的最终工程交付习惯。']},
 {'轨迹任务':['目标 / 控制输入','状态更新'],'误差分析':['位置误差','方向误差 / 指标'],'可视化':['轨迹图','误差-时间图 / 标注'],'综合验收':['README / 依赖','测试 / 版本 / 个人解释']},
 '轨迹与误差表达；综合项目复现与验收。','避免只展示“漂亮轨迹”而不展示误差、失败和参数；个人必须能解释本人负责内容。',
 '教师展示两份结果：A轨迹图很漂亮但无单位/误差/参数；B图较朴素但有目标/实际轨迹、误差曲线、参数和测试。让学生按工程验收判断哪份证据更完整。',
 [('轨迹/误差实现','完成给定速度或简单跟踪逻辑，记录目标与实际状态，计算位置/方向误差。'),
  ('工程可视化','制作目标/实际轨迹、误差—时间图和关键参数标注；核对坐标比例、单位、图例和结论。'),
  ('最终验收','按README从干净环境/新目录运行；检查代码、数据、参数、依赖、测试、图表和版本记录；个人说明模型假设、一个失败案例及AI/开源工具核验。')],
 '上机四完整项目＋期末综合项目可验收版本：代码、数据、图表、README、测试与个人说明。','运行日志、轨迹/误差图、参数配置、测试结果、版本记录、AI工具核验记录、个人现场解释。','课程作业（第五单元20%模块完成）＋期末综合项目60%核心验收；目标1.4、1.5、1.6、2.3、2.5、3.2、3.3、3.4、3.5。',
 ['轨迹图之外为什么还需要误差曲线？','怎样证明项目在另一环境/新目录中可复现？','如果AI建议的控制公式能运行但基础测试失败，应如何处理？'],
 '整理课程最终作品与学习反思；在不预填课堂事实的前提下，根据实际项目记录总结模型边界、失败经验和后续可改进点。',
 '结合移动平台安全与AI辅助工具使用，强调工程人员对模型、软件和结果承担最终验证责任。',
 '将最终项目作为机器人/AI/数据类后续课程的可复用计算与可视化基座，形成个人工程作品雏形。',
 '机房、Python、NumPy、Pandas、Matplotlib、Git/等价工具。',REFS['main']+'\n'+REFS['sci']+'\n'+REFS['python'])
]
assert len(lessons)==16
# exact hour closure: each session has 2 nominal class hours; split labels below encode 16 theory + 16 lab.
THEORY_HOURS=sum(2 if l['split']=='理论2课时' else 1 if '理论1课时' in l['split'] else 0 for l in lessons)
LAB_HOURS=sum(2 if l['split'].startswith('上机2课时') else 1 if '上机1课时' in l['split'] else 0 for l in lessons)
assert THEORY_HOURS==16,(THEORY_HOURS,LAB_HOURS)
assert LAB_HOURS==16,(THEORY_HOURS,LAB_HOURS)

COURSE_MAP={
 '工程建模':['问题定义 / 假设','模型 / 求解 / 验证'],
 'Python基础':['数据类型 / 控制','函数 / 模块 / 文件 / 异常'],
 '科学计算与数据':['NumPy / SciPy','Pandas / 数据质量'],
 '可视化与机构':['Matplotlib / 图表规范','机械臂运动学 / 工作空间'],
 '综合仿真与工程化':['双轮小车状态模型','版本 / 测试 / AI工具边界']
}

def render_map(path:Path,title:str,branches:Dict[str,List[str]]):
    dot=Digraph('G',format='png')
    dot.attr(rankdir='TB',bgcolor='white',margin='0.02',pad='0.05',nodesep='0.20',ranksep='0.36',dpi='190')
    dot.attr('node',shape='box',style='rounded,filled',fontname='Noto Sans CJK SC',fontsize='11',color='#6B7280',penwidth='1.0',fillcolor='#F8FAFC',margin='0.10,0.06')
    dot.attr('edge',color='#94A3B8',arrowsize='0.55',penwidth='0.9')
    dot.node('root',title,fillcolor='#E6F1F8',color='#2F6F8F',fontsize='14',penwidth='1.4')
    for bi,(b,leaves) in enumerate(branches.items(),1):
        bid=f'b{bi}'; dot.node(bid,b,fillcolor='#EEF6F1',color='#5D846E',fontsize='12'); dot.edge('root',bid)
        for li,leaf in enumerate(leaves,1):
            nid=f'{bid}_{li}'; dot.node(nid,leaf,fillcolor='#FFFFFF',color='#A5B4C2',fontsize='10'); dot.edge(bid,nid)
    outbase=str(path.with_suffix('')); dot.render(outbase,cleanup=True)
    gen=Path(outbase+'.png')
    if gen!=path: shutil.move(gen,path)

def obj_md(g):
    def B(label,arr): return f'**{label}：**\n'+'\n'.join(f'{i+1}. {x}' for i,x in enumerate(arr))
    return '\n\n'.join([B('知识目标',g['k']),B('能力目标',g['a']),B('价值目标',g['v'])])

def process_md(l):
    if l['kind']=='上机': times=[('一、任务说明与模型/代码骨架',20),('二、学生独立/小组实现',35),('三、测试、调试与证据固化',20)]
    elif l['kind']=='理实一体': times=[('一、前段收口/新知引入',25),('二、核心推导与实现',30),('三、验证、练习与证据固化',20)]
    else: times=[('一、核心概念与工程案例',30),('二、学生推演/建模/代码阅读',25),('三、对比、验证与错误复盘',20)]
    out=[]
    for (label,mins),(title,detail) in zip(times,l['blocks']):
        if l['kind']=='上机': student='【学生】按任务约束独立/小组实现，先写模型预期再运行，保留代码、参数、错误信息、测试与图表证据。'
        elif l['kind']=='理实一体': student='【学生】完成前一任务验收并进入新知识练习；所有结论用数值、图表或测试证据支持，记录理论/上机转换点。'
        else: student='【学生】完成公式/代码/图表推演或场景判断，先写预期与依据，再用课堂示例验证并修正。'
        out.append(f'**{label}（约{mins} min）—{title}**\n\n【教师】{detail}\n\n{student}')
    out += [f'**独立学习产物**：{l["product"]}',f'**学习证据**：{l["evidence"]}',f'**评价映射**：{l["assessment"]}']
    return '\n\n'.join(out)

def preclass(l):
    if l['kind']=='上机': third='检查Python环境、上一任务源文件/数据与依赖是否可运行；不得使用未经授权的真实敏感工程数据。'
    elif l['kind']=='理实一体': third='准备上一任务验收证据，并预读本次新概念/公式/数据文件；写出一个预期结果。'
    else: third='预读一个公式、代码或工程图表片段，写出预期结果与至少一个疑问。'
    return f'''【教师】发布“{l['topic']}”对应大纲/教材范围和一个最小工程问题，不提前给出完整代码答案。\n\n1. 阅读对应大纲与教材内容；\n2. 圈出3个关键术语并记录至少1个疑问；\n3. {third}\n\n【学生】完成准备并带着“预期结果/疑问/已有证据”进入课堂。'''

def reflection(l):
    return f'''【课后填写，不预填事实】\n\n1. 教学流程：记录100 min各环节实际用时、理论/上机切换及调整点；\n2. 教学内容：重点记录“{l['diff']}”的真实掌握证据和常见错误；\n3. 学生参与：仅依据实际课堂记录填写推演、独立编码、调试、图表互审和讨论情况，不补写推测；\n4. 评价证据：检查本课代码、数据、参数、测试、图表与说明是否完整，记录缺项原因；\n5. 后续改进：依据真实课堂证据决定下一次课的补救、复现或拓展。'''

def clean_md(s):
    s=re.sub(r'!\[[^\]]*\]\([^\)]*\)','',s); s=s.replace('**','').replace('`',''); s=re.sub(r'^- ','',s,flags=re.M); s=re.sub(r'\n{3,}','\n\n',s); return s.strip()

def extract_info(s):
    d={}
    for line in s.splitlines():
        m=re.match(r'- \*\*(.+?)\*\*：(.+)',line.strip())
        if m: d[m.group(1)]=m.group(2).strip()
    return d

def parse_md(path:Path):
    text=path.read_text(encoding='utf-8'); parts=re.split(r'(?=^## 第\d+次课\s)',text,flags=re.M); parts=[p for p in parts if re.match(r'^## 第\d+次课\s',p)]
    heads=['课次信息','教学目标','课程思政','专创融合','教学重难点','教学方法与用具','教学设计','课前任务','互动导入','传授新知与课堂训练','过关检测','课堂小结','作业布置','考勤','课后教学反思','本章节参考文献']
    parsed=[]
    for p in parts:
        h=re.match(r'^## 第(\d+)次课\s+(.+)$',p,re.M); fields={}
        for i,hd in enumerate(heads):
            st=re.search(rf'^### {re.escape(hd)}\s*$',p,re.M)
            if not st: continue
            pos=st.end(); nxt=[]
            for hd2 in heads[i+1:]:
                m2=re.search(rf'^### {re.escape(hd2)}\s*$',p[pos:],re.M)
                if m2: nxt.append(pos+m2.start())
            fields[hd]=p[pos:min(nxt) if nxt else len(p)].strip()
        img=re.search(r'!\[知识脉络图\]\((images/lesson-\d+-knowledge-map\.png)\)',p); fields['image']=img.group(1) if img else None
        parsed.append({'no':int(h.group(1)),'topic':h.group(2).strip(),'fields':fields})
    return parsed,text

def remove_table(table):
    e=table._element; e.getparent().remove(e)

def set_cell_text(cell,text,font_size=10.5,bold_prefixes=None,align=None):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    if align is not None: p.alignment=align
    for i,line in enumerate(text.split('\n')):
        if i>0:p.add_run().add_break()
        r=p.add_run(line); r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(font_size)
        if bold_prefixes and any(line.strip().startswith(x) for x in bold_prefixes): r.bold=True

def set_cell_margins(cell,top=60,start=70,bottom=60,end=70):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        n=tcMar.find(qn(f'w:{m}'))
        if n is None:n=OxmlElement(f'w:{m}'); tcMar.append(n)
        n.set(qn('w:w'),str(v)); n.set(qn('w:type'),'dxa')

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); h=OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); trPr.append(h)

def sha256(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for c in iter(lambda:f.read(1<<20),b''):h.update(c)
    return h.hexdigest()

def main():
    if OUT_ROOT.exists(): shutil.rmtree(OUT_ROOT)
    IMG_DIR.mkdir(parents=True)
    render_map(IMG_DIR/'course-knowledge-map.png','工程建模与科学计算可视化基础（Python）知识主线',COURSE_MAP)
    for l in lessons: render_map(IMG_DIR/f'lesson-{l["no"]:02d}-knowledge-map.png',l['topic'],l['map'])

    md=['# 《工程建模与科学计算可视化基础（Python）》教学设计（16次课）','', '> **Canonical source / 内容权威源**：本 Markdown 与 `images/`、`manifest.json` 共同构成教案语义主源。Word 仅由本源包编译，负责学校格式呈现。','', '## 课程基本信息']
    md += [f'- **课程名称**：{COURSE["name"]}',f'- **课程英文名称**：{COURSE["english"]}',f'- **课程代码**：{COURSE["code"]}',f'- **课程类别**：{COURSE["category"]}',f'- **课程性质**：{COURSE["nature"]}',f'- **授课学期**：{COURSE["semester"]}',f'- **学分**：{COURSE["credits"]}',f'- **总学时**：32（理论16，上机16）',f'- **适用专业**：{COURSE["major"]}',f'- **授课学院**：{COURSE["college"]}',f'- **先修课程**：{COURSE["prereq"]}',f'- **后续课程**：{COURSE["followup"]}',f'- **选用教材**：{COURSE["textbook"]}',f'- **课程评价**：{COURSE["assessment"]}',f'- **内容依据**：{COURSE["version"]}现行课程教学大纲（大纲更新时间：{COURSE["syllabus_date"]}）']
    md += ['', '## 课程总体设计',
           '课程以“问题定义—假设与模型—程序求解—数据验证—图形表达”为主线，五个理论单元分别覆盖工程建模、Python基础、NumPy/SciPy/Pandas科学计算、Matplotlib工程可视化以及综合项目/版本/AI工具边界；四个上机项目分别为蒙特卡罗圆周率、NumPy矩阵与Pandas数据处理、两连杆机械臂运动学与工作空间、双轮小车轨迹仿真。由于现行大纲上机项目学时为3/4/5/4，本教案使用4次理实一体课准确闭合理论16＋上机16，共16次课×2课时=32学时，不人为改写大纲学时。教学采用“短讲解—示例推导—分段编码—上机验证—结果复盘”，所有任务要求保留原始数据、代码、参数、版本、测试结果、图表和人工核验记录。','', '![课程知识脉络](images/course-knowledge-map.png)','', '### 课次总览','| 次数 | 课型 | 理论/上机分配 | 单元/项目 | 授课题目 | 主要评价 |','|---:|---|---|---|---|---|']
    for l in lessons: md.append(f'| {l["no"]} | {l["kind"]} | {l["split"]} | {l["section"]} | {l["topic"]} | {l["assessment"].split("；")[0]} |')
    md.append('')
    for l in lessons:
        md += [f'## 第{l["no"]}次课 {l["topic"]}','','### 课次信息',f'- **章节/单元**：{l["section"]}','- **周次**：按实际课表填写','- **课时安排**：2课时（100 min）',f'- **课型**：{l["kind"]}',f'- **理论/上机分配**：{l["split"]}','','![知识脉络图](images/lesson-%02d-knowledge-map.png)'%l['no'],'','### 教学目标',obj_md(l['goals']),'','### 课程思政',l['ideology'],'','### 专创融合',l['innovation'],'','### 教学重难点',f'**教学重点：**{l["focus"]}\n\n**教学难点：**{l["diff"]}','','### 教学方法与用具',f'**教学方法：**问题驱动法、建模推演法、代码阅读/演示法、预测—执行—验证法、任务驱动法、测试验收与错误复盘\n\n**教学用具：**{l["tools"]}','','### 教学设计','课前任务 → 互动导入（10 min）→ 传授新知与课堂训练（75 min）→ 过关检测（10 min）→ 课堂小结（3 min）→ 作业布置（1 min）→ 考勤（1 min）','','### 课前任务',preclass(l),'','### 互动导入',f'【教师】{l["intro"]}\n\n【学生】先独立预测/判断，再与同伴交换依据；教师收集典型分歧后进入本课。','','### 传授新知与课堂训练',process_md(l),'','### 过关检测','【教师】组织当堂检测：\n\n'+'\n'.join(f'{i+1}. {q}' for i,q in enumerate(l['checks']))+'\n\n【学生】独立作答/运行/推演验证；【教师】按错误类型讲解，要求说明模型、数值或图表依据。','','### 课堂小结',f'【教师】回到本课知识脉络图，用“问题—模型—计算/数据—验证—表达—边界”总结“{l["topic"]}”，再次强调教学重点：{l["focus"]}\n\n【学生】写下“本课一条最重要的模型/计算规则 + 一个最容易出错的边界”。','','### 作业布置',f'【教师】布置课后任务：{l["homework"]}\n\n【学生】按课程文件命名与证据要求整理提交；任何AI/开源工具建议均需保留人工核验记录。','','### 考勤','【教师】利用学校/课程实际使用的平台进行签到。\n\n【学生】按课程要求完成签到。','','### 课后教学反思',reflection(l),'','### 本章节参考文献',l['refs'],'']
    md_path=OUT_ROOT/MD_NAME; md_path.write_text('\n'.join(md),encoding='utf-8')

    manifest={'course':{'name':COURSE['name'],'code':COURSE['code'],'hours':32,'theory_hours':16,'lab_hours':16},'canonical_markdown':MD_NAME,'lesson_count':16,'syllabus':{'path':COURSE['syllabus_path'],'sha':COURSE['syllabus_sha'],'updated':COURSE['syllabus_date'],'version':COURSE['version']},'template_identifier':'程序设计基础（C&C++）Ⅰ最终校版学校教学设计布局基线（仅作Word格式模板）','course_map':'images/course-knowledge-map.png','lessons':[{'no':l['no'],'type':l['kind'],'hour_split':l['split'],'section':l['section'],'title':l['topic'],'image':f'images/lesson-{l["no"]:02d}-knowledge-map.png','assessment':l['assessment']} for l in lessons]}
    (OUT_ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT_ROOT/'README.md').write_text(f'''# 工程建模与科学计算可视化基础（Python）教案源码包\n\n本包采用 Markdown-first 工作流。\n\n- `{MD_NAME}`：唯一课程内容主源。\n- `images/`：课程总图与16次课知识脉络图。\n- `manifest.json`：课程、课次、图片、大纲版本及理论/上机分配映射。\n- `build_docx.py`：从内容主源编译学校格式Word的脚本。\n\n## 内容/格式边界\nMarkdown负责课程内容正确性；Word负责学校格式、表格、字体、分页和图片尺寸。Word视觉检查发现语义问题时，必须回到Markdown修改后重新编译。\n\n## 权威依据\n`{COURSE['syllabus_path']}`，SHA `{COURSE['syllabus_sha']}`，大纲更新时间 `{COURSE['syllabus_date']}`。\n\n## 学时说明\n现行大纲理论16学时、上机16学时；上机项目学时为3/4/5/4。为保持2课时一次课且不篡改大纲学时，本教案安排4次理实一体课次，实现总量精确闭合。\n''',encoding='utf-8')

    # CONTENT QA GATE (must pass before Word compilation)
    parsed,text=parse_md(md_path)
    refs=re.findall(r'!\[[^\]]*\]\((images/[^\)]+)\)',text); missing=[r for r in refs if not (OUT_ROOT/r).exists()]
    assert len(parsed)==16 and not missing
    assert THEORY_HOURS==16 and LAB_HOURS==16 and THEORY_HOURS+LAB_HOURS==32
    assert '平时表现20%＋课程作业20%＋期末综合项目60%' in text
    assert len(set(p['fields']['image'] for p in parsed))==16
    assert all('学习证据' in p['fields']['传授新知与课堂训练'] and '评价映射' in p['fields']['传授新知与课堂训练'] for p in parsed)
    assert all('【课后填写，不预填事实】' in p['fields']['课后教学反思'] for p in parsed)
    assert all(p['fields']['课次信息'].find('理论/上机分配')>=0 for p in parsed)
    # each lab project hours encoded exactly in schedule
    assert sum(1 for l in lessons if l['kind']=='理实一体')==4
    shutil.copy2(md_path,FINAL_MD)

    # WORD COMPILATION from content-approved Markdown
    doc=Document(str(TEMPLATE)); proto=copy.deepcopy(doc.tables[2]._element)
    for t in list(doc.tables[2:]): remove_table(t)
    for par in list(doc.paragraphs[24:]): par._element.getparent().remove(par._element)
    cover=doc.tables[0]
    for i,val in enumerate([COURSE['college'],COURSE['major'],COURSE['name'],'（按实际填写）','（按实际填写）',COURSE['semester'],COURSE['major']]): set_cell_text(cover.cell(i,1),val,13.5)
    info=doc.tables[1]
    set_cell_text(info.cell(0,1),COURSE['name'],10.4); set_cell_text(info.cell(1,1),COURSE['english'],9.3); set_cell_text(info.cell(2,1),COURSE['category']); set_cell_text(info.cell(2,3),COURSE['nature']); set_cell_text(info.cell(2,5),COURSE['language'])
    set_cell_text(info.cell(3,1),COURSE['semester']); set_cell_text(info.cell(3,5),COURSE['credits'])
    for c,v in zip(range(1,6),[32,16,0,16,0]): set_cell_text(info.cell(5,c),str(v))
    set_cell_text(info.cell(6,1),COURSE['major'],9.3); set_cell_text(info.cell(7,1),COURSE['textbook'],9.2); set_cell_text(info.cell(8,1),COURSE['college']); set_cell_text(info.cell(9,1),COURSE['prereq'],9.4); set_cell_text(info.cell(10,1),COURSE['followup'],9.2)
    set_cell_text(info.cell(11,1),'考试课（ ）；考查课（√）',10.5)
    set_cell_text(info.cell(12,1),'闭卷（ ）；开卷（ ）；笔试（ ）；机试（ ）；口试（ ）；其它（√）：期末综合项目',9.8)
    set_cell_text(info.cell(13,1),'平时表现（√）；课程作业（√）；期末综合项目（√）',10)
    set_cell_text(info.cell(14,1),COURSE['assessment'],10)
    intro='课程基本定位：本课程以Python为主要工具，面向工程问题抽象建模、科学计算、数据处理和结果可视化，建立“问题定义—假设与模型—程序求解—数据验证—图形表达”主线。\n核心学习结果：能够使用Python、NumPy、SciPy、Pandas和Matplotlib完成基础工程计算与可视化，完成两连杆机械臂和双轮小车典型建模/仿真任务，并说明假设、误差和工具边界。\n主要教学方法：短讲解—示例推导—分段编码—上机验证—结果复盘；要求保留原始数据、代码、参数、版本、测试结果、图表和人工核验记录。'
    set_cell_text(info.cell(15,1),intro,9.2,bold_prefixes=['课程基本定位','核心学习结果','主要教学方法'])

    for p in parsed:
        sep=doc.add_paragraph(); sep.paragraph_format.page_break_before=True; sep.paragraph_format.space_before=Pt(0); sep.paragraph_format.space_after=Pt(0); new_tbl=copy.deepcopy(proto); sep._p.addnext(new_tbl); t=doc.tables[-1]
        f=p['fields']; li=extract_info(f['课次信息'])
        set_cell_text(t.cell(0,1),li.get('章节/单元',''),9.5); set_cell_text(t.cell(0,3),p['topic'],9.7); set_cell_text(t.cell(1,1),li.get('周次','按实际课表填写')); set_cell_text(t.cell(1,3),li.get('课时安排','2课时（100 min）')+'；'+li.get('理论/上机分配',''))
        c=t.cell(2,1); c.text=''; pp=c.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.add_run().add_picture(str(OUT_ROOT/f['image']),width=Cm(12.4))
        set_cell_text(t.cell(3,1),clean_md(f['教学目标']),9.5,bold_prefixes=['知识目标','能力目标','价值目标']); set_cell_text(t.cell(4,1),clean_md(f['课程思政']),9.2); set_cell_text(t.cell(5,1),clean_md(f['专创融合']),9.2); set_cell_text(t.cell(6,1),clean_md(f['教学重难点']),9.2,bold_prefixes=['教学重点','教学难点'])
        meth=clean_md(f['教学方法与用具']); mm=re.search(r'教学方法：(.*?)(?:\n\n|\n)教学用具：(.*)',meth,re.S)
        if mm: set_cell_text(t.cell(7,1),mm.group(1).strip(),8.9); set_cell_text(t.cell(8,1),mm.group(2).strip(),8.9)
        else: set_cell_text(t.cell(7,1),meth,9.0); set_cell_text(t.cell(8,1),'电脑、Python课程环境。',9.0)
        set_cell_text(t.cell(9,1),clean_md(f['教学设计']),9.0); set_cell_text(t.cell(10,1),'教学步骤及主要教学内容',10.3); set_cell_text(t.cell(10,4),'设计意图',10.3)
        set_cell_text(t.cell(11,1),clean_md(f['课前任务']),8.8,bold_prefixes=['【教师】','【学生】']); set_cell_text(t.cell(11,4),'通过预读公式/代码/图表、环境检查和预期结果建立先备认知；不使用未经授权的敏感工程数据。',8.45)
        combo='【互动导入 10 min】\n'+clean_md(f['互动导入'])+'\n\n【传授新知与课堂训练 75 min】\n'+clean_md(f['传授新知与课堂训练'])
        set_cell_text(t.cell(12,0),'互动导入\n（10 min）\n\n传授新知与课堂训练\n（75 min）',9.0); set_cell_text(t.cell(12,1),combo,8.5,bold_prefixes=['【互动导入','【传授','一、','二、','三、','独立学习产物','学习证据','评价映射']); set_cell_text(t.cell(12,4),'采用“工程问题 → 模型/代码预期 → 实现/计算 → 边界/对照测试 → 数据/图表证据 → 结论边界”的闭环；理实一体课明确前后任务切换和证据收口。',8.4)
        set_cell_text(t.cell(13,1),clean_md(f['过关检测']),8.55); set_cell_text(t.cell(13,4),'检测模型定义、代码/数值语义、图表表达和验证依据，要求说明为什么而不只给结果。',8.4)
        set_cell_text(t.cell(14,1),clean_md(f['课堂小结']),8.55); set_cell_text(t.cell(14,4),'回到知识脉络图，建立“问题—模型—计算/数据—验证—表达—边界”的可迁移结构。',8.4)
        set_cell_text(t.cell(15,1),clean_md(f['作业布置']),8.55); set_cell_text(t.cell(15,4),'固化代码、数据、参数、测试和图表证据，为后续综合建模任务提供可复用输入。',8.4)
        set_cell_text(t.cell(16,1),clean_md(f['考勤']),8.55); set_cell_text(t.cell(16,4),'保留学校模板考勤环节，不预填实际出勤事实。',8.4)
        set_cell_text(t.cell(17,1),clean_md(f['课后教学反思']),8.25); set_cell_text(t.cell(17,4),'课后反思必须基于真实课堂、代码运行和图表/测试证据；课前只保留填写框架。',8.25)
        set_cell_text(t.cell(18,0),'本章节参考文献',8.2,align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(18,1),clean_md(f['本章节参考文献']),8.1)
        for row in t.rows:
            for cc in row.cells: cc.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cc)
        for cc in t.rows[17].cells: set_cell_margins(cc,top=35,start=70,bottom=35,end=70)
        for cc in t.rows[18].cells: set_cell_margins(cc,top=30,start=70,bottom=30,end=70)
        set_repeat_header(t.rows[0])

    cp=doc.core_properties
    cp.title='工程建模与科学计算可视化基础（Python） 教学设计（16次课）'
    cp.subject='32学时（理论16、上机16）；适用智能制造工程、机械工程、能源与动力工程、车辆工程'
    cp.keywords='工程建模,Python,NumPy,SciPy,Pandas,Matplotlib,教案,教学设计,科学计算,可视化'
    cp.comments='依据现行工程建模与科学计算可视化基础（Python）课程教学大纲编制；Markdown为内容主源，Word为学校格式编译产物。'
    doc.save(FINAL_DOCX)
    scrub=Path('/home/oai/skills/docx/scripts/privacy_scrub.py'); clean=BASE/'工程建模与科学计算可视化基础_Python_教案_16次课_最终版_clean.docx'
    subprocess.run(['python',str(scrub),str(FINAL_DOCX),'--out',str(clean)],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True); shutil.move(clean,FINAL_DOCX)
    shutil.copy2(Path(__file__),OUT_ROOT/'build_docx.py')
    manifest['files']={MD_NAME:sha256(md_path),'docx':sha256(FINAL_DOCX),'images_count':len(list(IMG_DIR.glob('*.png')))}; (OUT_ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    for z in [SOURCE_ZIP,COMPLETE_ZIP]:
        if z.exists():z.unlink()
    with zipfile.ZipFile(SOURCE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for fp in OUT_ROOT.rglob('*'):
            if fp.is_file(): z.write(fp,fp.relative_to(OUT_ROOT.parent))
    comp=BASE/'工程建模与科学计算可视化基础_Python_教案_完整交付包'; shutil.rmtree(comp,ignore_errors=True); shutil.copytree(OUT_ROOT,comp/'source'); shutil.copy2(FINAL_DOCX,comp/DOCX_NAME); shutil.copy2(FINAL_MD,comp/MD_NAME); (comp/'README.md').write_text('本包包含canonical Markdown源包与由其编译的学校格式Word。内容修改请先修改source中的Markdown和语义图片，再重新编译Word。\n',encoding='utf-8')
    with zipfile.ZipFile(COMPLETE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for fp in comp.rglob('*'):
            if fp.is_file(): z.write(fp,fp.relative_to(comp.parent))
    print(json.dumps({'md':str(FINAL_MD),'docx':str(FINAL_DOCX),'source_zip':str(SOURCE_ZIP),'complete_zip':str(COMPLETE_ZIP),'source_dir':str(OUT_ROOT),'theory_hours':THEORY_HOURS,'lab_hours':LAB_HOURS},ensure_ascii=False,indent=2))
if __name__=='__main__': main()

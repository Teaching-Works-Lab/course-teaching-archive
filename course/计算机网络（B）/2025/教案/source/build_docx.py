from __future__ import annotations
from pathlib import Path
from copy import deepcopy
import re, json, shutil, zipfile, hashlib, os, subprocess
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

BASE=Path('/mnt/data')
TEMPLATE=BASE/'机械测试技术教案.docx'
SYLLABUS=BASE/'25JX21801-计算机网络（B）-课程教学大纲.md'
ROOT=BASE/'25JX21801-计算机网络（B）-教案_源码包'
IMGDIR=ROOT/'images'
MD_NAME='25JX21801-计算机网络（B）-教案.md'
DOCX_NAME='25JX21801-计算机网络（B）-教案.docx'
FINAL_MD=BASE/MD_NAME
FINAL_DOCX=BASE/DOCX_NAME
SOURCE_ZIP=BASE/'25JX21801-计算机网络（B）-教案_源码包.zip'
COMPLETE_DIR=BASE/'25JX21801-计算机网络（B）-教案_完整交付包'
COMPLETE_ZIP=BASE/'25JX21801-计算机网络（B）-教案_完整交付包.zip'

COURSE={
 'name':'计算机网络（B）','english':'Computer Networks (B)','code':'25JX21801','category':'专业类','nature':'选修','language':'中文','semester':'第6学期','credits':'3',
 'hours_total':48,'hours_theory':48,'major':'智能制造工程','college':'计算机与信息技术学院',
 'prereq':'数据结构（C）、程序设计基础（C&C++）Ⅱ、工程建模与科学计算可视化基础（Python）',
 'followup':'先进装备系统设计与开发（ROS）、智能生产计划管理（MES/ERP）、毕业设计（论文）',
 'textbook':'胡亮, 徐高潮, 张宗升, 霍严梅, 车喜龙. 计算机网络[M]. 北京：高等教育出版社，2024. ISBN 9787040612875',
 'assessment':'总评成绩 = 平时学习与课堂分析×20% + 课程作业（协议与报文分析）×20% + 期末综合考查×60%',
 'assessment_detail':'期末综合考查采用网络设计与故障分析综合报告 + 个人核验答辩；不安排期末笔试或机考。',
}
REFS='''[1] 胡亮, 徐高潮, 张宗升, 霍严梅, 车喜龙. 计算机网络[M]. 北京：高等教育出版社，2024. ISBN 9787040612875\n[2] 袁华, 王昊翔, 黄敏. 深入理解计算机网络[M]. 北京：清华大学出版社，2024. ISBN 9787302662709\n[3] 冯博琴, 陈妍. 计算机网络（第4版）[M]. 北京：高等教育出版社，2023. ISBN 9787040599909'''

# no, unit, topic, branches, focus, diff, intro, teach bullets, student product, evidence, checks, homework, ideology, innovation, objectives, assessment mapping
lessons=[]
def add(no,unit,topic,branches,focus,diff,intro,teach,product,evidence,checks,homework,ideology,innovation,k,a,v,ass):
    lessons.append(dict(no=no,unit=unit,topic=topic,branches=branches,focus=focus,diff=diff,intro=intro,teach=teach,product=product,evidence=evidence,checks=checks,homework=homework,ideology=ideology,innovation=innovation,obj={'k':k,'a':a,'v':v},assessment=ass))

# Unit 1: 3 lessons / 6h
add(1,'第1单元 计算机网络概论、体系结构与性能','课程导入、网络组成与互联网结构',
 {'网络基本概念':['计算机网络','协议与资源共享'],'网络组成':['主机/终端','链路/交换节点'],'互联网结构':['网络边缘','接入网/核心网'],'分类与拓扑':['LAN/WAN','星型/树型/网状']},
 '计算机网络组成、互联网边缘/接入/核心结构。','从真实智能制造通信场景抽象网络组成，不把“互联网”简单等同于某一种物理拓扑。',
 '给出“机器人—边缘服务器—MES—云平台”通信图，让学生先圈出终端、链路、网络设备和服务，并判断哪些属于网络边缘。',
 ['从资源共享和通信角度定义计算机网络，区分互联网、局域网和通信系统。','解析网络边缘、接入网、核心网以及路由器/交换机的基本职责。','结合智能装备、ROS和MES场景，完成一张“端系统—接入—核心—服务”的通信结构图。'],
 '智能制造场景网络组成图及术语标注。','课堂分析图 + 关键术语解释 + 1条修正记录。',
 ['计算机网络最核心的功能是什么？','网络边缘与核心网分别包含哪些典型对象？','交换机与路由器为什么不能简单视为同一种设备？'],
 '选择一个熟悉的智能装备系统，画出至少6个节点的通信结构图，标明端系统、网络设备和服务。',
 '结合我国网络基础设施和工业互联网发展，强调网络是关键基础设施，工程人员应遵守标准、保障可靠连接。',
 '把网络组成图作为后续ROS/MES系统通信架构的基础设计文档，培养跨设备、软件与服务的系统视角。',
 ['理解计算机网络、互联网、网络边缘、接入网和核心网的基本概念。','认识常见网络分类和拓扑。'],
 ['能够从智能制造通信场景识别端系统、链路和网络节点。','能够用结构图说明基本通信关系。'],
 ['形成以标准术语描述网络系统、先画结构再分析问题的工程习惯。'],'平时学习与课堂分析；目标1.1、2.1、3.1。')

add(2,'第1单元 计算机网络概论、体系结构与性能','分组交换、时延、吞吐量与丢包',
 {'交换方式':['电路交换','分组交换'],'时延组成':['发送时延','传播/排队/处理时延'],'性能指标':['带宽','吞吐量/RTT'],'拥塞现象':['队列','丢包']},
 '分组交换；时延组成；吞吐量和丢包。','区分发送时延与传播时延，并用瓶颈链路解释端到端吞吐量。',
 '同样发送1 MB文件，分别改变链路带宽和链路长度，要求学生判断哪一种会改变发送时延、哪一种会改变传播时延。',
 ['比较电路交换和分组交换的资源使用方式与适用场景。','推导发送、传播、处理、排队时延的含义，结合RTT和丢包解释拥塞。','学生完成2道时延计算和1张“瓶颈链路—吞吐量”判断表，并说明单位。'],
 '时延/吞吐量计算单及瓶颈链路判断表。','计算过程 + 单位检查 + 错误修正说明。',
 ['发送时延由哪些量决定？','传播时延与链路带宽是否直接相关？','端到端多段链路的吞吐量通常由什么决定？'],
 '完成3道时延与吞吐量计算题，并解释一次“带宽高但访问仍慢”的可能原因。',
 '通过共享链路拥塞讨论资源公平使用、容量规划和对他人通信质量负责的工程意识。',
 '把时延、吞吐量、丢包转化为网络服务SLA的基本指标，为机器人控制、边缘计算和MES通信需求分析做准备。',
 ['掌握分组交换、电路交换和主要网络性能指标。','理解发送、传播、排队和处理时延。'],
 ['能够估算基本端到端时延和瓶颈吞吐量。','能够用可测指标解释网络性能现象。'],
 ['形成使用单位、公式和可测指标说明性能结论的严谨习惯。'],'平时学习与课堂分析；课程作业过程证据；目标1.1、2.1、3.1。')

add(3,'第1单元 计算机网络概论、体系结构与性能','协议、分层体系结构与封装/解封装',
 {'协议三要素':['语法','语义/同步'],'分层思想':['服务','接口'],'体系结构':['OSI','TCP/IP'],'封装过程':['应用数据','段/包/帧']},
 '协议、服务、接口；OSI/TCP-IP；封装与解封装。','沿端到端通信路径区分“同层协议关系”和“相邻层服务关系”。',
 '给出“浏览器访问网站”的数据流，让学生判断HTTP、TCP、IP、以太网分别在哪一层，并把“报文、段、包、帧”排序。',
 ['解释协议的语法、语义和同步，以及服务/接口与协议的区别。','比较OSI参考模型和TCP/IP体系，强调分层是降低复杂度的工程方法。','学生完成端到端封装图：应用数据→TCP段→IP包→以太网帧→接收端逆向解封装。'],
 '端到端分层通信与封装/解封装图。','分层图 + PDU名称 + 同层/相邻层关系说明。',
 ['协议和服务有什么区别？','TCP/IP体系中IP位于哪一层？','数据在发送端为什么要逐层增加首部？'],
 '画出“MES客户端访问服务器”的分层通信图，标明每层至少一个协议和对应PDU。',
 '通过开放标准和分层协作说明大型工程系统需要共同规则与接口契约。',
 '将分层思想迁移到软件架构和ROS接口设计，理解“稳定接口降低系统耦合”的价值。',
 ['掌握协议、服务、接口和分层体系结构。','理解封装与解封装过程。'],
 ['能够沿协议栈解释端到端数据流。','能够区分各层PDU和职责。'],
 ['形成接口契约、层次化分析和遵循标准的工程意识。'],'平时学习与课堂分析；课程作业过程证据；目标1.1、2.1、3.1。')

# Unit2 4 lessons
add(4,'第2单元 物理层、数据链路层、以太网与无线局域网','数据通信基础、信号、带宽与传输介质',
 {'信号基础':['模拟/数字','频率/周期'],'信道能力':['带宽','码元/比特率'],'传输介质':['双绞线/光纤','无线'],'物理层职责':['比特传输','接口/编码概念']},
 '信号、带宽、传输介质和物理层职责。','理解物理带宽、数据率和实际吞吐量不是同一个概念。',
 '对比“千兆以太网端口”和“1 GHz无线频段”中的“带宽”含义，要求学生判断两个术语是否完全相同。',
 ['回顾周期、频率、模拟/数字信号和数字通信基本概念。','比较双绞线、光纤和无线介质在距离、抗干扰、容量、成本方面的差异。','学生根据三种智能制造场景选择介质并写出至少3条工程约束。'],
 '传输介质选型表：场景、距离、干扰、带宽、安全、成本。','选型表 + 依据说明 + 一条反例修正。',
 ['物理层传输的基本单位是什么？','光纤相较双绞线有哪些典型优势？','为什么端口标称速率不等于应用吞吐量？'],
 '为“产线相机—边缘服务器”和“移动机器人—AP”分别选择介质并说明理由。',
 '通过电磁干扰、布线质量和通信失效案例强调基础设施质量和施工规范直接影响系统安全。',
 '从智能产线布线和无线覆盖需求出发，训练网络基础设施选型能力。',
 ['理解数据通信基础、信号和常见传输介质。'],['能够根据工程约束选择合适介质并解释理由。'],['形成面向环境、可靠性和维护成本综合选型的意识。'],'平时学习与课堂分析；目标1.2、2.2、3.2。')

add(5,'第2单元 物理层、数据链路层、以太网与无线局域网','数据链路层：成帧、差错检测、MAC与ARP',
 {'链路层职责':['成帧','透明传输'],'差错检测':['奇偶/校验概念','CRC思想'],'MAC寻址':['MAC地址','单播/广播'],'ARP':['IP→MAC','ARP缓存']},
 '成帧、差错检测、MAC地址与ARP。','区分IP地址与MAC地址的作用，并解释ARP仅用于本地链路邻居解析。',
 '给出“已知目标IP但不知道目标MAC”的局域网通信场景，学生先画主机需要广播什么、谁应该回答。',
 ['说明数据链路层如何把网络层数据组织为帧并进行差错检测。','解释MAC地址与IP地址分别解决“链路交付”和“网络寻址”的不同问题。','用教师提供ARP请求/响应字段或抓包截图，学生完成源/目的MAC、源/目的IP和缓存变化表。'],
 'ARP请求/响应分析表及帧字段标注图。','教师提供抓包截图/字段表 + 学生标注 + 缓存变化说明。',
 ['ARP请求为什么通常使用广播？','路由器外的远端主机MAC是否直接由本机ARP获得？','CRC检测到差错是否等于自动纠错？'],
 '根据给定ARP表和主机地址，推演同网段通信前的ARP过程。',
 '通过MAC欺骗和ARP异常案例强调授权环境、网络边界和规范诊断。',
 '将ARP/MAC分析作为设备接入和边缘节点故障定位的第一层工具。',
 ['掌握成帧、差错检测、MAC与ARP基本机制。'],['能够依据MAC/ARP信息分析局域网通信。'],['形成不越权抓包、用证据定位链路问题的习惯。'],'平时学习与课堂分析；课程作业过程证据；目标1.2、2.2、3.2。')

add(6,'第2单元 物理层、数据链路层、以太网与无线局域网','以太网帧、交换机学习转发、广播域与VLAN',
 {'以太网帧':['源/目的MAC','类型/FCS'],'交换表':['源地址学习','目的查表'],'转发行为':['单播转发','未知单播/广播泛洪'],'网络边界':['广播域','VLAN概念']},
 '以太网帧与交换机学习/转发；广播域与VLAN。','根据交换表状态动态推演帧转发，而不是死记“交换机转发端口”。',
 '给出一台4口交换机空MAC表和A→B、B→A、C→A三次通信，让学生逐步填写MAC表并预测每一步转发端口。',
 ['解析以太网帧主要字段及交换机源地址学习机制。','通过已知单播、未知单播和广播三类情况推演转发。','引入广播域和VLAN概念，学生完成“端口—VLAN—广播可达范围”示意表。'],
 '交换机MAC表演化与帧转发推演表。','三轮转发记录 + MAC表 + VLAN广播范围图。',
 ['交换机学习的是源MAC还是目的MAC？','未知单播帧如何处理？','VLAN为什么可以划分广播域？'],
 '完成一个双VLAN网络的广播可达性分析，不要求厂商命令。',
 '通过二层环路和广播风暴后果强化规范配置、边界意识和变更责任。',
 '把VLAN与设备类型/生产单元逻辑分区联系，训练中小型网络分区设计。',
 ['掌握以太网帧、交换机学习与转发机制。','理解广播域和VLAN概念。'],['能够根据MAC表推演帧传递并分析常见局域网问题。'],['形成规范配置和变更可追溯意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.2、2.2、3.2。')

add(7,'第2单元 物理层、数据链路层、以太网与无线局域网','802.11无线局域网与局域网综合故障分析',
 {'WLAN组成':['STA/AP','BSS/SSID'],'共享介质':['信道','竞争接入'],'无线问题':['干扰/覆盖','漫游/认证概念'],'综合诊断':['链路','MAC/ARP/VLAN']},
 '802.11基本组成、共享介质特性与局域网分层诊断。','将“信号问题、接入问题、地址问题、交换问题”分层，不用一个“网络不好”概括全部故障。',
 '给出“移动机器人可以连上Wi-Fi但访问不了边缘服务器”的故障描述，学生先列出至少四层不同原因。',
 ['介绍STA、AP、SSID/BSS、信道共享和无线覆盖/干扰的基本概念。','比较有线以太网与无线局域网在介质共享、稳定性和移动性上的差异。','用故障树整合本单元：物理信号→关联→VLAN→MAC/ARP→上层，学生完成分层排障顺序。'],
 'WLAN与以太网局域网综合故障树。','故障树 + 每层一个可观察证据 + 修复优先级。',
 ['AP与无线终端分别承担什么角色？','为什么信号强不等于应用一定可达？','局域网排障为什么应先确认链路和地址边界？'],
 '完善“机器人无法访问边缘服务器”故障树，至少包含8个检查点。',
 '强调无线接入安全、最小权限和禁止未经授权扫描/抓包。',
 '面向移动机器人和AGV无线接入，训练覆盖、漫游、VLAN和故障定位的系统思维。',
 ['理解802.11无线局域网基本组成和共享介质特性。'],['能够综合MAC、ARP、VLAN与无线信息分析局域网故障。'],['形成授权诊断和分层排错的工程规范。'],'平时学习与课堂分析；课程作业重点证据；目标1.2、2.2、3.2。')

# Unit3 5 lessons
add(8,'第3单元 网络层、IPv4/IPv6、路由与网络互联','IPv4数据报、地址结构、CIDR与前缀',
 {'IPv4数据报':['首部关键字段','TTL/协议号'],'IPv4地址':['网络/主机部分','点分十进制'],'CIDR':['前缀长度','无类地址'],'网络判断':['网络地址','广播地址']},
 'IPv4数据报、CIDR前缀和网络地址判断。','把“地址数、可用主机数、网络/广播地址”与CIDR前缀统一理解。',
 '给出192.168.10.130/26，要求学生不查表先判断网络地址、广播地址和可用主机范围。',
 ['解析IPv4首部中版本、总长度、TTL、协议、源/目的地址等字段。','从二进制位解释CIDR前缀、网络部分和主机部分。','学生完成4个地址的网络/广播/主机范围判断，并互相检查边界。'],
 'CIDR地址边界计算表。','二进制/十进制计算过程 + 边界检查。',
 ['/26对应多少个地址？','TTL的主要作用是什么？','同一IP地址在不同前缀长度下网络边界是否可能不同？'],
 '完成5个CIDR边界计算题，并写出1条“前缀写错导致网络故障”的说明。',
 '结合IPv4地址资源有限性，培养规范规划公共/私有地址资源的意识。',
 '将CIDR作为后续产线分区、设备子网和路由聚合的基础能力。',
 ['掌握IPv4数据报主要字段和CIDR地址表示。'],['能够判断网络地址、广播地址和主机范围。'],['形成地址规划精确、计算可复核的习惯。'],'平时学习与课堂分析；课程作业重点证据；目标1.3、2.3、3.3。')

add(9,'第3单元 网络层、IPv4/IPv6、路由与网络互联','IPv4子网划分与地址规划',
 {'需求分析':['主机数量','网络数量'],'子网掩码':['前缀选择','容量计算'],'地址分配':['网络/网关','主机区间'],'规划质量':['可扩展','不重叠/可追溯']},
 '子网划分、地址规划和容量校验。','从业务需求选择前缀并验证不重叠、容量和扩展余量。',
 '给出“3条产线分别需要50、20、10个地址”的192.168.100.0/24地址块，让学生比较等长子网和按需规划。',
 ['从主机规模和网络数量确定前缀，说明容量与浪费的权衡。','示范地址规划表：用途、网段、前缀、网关、主机范围、保留空间。','学生独立完成三产线地址规划，交换检查重叠和容量。'],
 '智能制造车间IPv4地址规划表。','地址表 + 容量计算 + 冲突检查记录。',
 ['一个需要50个主机地址的子网最小可选什么前缀？','两个子网为什么不能重叠？','地址规划为什么要保留用途和责任信息？'],
 '完成一个包含办公、机器人、相机、服务器四区的地址规划方案。',
 '通过地址资源规划培养长期演进、配置可追溯和公共资源规范使用意识。',
 '将地址规划表作为中小型智能制造网络方案的核心交付件。',
 ['掌握IPv4子网划分和地址规划方法。'],['能够根据规模需求完成子网方案并验证容量。'],['形成可扩展、可追溯的地址管理意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.3、2.3、3.3。')

add(10,'第3单元 网络层、IPv4/IPv6、路由与网络互联','ARP/ICMP、NAT、MTU与分片',
 {'邻接与诊断':['ARP','ICMP'],'地址转换':['NAT','私网/公网'],'报文大小':['MTU','分片概念'],'故障现象':['不可达','路径/大小问题']},
 'ICMP、NAT、MTU/分片及其故障现象。','理解ARP、ICMP和NAT处在不同问题层面，不能把“ping不通”简单归因于ICMP。',
 '给出“同网段能访问、跨网段不能访问”“能ping小包、传大数据异常”两类现象，要求学生判断检查方向。',
 ['梳理ARP与IP交付关系，介绍ICMP差错/诊断信息和Ping/Traceroute基本原理。','解释私网地址、NAT映射、端口复用概念。','介绍MTU与IPv4分片概念，学生完成一张“现象—可能层次—证据”故障表。'],
 'ARP/ICMP/NAT/MTU综合故障判断表。','故障现象分类 + 证据字段 + 分层判断理由。',
 ['Ping使用哪类协议支持？','NAT的主要目的是什么？','MTU过小可能引发什么网络现象？'],
 '分析3个网络层现象，分别给出第一检查证据和下一步。',
 '强调诊断工具必须在授权网络使用，不因“能运行工具”就越过网络边界。',
 '将ICMP/NAT/MTU故障分析迁移到边缘网关和跨网段设备通信场景。',
 ['理解ARP/ICMP关系、NAT、MTU和分片基本概念。'],['能够根据网络现象选择合适诊断证据。'],['形成授权诊断、逐层求证的习惯。'],'平时学习与课堂分析；课程作业重点证据；目标1.3、2.3、3.3。')

add(11,'第3单元 网络层、IPv4/IPv6、路由与网络互联','路由表、最长前缀匹配与路由选择',
 {'路由表':['目的前缀','下一跳/接口'],'转发原则':['最长前缀匹配','默认路由'],'路由来源':['直连/静态','动态路由概念'],'路径判断':['可达性','路由环路概念']},
 '路由表和最长前缀匹配。','同时存在多条匹配路由时，正确应用最长前缀匹配并推演下一跳。',
 '给出含/0、/16、/24、/25的路由表和4个目的地址，要求学生只按规则选路，不凭“看起来更近”。',
 ['解释路由表中目的前缀、下一跳、接口等基本字段。','用二进制前缀说明最长前缀匹配和默认路由。','学生完成4个目的地址的逐项匹配过程，并检查是否存在不可达/错误下一跳。'],
 '路由表查找与最长前缀匹配推演表。','逐条匹配过程 + 最终下一跳 + 错误路由修正。',
 ['什么是默认路由？','为什么选择最长前缀而不是最短前缀？','直连路由与静态路由有什么区别？'],
 '完成一张包含两个子网和一个默认出口的路由表推演题。',
 '通过路由错误导致大范围中断案例强调配置审核、变更记录和回滚责任。',
 '训练网络方案中路由可达性和故障影响分析能力。',
 ['掌握路由表和最长前缀匹配。'],['能够依据路由表分析IP转发与可达性。'],['形成配置变更可审查、错误可回滚的系统意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.3、2.3、3.3。')

add(12,'第3单元 网络层、IPv4/IPv6、路由与网络互联','RIP/OSPF/BGP概念、IPv6与网络层综合',
 {'动态路由':['RIP距离向量','OSPF链路状态'],'域间路由':['AS','BGP概念'],'IPv6':['地址表示','基本首部'],'综合互联':['IPv4/IPv6','路由/诊断']},
 '动态路由协议分工、IPv6地址/首部及网络层综合。','理解RIP/OSPF/BGP的适用层次而非记厂商配置；正确压缩IPv6地址。',
 '给出“园区内部选路”和“不同运营网络之间选路”两种问题，学生判断OSPF/BGP为何不是同一层次的协议。',
 ['概念性比较RIP、OSPF、BGP的作用范围和基本思路，不展开设备命令。','讲解IPv6地址表示、压缩规则、首部简化和IPv4向IPv6演进。','学生完成2个IPv6压缩/还原和一张网络层综合故障链：地址→网关→路由→ICMP/NAT→服务。'],
 '动态路由协议对比表 + IPv6地址练习 + 网络层综合故障链。','对比表、IPv6练习和分层故障路径。',
 ['OSPF与BGP的主要使用范围有什么不同？','IPv6地址中::最多出现几次？','网络层排障至少应检查哪三类信息？'],
 '整理网络层课程作业部分：地址规划、路由分析、ICMP/NAT或IPv6案例。',
 '结合IPv6演进说明工程技术需要兼顾存量兼容与长期发展。',
 '将IPv6和动态路由作为面向未来边缘/物联网网络方案的基础认知。',
 ['理解RIP/OSPF/BGP概念和IPv6基本机制。'],['能够识读IPv6并综合分析网络层故障。'],['形成持续学习、长期演进和系统兼容意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.3、2.3、3.3。')

# Unit4 4 lessons
add(13,'第4单元 运输层、UDP/TCP与端到端可靠传输','端口、复用/分用、UDP与Socket概念',
 {'进程通信':['端口','Socket端点'],'复用分用':['多进程→网络','网络→目标进程'],'UDP':['无连接','报文边界'],'适用场景':['低开销','实时/容错权衡']},
 '端口、UDP和进程间端到端通信。','区分IP地址定位主机与端口定位进程，并理解UDP“无连接”不等于“无用途/不可靠系统”。',
 '给出同一服务器同时运行Web、DNS、MQTT三项服务，提问“只有IP地址够不够找到正确程序？”',
 ['说明运输层复用/分用以及端口号和Socket端点概念。','解析UDP报文基本字段、无连接和保留报文边界的特性。','学生对DNS查询、视频流、设备遥测三个场景选择UDP/TCP并写出依据。'],
 '运输层端口映射图 + UDP/TCP初步选型表。','端口图、协议选型理由和反例修正。',
 ['端口号主要用于标识什么？','UDP为什么称为无连接？','使用UDP是否意味着应用一定不能实现可靠机制？'],
 '列出5个常见应用协议及其典型运输层协议/端口。',
 '通过共享网络资源和端口服务暴露讨论最小暴露、合法使用网络服务。',
 '为MES、机器人遥测和实时数据流选择运输层协议，训练需求驱动的协议选型。',
 ['掌握端口、复用/分用和UDP基本机制。'],['能够根据应用需求初步选择UDP/TCP。'],['形成端口服务最小暴露和需求驱动选型意识。'],'平时学习与课堂分析；课程作业过程证据；目标1.4、2.4、3.4。')

add(14,'第4单元 运输层、UDP/TCP与端到端可靠传输','TCP报文段、序号确认、三次握手与连接释放',
 {'TCP首部':['端口','Seq/Ack/标志位'],'建立连接':['SYN','三次握手'],'数据确认':['序号','累计确认'],'释放连接':['FIN/ACK','状态转换概念']},
 'TCP序号确认、三次握手与连接释放。','按字节序号理解Seq/Ack，并正确推演握手/释放而非只背报文个数。',
 '给出一组乱序的SYN、SYN-ACK、ACK报文，让学生按Seq/Ack关系排序并说明每一步确认了什么。',
 ['解析TCP首部中的序号、确认号和主要标志位。','用状态/报文序列解释三次握手为何建立双向通信状态。','学生完成教师提供握手与释放抓包/字段表，填Seq、Ack、Flags和方向。'],
 'TCP握手与释放报文序列表。','字段表 + 流程图 + 1条异常握手现象说明。',
 ['三次握手第二个报文通常包含哪些标志位？','ACK号表示确认到哪个位置？','TCP连接为什么要维护状态？'],
 '根据给定初始序号完成一次TCP握手与两段数据确认计算。',
 '通过连接状态和异常重试强调通信系统行为必须用报文证据解释。',
 '将TCP连接状态分析用于工业服务端口、API和边缘服务器故障定位。',
 ['掌握TCP首部、序号确认、连接建立与释放。'],['能够依据报文字段分析TCP连接状态。'],['形成基于抓包/日志证据解释系统状态的习惯。'],'平时学习与课堂分析；课程作业重点证据；目标1.4、2.4、3.4。')

add(15,'第4单元 运输层、UDP/TCP与端到端可靠传输','可靠传输、滑动窗口、超时重传与流量控制',
 {'可靠机制':['校验/序号','确认/重传'],'滑动窗口':['发送窗口','累计确认'],'超时机制':['RTT估计','超时重传'],'流量控制':['接收窗口','发送速率约束']},
 '可靠传输、滑动窗口、超时和流量控制。','区分流量控制与拥塞控制；理解窗口、RTT和重传共同影响性能。',
 '展示“网络不拥塞但接收端处理很慢”的场景，学生判断应由流量控制还是拥塞控制解决。',
 ['回顾可靠传输所需的序号、确认、超时和重传机制。','通过滑动窗口示意图解释连续发送与累计确认。','学生完成一组丢包/超时/重复ACK情景的发送窗口变化推演。'],
 'TCP可靠传输与窗口变化推演表。','报文序列 + 窗口变化 + 重传原因说明。',
 ['流量控制主要保护谁？','超时时间为什么不能固定得过小？','累计确认的含义是什么？'],
 '完成一个含丢包和接收窗口变化的可靠传输推演题。',
 '通过流量控制讨论系统设计中尊重接收能力、避免资源压垮的责任意识。',
 '把流控和超时参数理解为边缘设备、云服务之间性能调优的基础。',
 ['掌握TCP可靠传输、滑动窗口、超时重传和流量控制。'],['能够推演报文序号、确认与窗口变化。'],['形成从端到端可靠性和资源能力评价通信行为的意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.4、2.4、3.4。')

add(16,'第4单元 运输层、UDP/TCP与端到端可靠传输','拥塞控制、RTT与TCP/QUIC性能故障分析',
 {'拥塞现象':['队列/丢包','RTT升高'],'TCP拥塞控制':['慢启动概念','拥塞避免概念'],'性能证据':['RTT/重传','吞吐量'],'现代传输':['QUIC概览','TCP对比']},
 '拥塞控制基本思想；利用RTT/重传/吞吐量分析端到端性能。','区分接收端流控问题与网络路径拥塞，并避免把QUIC介绍为“永远更快”。',
 '给出两个抓包统计摘要：一个接收窗口持续变小，一个RTT和重传同时升高，学生判断哪一个更像接收端瓶颈、哪一个更像网络拥塞。',
 ['解释拥塞产生、慢启动/拥塞避免的基本思想，不展开具体实现版本细节。','用RTT、重传、窗口、吞吐量等证据构建性能诊断框架。','概览QUIC基于UDP实现现代传输能力；学生完成TCP/UDP/QUIC适用场景比较。'],
 '运输层性能诊断表 + TCP/UDP/QUIC对比表。','指标表、故障判断依据、协议选型说明。',
 ['拥塞控制主要保护谁？','RTT持续上升可能说明什么？','QUIC为什么不能简单等同于“UDP不可靠”？'],
 '完成课程作业中的TCP/UDP与可靠传输分析部分。',
 '通过拥塞控制讨论共享资源公平、网络稳定和性能优化中的公共责任。',
 '将现代传输协议视为智能制造云边协同和实时服务的技术选项，强调基于需求与证据选型。',
 ['理解TCP拥塞控制和QUIC基本概念。'],['能够依据RTT、窗口、重传等证据分析端到端性能问题。'],['形成兼顾性能、公平性与稳定性的工程意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.4、2.4、3.4。')

# Unit5 4 lessons
add(17,'第5单元 应用层协议与网络服务','客户端/服务器、P2P、DNS与DHCP',
 {'应用架构':['C/S','P2P'],'DNS':['层次命名','递归/迭代/缓存'],'DHCP':['地址租约','自动配置'],'服务依赖':['名称→地址','地址配置→可达']},
 'DNS查询链与DHCP自动配置；应用服务依赖。','区分“DNS解析失败”和“IP网络不可达”，并正确理解递归/迭代查询。',
 '给出“能ping服务器IP但访问域名失败”的现象，学生判断先查DNS还是TCP/HTTP。',
 ['介绍客户端/服务器与P2P基本架构。','讲解DNS层次、递归/迭代查询、缓存与常见记录概念。','讲解DHCP地址配置流程概念；学生完成“开机获得地址→解析域名→访问服务”的依赖链。'],
 'DNS/DHCP服务依赖与故障链图。','服务链图 + 2个故障现象的分层定位。',
 ['DNS解决什么问题？','递归查询与迭代查询有什么区别？','DHCP异常可能导致什么网络现象？'],
 '分析“域名失败、IP可达”和“未获得合法地址”两个案例。',
 '围绕域名、地址和网络服务强调基础服务的可用性与运维责任。',
 '把DNS/DHCP作为MES、边缘服务、机器人系统接入时的基础依赖服务。',
 ['掌握C/S、P2P、DNS和DHCP基本机制。'],['能够从域名和地址配置分析服务访问故障。'],['形成服务依赖和可用性意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.5、2.5、3.5。')

add(18,'第5单元 应用层协议与网络服务','HTTP/HTTPS、状态码、连接复用与Web缓存',
 {'HTTP模型':['请求/响应','方法/URL'],'状态信息':['状态码','首部'],'性能机制':['持久连接','缓存/CDN概念'],'HTTPS':['TLS承载','机密性/身份概念']},
 'HTTP请求响应、状态码、HTTPS与缓存。','区分HTTP应用语义与TLS安全机制，能从状态码/请求响应定位服务问题。',
 '给出200、301、404、500四个状态码和同一个“页面打不开”描述，学生判断这些现象分别指向客户端、资源还是服务端问题。',
 ['解析HTTP请求行、响应状态、常见首部和连接复用概念。','说明HTTPS是在HTTP与TLS等安全机制配合下实现安全通信，不把“有HTTPS”理解成应用绝对安全。','学生使用教师提供HTTP报文/抓包截图完成“请求—响应—状态码—缓存”分析表。'],
 'HTTP/HTTPS请求响应分析表。','教师提供报文 + 状态码解释 + 一条故障结论。',
 ['404与500通常分别表示什么？','HTTPS主要解决哪些通信安全问题？','Web缓存为什么能改善访问性能？'],
 '完成一组HTTP请求/响应分析，标出方法、主机、状态码和缓存相关字段。',
 '结合隐私泄露和证书错误案例强调个人信息保护、合法服务访问与安全验证。',
 '将HTTP/HTTPS和REST API作为MES/边缘平台服务接口的网络基础。',
 ['掌握HTTP/HTTPS请求响应与常见状态码。'],['能够依据报文和状态码分析Web服务访问问题。'],['形成隐私保护、接口规范和合法访问意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.5、2.5、3.5。')

add(19,'第5单元 应用层协议与网络服务','电子邮件、FTP/SFTP与Socket应用模型',
 {'邮件系统':['SMTP','POP3/IMAP'],'文件传输':['FTP概念','SFTP概念'],'应用端点':['Socket','IP+端口'],'服务安全':['明文风险','认证/权限']},
 '邮件协议分工、文件传输和Socket应用模型。','理解协议名称相似不代表安全机制相同；区分应用协议与底层Socket接口。',
 '给出“SMTP负责发信、IMAP负责收取/同步”的服务链，让学生解释为什么一个完整邮件系统需要多个协议。',
 ['说明SMTP、POP3、IMAP在邮件发送/接收中的分工。','比较FTP与SFTP的基本安全差异，不展开具体攻击操作。','回到Socket应用模型：应用协议最终通过运输层端点通信；学生画一张客户端/服务端端口关系图。'],
 '邮件/文件服务协议分工表 + Socket端点图。','协议对比表、端点图、安全边界说明。',
 ['SMTP主要负责发送还是接收？','SFTP与FTP最核心的安全差异是什么？','Socket端点通常由哪两个关键地址信息组成？'],
 '整理一个“协议—运输层—端口—安全注意事项”表，至少6项。',
 '通过明文传输与账号权限讨论敏感数据保护和最小权限。',
 '训练将协议、端口、认证和服务部署统一到应用系统接口清单中。',
 ['理解邮件、文件传输和Socket应用模型。'],['能够解释应用服务、运输层端点与安全机制的关系。'],['形成账号权限、敏感数据和协议安全意识。'],'平时学习与课堂分析；课程作业过程证据；目标1.5、2.5、3.5。')

add(20,'第5单元 应用层协议与网络服务','REST/API、WebSocket、MQTT、CDN与应用层综合故障链',
 {'现代接口':['REST/API','WebSocket'],'消息通信':['MQTT概念','发布订阅'],'内容分发':['缓存','CDN'],'综合诊断':['DNS','端口/HTTP/服务依赖']},
 '现代应用通信概念及从DNS到应用服务的故障链。','根据通信模式选择接口/协议，并从依赖链定位“服务不可用”。',
 '比较“周期性设备数据上报”“浏览器请求业务数据”“双向实时状态推送”三类需求，让学生选HTTP API、MQTT或WebSocket。',
 ['概览REST/API、WebSocket、MQTT的通信模式和典型适用场景。','解释CDN/内容分发和缓存的基本价值。','学生构建应用层故障链：DHCP/IP→DNS→TCP/UDP端口→TLS→HTTP/应用服务，并为每层给出一个证据。'],
 '应用协议选型表 + 服务访问故障链。','选型依据、分层证据和故障定位顺序。',
 ['MQTT通常采用什么通信模式？','WebSocket适合什么类型交互？','应用访问失败为什么要同时检查DNS、端口和应用状态？'],
 '完成课程作业中的DNS/HTTP/DHCP或现代服务链分析部分。',
 '强调现代网络服务使用中的隐私、接口权限和数据合规。',
 '把REST、WebSocket、MQTT联系到MES、机器人状态推送和IoT边缘应用。',
 ['理解REST/API、WebSocket、MQTT和CDN基本概念。'],['能够从完整服务依赖链分析应用访问故障。'],['形成接口规范、隐私保护和服务可用性意识。'],'平时学习与课堂分析；课程作业重点证据；目标1.5、2.5、3.5。')

# Unit6 4 lessons
add(21,'第6单元 网络管理、安全、新型网络与智能制造应用','网络管理与分层诊断：Ping、Traceroute、表项与抓包证据',
 {'管理对象':['接口/链路','地址/路由/服务'],'诊断工具':['Ping','Traceroute'],'状态表项':['ARP表','路由表/DNS查询'],'证据链':['日志','教师提供抓包']},
 '分层诊断工具及其证据边界。','理解工具输出只能说明特定层面的现象，不能用一次Ping替代完整故障结论。',
 '给出“Ping通但网页打不开”“同网段不通”“Traceroute中间节点不回应但终点可达”三个现象，要求学生判断能说明什么、不能说明什么。',
 ['建立分层故障诊断顺序：物理/链路→地址→路由→运输→域名/应用。','说明Ping、Traceroute、ARP表、路由表、DNS查询和教师提供抓包的观察对象。','学生完成三案例“工具输出—可证明—不可证明—下一步证据”表。'],
 '网络诊断证据矩阵与分层故障树。','证据矩阵 + 故障树 + 不确定项标注。',
 ['Ping通能否证明HTTP服务正常？','Traceroute主要观察什么？','为什么抓包必须限定在授权环境？'],
 '对一个“设备能上网但访问MES失败”案例给出5步分层诊断计划。',
 '明确所有网络探测、抓包和配置只能在授权环境、教师提供数据或隔离仿真中进行。',
 '将诊断证据矩阵转化为企业网络运维检查单。',
 ['理解常见网络管理/诊断工具的作用。'],['能够按分层顺序定位常见网络故障并说明证据边界。'],['形成合法授权、证据诊断和不确定项说明意识。'],'平时学习与课堂分析；期末综合考查过程准备；目标1.6、2.6、3.6。')

add(22,'第6单元 网络管理、安全、新型网络与智能制造应用','网络安全目标、TLS、防火墙、VPN与访问控制',
 {'安全目标':['机密性','完整性/可用性'],'身份与加密':['认证','TLS概念'],'边界控制':['防火墙','访问控制/最小权限'],'安全隧道':['VPN','远程访问边界']},
 '网络安全目标、认证/加密、TLS、防火墙、VPN和最小权限。','区分加密、认证、访问控制等不同机制，避免认为“上了防火墙就安全”。',
 '给出“通信被窃听、身份被冒充、服务被阻断、普通用户访问管理员接口”四类风险，让学生分别匹配安全目标/控制。',
 ['解释机密性、完整性、可用性、认证等基本安全目标。','概念性介绍TLS、防火墙、VPN和访问控制，强调多层防护。','学生完成“威胁—资产—安全目标—控制措施—残余风险”表，不涉及攻击实施。'],
 '智能制造网络威胁—控制映射表。','风险表 + 最小权限规则 + 残余风险说明。',
 ['TLS主要保护什么？','防火墙与身份认证解决的是同一问题吗？','最小权限原则为什么重要？'],
 '为一个远程维护场景设计最小权限、VPN/认证和日志要求。',
 '结合网络空间法治和关键系统安全强调未经授权访问不可接受，安全是工程交付的一部分。',
 '把安全控制纳入设备远程运维和边缘服务设计，而不是功能完成后再补。',
 ['理解网络安全目标和常见防护机制。'],['能够为简单网络场景设计访问控制和安全边界。'],['形成最小权限、合法授权和可恢复意识。'],'平时学习与课堂分析；期末综合考查准备；目标1.6、2.6、3.6。')

add(23,'第6单元 网络管理、安全、新型网络与智能制造应用','SDN、网络虚拟化、云/边缘、IoT与智能制造网络',
 {'新型架构':['SDN','网络虚拟化'],'云边协同':['云网络','边缘网络'],'物联网':['设备接入','消息/数据流'],'智能制造':['ROS/MES','工业互联网关系']},
 'SDN、云边缘、IoT概念及通用网络原理在智能制造中的迁移。','明确本课程讲通用网络原理，不把工业网络专用协议或厂商平台当作主线。',
 '给出ROS机器人、边缘服务器、MES、云分析四层架构，让学生标出哪些通信仍然需要IP、TCP/UDP、DNS/HTTP等通用网络基础。',
 ['概览SDN控制/数据平面分离思想和网络虚拟化。','说明云网络、边缘网络和IoT中的设备接入、服务发现、消息传递仍建立在通用网络机制上。','学生完成“ROS/MES/边缘/云—网络层次—关键协议—性能/安全需求”映射表。'],
 '智能制造网络技术映射表。','系统图 + 协议映射 + 两项性能/安全约束。',
 ['SDN的核心思想是什么？','边缘计算为什么仍然需要完整网络设计？','本课程与《工业网络技术及应用》的内容边界是什么？'],
 '选择ROS或MES场景，写出一页通用网络基础需求说明。',
 '通过自主创新、开源技术和关键基础设施案例培养持续学习与跨专业协作意识。',
 '把通用网络原理迁移到ROS、MES、边缘和IoT系统，形成后续装备系统开发的网络基础。',
 ['理解SDN、网络虚拟化、云/边缘和IoT基本概念。'],['能够把通用网络原理迁移到智能制造系统分析。'],['形成持续学习、跨专业沟通和边界意识。'],'平时学习与课堂分析；期末综合考查准备；目标1.6、2.6、3.6。')

add(24,'第6单元 网络管理、安全、新型网络与智能制造应用','中小型网络设计、分层故障分析与课程综合验收',
 {'需求与拓扑':['设备/用户','性能/可用性'],'网络设计':['VLAN/WLAN','IPv4/IPv6/路由'],'服务与安全':['TCP/UDP/DNS/HTTP','访问控制/恢复'],'验证与故障':['分层诊断','证据/答辩']},
 '拓扑、地址、路由、服务、安全与验证形成完整网络方案。','在有限信息下形成可实施方案，同时清楚标注假设、边界和验证方法。',
 '发布一个“智能制造车间网络”综合案例：办公终端、机器人、工业相机、边缘服务器、MES和远程维护，要求学生先列需求而不是直接画拓扑。',
 ['从设备/用户、通信关系、规模、带宽、时延、可用性和安全要求形成需求清单。','组织局域网/VLAN/WLAN、IP地址/网关/路由、TCP/UDP与DNS/HTTP服务、安全边界和监测恢复设计。','学生用给定故障（错误网关/路由、DNS异常、端口关闭等）修改方案并说明验证步骤，模拟期末个人核验答辩。'],
 '期末综合考查方案草图：需求、拓扑、地址/路由、服务、安全、验证和故障处理。','综合方案草图 + 地址/路由表 + 故障修改记录 + 口头解释提纲。',
 ['完整网络方案至少需要哪些类型的交付物？','为什么故障分析要从需求和分层证据出发？','期末综合考查为什么需要个人核验答辩？'],
 '完成期末综合考查正式任务准备：整理网络设计与故障分析综合报告框架及个人核验要点。',
 '以负责任网络工程交付为课程收束：方案必须可验证、风险可说明、工具使用合法、个人贡献可核验。',
 '将课程知识整合为可评审的智能制造中小型网络方案，训练从技术分析到工程交付。',
 ['综合理解网络体系结构、链路/IP/运输/应用、安全和新型网络。'],['能够形成中小型网络设计、验证和分层故障分析方案。'],['形成合法授权、最小权限、可恢复、持续学习和跨专业沟通意识。'],'期末综合考查；目标1.1—1.6、2.1—2.6、3.1—3.6。')

assert len(lessons)==24

# ---------- map drawing ----------
FONT_REG='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FONT_BOLD='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
if not Path(FONT_BOLD).exists(): FONT_BOLD=FONT_REG
W,H=1800,620

def font(path,size): return ImageFont.truetype(path,size)
def fit(draw,text,maxw,start=50,minsize=24,bold=False):
    fp=FONT_BOLD if bold else FONT_REG
    for s in range(start,minsize-1,-1):
        f=font(fp,s); b=draw.textbbox((0,0),text,font=f)
        if b[2]-b[0]<=maxw: return f
    return font(fp,minsize)
def box(draw,xy,fill,outline='#8AA1B5',radius=20,width=3): draw.rounded_rectangle(xy,radius=radius,fill=fill,outline=outline,width=width)
def ctext(draw,xy,text,f,fill='#26394C'):
    x0,y0,x1,y1=xy; b=draw.textbbox((0,0),text,font=f); tw,th=b[2]-b[0],b[3]-b[1]
    draw.text(((x0+x1-tw)/2,(y0+y1-th)/2-b[1]),text,font=f,fill=fill)
def draw_map(path,title,branches):
    im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im); line='#819AB5'
    rb=(300,18,1500,112); box(d,rb,'#DCECF7',outline='#4D7894',radius=28,width=4); ctext(d,rb,title,fit(d,title,1120,56,32,True),'#234C68')
    n=len(branches); margin=42; gap=28; cw=(W-2*margin-gap*(n-1))//n; bus=138
    centers=[]
    for i in range(n): centers.append(margin+i*(cw+gap)+cw//2)
    d.line((W//2,112,W//2,bus),fill=line,width=4); d.line((centers[0],bus,centers[-1],bus),fill=line,width=4)
    for i,(bl,leaves) in enumerate(branches.items()):
        left=margin+i*(cw+gap); cx=left+cw//2; d.line((cx,bus,cx,160),fill=line,width=4)
        bb=(left+8,160,left+cw-8,236); box(d,bb,'#EDF4F8'); ctext(d,bb,bl,fit(d,bl,cw-50,40,26,True),'#294D68')
        trunk=left+28; d.line((cx,236,cx,258),fill=line,width=3); d.line((cx,258,trunk,258),fill=line,width=3)
        ys=[300,404,508]; d.line((trunk,258,trunk,ys[min(len(leaves),3)-1]+35),fill=line,width=3)
        for j,leaf in enumerate(leaves[:3]):
            y=ys[j]; lx=left+60; rx=left+cw-5; d.line((trunk,y+35,lx,y+35),fill=line,width=3)
            lb=(lx,y,rx,y+70); box(d,lb,'#FBFCFD'); ctext(d,lb,leaf,fit(d,leaf,rx-lx-24,34,23,False),'#343B43')
    im.save(path,dpi=(300,300),optimize=True)

# ---------- markdown ----------
def obj_md(o):
    out=[]
    for label,key in [('知识目标','k'),('能力目标','a'),('价值目标','v')]:
        out.append(f'**{label}：**'); out += [f'- （{i+1}）{x}' for i,x in enumerate(o[key])]
    return '\n'.join(out)
def process_md(l):
    items='\n'.join(f'【教师】{x}\n【学生】完成对应分析/推演并记录判断依据。' for x in l['teach'])
    return f'''【互动导入 15 min】\n【教师】{l['intro']}\n【学生】先独立判断/计算，再交换依据。\n\n【传授新知与课堂训练 70 min】\n{items}\n\n【课堂产物】{l['product']}\n【学习证据】{l['evidence']}\n【评价关联】{l['assessment']}'''
def reflect(l): return f'''【课后填写，不预填事实】\n- 1. 教学流程：记录100 min各环节实际用时及需调整位置；\n- 2. 教学内容：重点记录“{l['diff']}”的真实掌握证据与常见错误；\n- 3. 学生参与：仅依据实际课堂记录填写独立分析、讨论、互审情况，不补写推测；\n- 4. 评价证据：检查本课分析表/流程图/计算单/方案证据是否完整，记录缺项；\n- 5. 后续改进：依据真实课堂证据确定下一轮调整。'''

# ---------- docx helpers ----------
def set_cell_text(cell,text,size=9,bold=False,align=None):
    cell.text=''; lines=str(text).split('\n')
    for i,line in enumerate(lines):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
        if align is not None: p.alignment=align
        r=p.add_run(line); r.bold=bold; r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(size)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def set_margins(cell,top=50,start=65,bottom=50,end=65):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=tcMar.find(qn('w:'+k))
        if el is None: el=OxmlElement('w:'+k); tcMar.append(el)
        el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa')
def clean_md(s):
    s=s.replace('**','').replace('`',''); s=re.sub(r'^- ','',s,flags=re.M); s=re.sub(r'!\[[^\]]*\]\([^\)]*\)','',s); return s.strip()

def parse_md(md_path):
    txt=md_path.read_text(encoding='utf-8'); parts=re.split(r'(?=^## 第\d{2}次课 )',txt,flags=re.M); parts=[p for p in parts if p.startswith('## 第')]
    out=[]
    for p in parts:
        h=re.match(r'^## 第(\d{2})次课 (.+)$',p,re.M); no=int(h.group(1)); topic=h.group(2).strip()
        def sect(name,next_names):
            m=re.search(rf'^### {re.escape(name)}\s*$',p,re.M)
            if not m:return ''
            start=m.end(); ends=[]
            for n in next_names:
                m2=re.search(rf'^### {re.escape(n)}\s*$',p[start:],re.M)
                if m2: ends.append(start+m2.start())
            return p[start:min(ends) if ends else len(p)].strip()
        names=['知识脉络图','教学目标','课程思政','专创融合 / 双创融合','教学重难点','教学方法','教学用具','教学设计','教学过程','过关检测','课堂小结','作业布置','考勤','课后教学反思','本章节参考文献']
        fields={n:sect(n,names[i+1:]) for i,n in enumerate(names)}
        info={}
        for k in ['章节/单元','周次','课时安排']:
            mm=re.search(rf'- \*\*{re.escape(k)}\*\*：(.+)',p)
            if mm: info[k]=mm.group(1).strip()
        img=re.search(r'!\[[^\]]*\]\((images/lesson-\d{2}-knowledge-map\.png)\)',fields['知识脉络图']); fields['image']=img.group(1) if img else ''
        out.append(dict(no=no,topic=topic,info=info,fields=fields))
    return txt,out

# ---------- build ----------
def main():
    for d in [ROOT,COMPLETE_DIR]:
        if d.exists(): shutil.rmtree(d)
    IMGDIR.mkdir(parents=True)
    for z in [SOURCE_ZIP,COMPLETE_ZIP]:
        if z.exists(): z.unlink()
    # images first as semantic assets
    for l in lessons: draw_map(IMGDIR/f'lesson-{l["no"]:02d}-knowledge-map.png',l['topic'],l['branches'])
    # canonical markdown
    md=['# 计算机网络（B）教案（24次课）','', '> **Canonical Markdown / 内容权威源**：本文件与 `images/`、`manifest.json` 共同构成课程教案内容主源；Word 仅由此编译。','', '## 教案封面信息',f'- **学院**：{COURSE["college"]}','- **专业（教研室）**：智能制造工程',f'- **课程名称**：{COURSE["name"]}','- **主讲教师**：（按实际填写）','- **职称**：（按实际填写）',f'- **授课学期**：{COURSE["semester"]}',f'- **授课专业**：{COURSE["major"]}','','## 课程基本信息']
    items=[('课程代码',COURSE['code']),('课程名称（中文）',COURSE['name']),('课程名称（英文）',COURSE['english']),('课程类别',f'{COURSE["category"]}；课程性质：{COURSE["nature"]}；授课语言：{COURSE["language"]}'),('授课学期',f'{COURSE["semester"]}；学分：{COURSE["credits"]}'),('课程学时及分配','总学时48；理论48；实验0；上机0；项目式0'),('适用专业',COURSE['major']),('教材',COURSE['textbook']),('授课学院',COURSE['college']),('先修课程',COURSE['prereq']),('后续课程',COURSE['followup']),('考核类型','考查课'),('考核形式','期末综合考查：网络设计与故障分析综合报告 + 个人核验答辩；不安排期末笔试或机考'),('考核方式','平时学习与课堂分析、课程作业（协议与报文分析）、期末综合考查'),('总评成绩比例',COURSE['assessment'])]
    md += [f'- **{k}**：{v}' for k,v in items]
    md += ['', '### 课程简介','', '课程基本定位：本课程是智能制造工程专业的专业选修课，建立“网络体系结构—局域网与数据链路—IP网络—运输层—应用层—网络管理与安全”的知识主线，侧重通用计算机网络原理与网络系统分析。工业网络专用协议不作为本课程主线。','', '核心学习结果：学生能够解释分层、封装、分组交换、时延、吞吐量和可靠传输等概念，识读典型协议数据流，完成IPv4子网划分、路由分析、TCP/UDP与应用服务故障分析，并形成含拓扑、地址、路由、服务、安全边界和验证方法的中小型网络设计方案。','', '学情分析：学生已完成程序设计、数据结构和Python基础课程，具备一定软件与算法基础。本课程48学时全部为理论教学，Wireshark、网络仿真、Socket和路由配置仅作为课堂演示、课内分析或课后材料，不另计实践学时。','', '主要教学方法：采用原理讲授、协议报文解析、分层通信过程推演、教师提供抓包数据演示、网络拓扑与地址规划练习、故障案例研讨和方案评审。坚持“先理解协议行为，再使用工具验证”，所有网络探测、抓包和配置限定在授权环境、教师提供数据或隔离仿真环境。','', '## 24次课教学设计','']
    for l in lessons:
        md += [f'## 第{l["no"]:02d}次课 {l["topic"]}','',f'- **章节/单元**：{l["unit"]}','- **周次**：按实际课表填写','- **课时安排**：2课时（100 min）','', '### 知识脉络图','',f'![第{l["no"]:02d}次课知识脉络图](images/lesson-{l["no"]:02d}-knowledge-map.png)','', '### 教学目标','',obj_md(l['obj']),'','### 课程思政','',l['ideology'],'','### 专创融合 / 双创融合','',l['innovation'],'','### 教学重难点','',f'教学重点：{l["focus"]}\n教学难点：{l["diff"]}','','### 教学方法','','讲授法、问题引导法、协议推演法、案例分析法、教师演示法、分组讨论法、方案评审与证据复盘','','### 教学用具','','电脑、投影仪、多媒体课件、教材、网络协议结构图、教师提供的抓包/报文字段资料、地址/路由练习表','','### 教学设计','','课前任务 → 互动导入（15 min）→ 传授新知与课堂训练（70 min）→ 过关检测（10 min）→ 课堂小结（3 min）→ 作业布置（1 min）→ 考勤（1 min）','','### 教学过程','', '#### 课前任务','',f'【教师】发布“{l["topic"]}”预习材料和一个最小判断问题；要求只使用教材、课程资料或教师提供数据。\n【学生】阅读对应内容，记录3个核心术语和至少1个疑问；准备课堂分析表。','', '**设计意图：** 通过阅读和最小问题建立先备认知，让学生带着可验证问题进入课堂。','', '#### 互动导入与传授新知 / 课堂训练','',process_md(l),'', '**设计意图：** 采用“先判断/计算 → 最小讲解 → 协议/数据流推演 → 独立分析产出 → 证据复核”的理论课闭环，把抽象网络原理落实到可检查的图表、计算和故障证据。','', '### 过关检测','', '【教师】组织当堂检测：\n'+'\n'.join(f'- {i+1}. {q}' for i,q in enumerate(l['checks']))+'\n【学生】独立作答；【教师】要求说明判断依据。','', '### 课堂小结','',f'【教师】回到知识脉络图，用“概念—关系—协议行为—工程边界”总结“{l["topic"]}”。\n【学生】写下本课最重要的一条规则和一个最容易误判的边界。','', '### 作业布置','',f'【教师】布置课后任务：{l["homework"]}\n【学生】按课程文件命名要求整理分析图、计算/协议表和必要说明。','', '### 考勤','', '【教师】利用学校/课程实际使用的平台进行签到。\n【学生】按课程要求完成签到。','', '### 课后教学反思','',reflect(l),'', '### 本章节参考文献','',REFS,'']
    md_path=ROOT/MD_NAME; md_path.write_text('\n'.join(md),encoding='utf-8'); shutil.copy2(md_path,FINAL_MD)
    # manifest/readme
    def sha(p):
        h=hashlib.sha256(); h.update(Path(p).read_bytes()); return h.hexdigest()
    manifest={'course':COURSE['name'],'course_code':COURSE['code'],'canonical_markdown':MD_NAME,'lesson_count':24,'hours':{'total':48,'theory':48,'experiment':0,'computer':0,'project':0},'assessment':{'平时学习与课堂分析':20,'课程作业':20,'期末综合考查':60},'lessons':[{'lesson':l['no'],'title':l['topic'],'unit':l['unit'],'image':f'images/lesson-{l["no"]:02d}-knowledge-map.png'} for l in lessons],'syllabus':'25JX21801-计算机网络（B）-课程教学大纲.md'}
    (ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'README.md').write_text(f'# {COURSE["name"]}教案源码包\n\n本包采用 Markdown-first 工作流。`{MD_NAME}` 是唯一内容主源，`images/` 含24张课次知识脉络图，Word由Markdown内容和学校教案模板编译。课程48学时全部为理论教学，不包含实验/上机/项目式课时。\n',encoding='utf-8')
    # content gate from markdown
    txt,parsed=parse_md(md_path)
    assert len(parsed)==24
    assert len(re.findall(r'images/lesson-\d{2}-knowledge-map\.png',txt))==24
    assert txt.count('【课后填写，不预填事实】')==24
    assert all((ROOT/p['fields']['image']).exists() for p in parsed)
    assert COURSE['hours_total']==24*2
    assert '总学时48；理论48；实验0；上机0；项目式0' in txt
    # compile Word from parsed markdown
    src=Document(str(TEMPLATE)); proto=deepcopy(src.tables[2]._tbl); body=src._element.body
    # remove old lesson tables and all following body content after note 5 (body idx 25)
    for ch in list(body)[26:]:
        if not ch.tag.endswith('}sectPr'): body.remove(ch)
    # cover
    cover=src.tables[0]
    vals=[COURSE['college'],'智能制造工程',COURSE['name'],'（按实际填写）','（按实际填写）',COURSE['semester'],COURSE['major']]
    for i,v in enumerate(vals): set_cell_text(cover.cell(i,1),v,12,align=WD_ALIGN_PARAGRAPH.CENTER)
    # basic info
    ci=src.tables[1]
    set_cell_text(ci.cell(0,1),COURSE['name'],9); set_cell_text(ci.cell(1,1),COURSE['english'],9)
    set_cell_text(ci.cell(2,1),COURSE['category'],9,align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(ci.cell(2,3),COURSE['nature'],9,align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(ci.cell(2,5),COURSE['language'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(ci.cell(3,1),COURSE['semester'],9,align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(ci.cell(3,5),COURSE['credits'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
    # Normalize hours header for this course; no practice/experiment/computer hours are invented.
    for col,val in enumerate(['课程学时及分配','总学时','理论','实验','上机','其他']): set_cell_text(ci.cell(4,col),val,8.5,bold=(col==0),align=WD_ALIGN_PARAGRAPH.CENTER)
    for col,val in enumerate(['课程学时及分配','48','48','0','0','0']): set_cell_text(ci.cell(5,col),val,9,bold=(col==0),align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(ci.cell(6,1),COURSE['major'],9); set_cell_text(ci.cell(7,1),COURSE['textbook'],8.5); set_cell_text(ci.cell(8,1),COURSE['college'],9); set_cell_text(ci.cell(9,1),COURSE['prereq'],8.5); set_cell_text(ci.cell(10,1),COURSE['followup'],8.5)
    set_cell_text(ci.cell(11,1),'考试课（ ）；考查课（√）',9); set_cell_text(ci.cell(12,1),'闭卷（ ）；开卷（ ）；笔试（ ）；机试（ ）；口试（√，个人核验答辩）；其它（√，综合报告）',8.2)
    set_cell_text(ci.cell(13,1),'作业（√）；报告（√）；随堂分析（√）；线上考试（ ）；线下考试（ ）；其它（个人核验答辩）',8.2); set_cell_text(ci.cell(14,1),COURSE['assessment'],8.6)
    intro='课程基本定位：本课程是智能制造工程专业的专业选修课，系统讲授通用计算机网络原理和网络系统分析，建立网络体系结构—局域网与数据链路—IP网络—运输层—应用层—网络管理与安全的知识主线。\n核心学习结果：学生能够识读典型协议数据流，完成IPv4子网与路由分析、TCP/UDP和应用服务故障分析，并形成中小型网络设计与验证方案。\n学情分析：学生已具备程序设计、数据结构和Python基础；48学时全部为理论教学，抓包、网络仿真、Socket和路由配置仅作课堂演示、课内分析或课后材料。\n主要教学方法：原理讲授、协议报文解析、分层通信推演、教师提供抓包演示、地址规划练习、故障研讨和方案评审。所有网络探测、抓包和配置限定在授权环境、教师提供数据或隔离仿真环境。'
    set_cell_text(ci.cell(15,1),intro,8.1)
    # append lesson tables, one page break before each lesson
    for p in parsed:
        br=src.add_paragraph(); br.add_run().add_break(WD_BREAK.PAGE); new=deepcopy(proto); br._p.addnext(new); t=Table(new,src)
        info=p['info']; f=p['fields']
        set_cell_text(t.cell(0,0),'章节',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(0,1),info['章节/单元'],8.6,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(0,2),'授课题目',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(0,3),p['topic'],8.6,True,WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(t.cell(1,0),'周次',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(1,1),info['周次'],9,False,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(1,2),'课时安排',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(1,3),info['课时安排'],9,False,WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(t.cell(2,0),'知识脉络图',9,True,WD_ALIGN_PARAGRAPH.CENTER); mc=t.cell(2,1); tc=mc._tc; tcPr=tc.tcPr
        for child in list(tc):
            if child is not tcPr: tc.remove(child)
        pe=OxmlElement('w:p'); tc.append(pe); pp=mc.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.add_run().add_picture(str(ROOT/f['image']),width=Inches(4.75)); t.rows[2].height=Inches(1.75); t.rows[2].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        set_cell_text(t.cell(3,0),'教学目标',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(3,1),clean_md(f['教学目标']),8.5)
        set_cell_text(t.cell(4,0),'课程思政',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(4,1),clean_md(f['课程思政']),8.3)
        set_cell_text(t.cell(5,0),'专创融合 /\n双创融合',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(5,1),clean_md(f['专创融合 / 双创融合']),8.3)
        set_cell_text(t.cell(6,0),'教学重难点',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(6,1),clean_md(f['教学重难点']),8.4)
        set_cell_text(t.cell(7,0),'教学方法',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(7,1),clean_md(f['教学方法']),8.3)
        set_cell_text(t.cell(8,0),'教学用具',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(8,1),clean_md(f['教学用具']),8.2)
        set_cell_text(t.cell(9,0),'教学设计',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(9,1),clean_md(f['教学设计']),8.2)
        set_cell_text(t.cell(10,0),'教学过程',11,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(10,1),'教学步骤及主要教学内容',10,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(10,4),'设计意图',9,True,WD_ALIGN_PARAGRAPH.CENTER)
        # parse process subsections embedded in f['教学过程']
        proc=f['教学过程']
        def sub(name,nextname=None):
            m=re.search(rf'^#### {re.escape(name)}\s*$',proc,re.M)
            if not m:return ''
            end=len(proc)
            if nextname:
                m2=re.search(rf'^#### {re.escape(nextname)}\s*$',proc[m.end():],re.M)
                if m2:end=m.end()+m2.start()
            return proc[m.end():end].strip()
        pre=sub('课前任务','互动导入与传授新知 / 课堂训练'); main=sub('互动导入与传授新知 / 课堂训练')
        # strip design intent lines from cells
        pre=re.sub(r'\*\*设计意图：\*\*.*','',pre,flags=re.S).strip(); main_clean=re.sub(r'\*\*设计意图：\*\*.*','',main,flags=re.S).strip()
        set_cell_text(t.cell(11,0),'课前任务',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(11,1),clean_md(pre),8.2); set_cell_text(t.cell(11,4),'通过课前阅读和最小判断问题建立先备认知，让学生带着可验证问题进入课堂。',8.0)
        set_cell_text(t.cell(12,0),'互动导入\n（15 min）\n\n传授新知与课堂训练\n（70 min）',8.5,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(12,1),clean_md(main_clean),7.9); set_cell_text(t.cell(12,4),'采用“先判断/计算→最小讲解→协议/数据流推演→独立分析产出→证据复核”的理论课闭环。',7.8)
        set_cell_text(t.cell(13,0),'过关检测\n（10 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(13,1),clean_md(f['过关检测']),8.0); set_cell_text(t.cell(13,4),'即时检验关键概念和边界，要求说明依据。',7.8)
        set_cell_text(t.cell(14,0),'课堂小结\n（3 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(14,1),clean_md(f['课堂小结']),8.0); set_cell_text(t.cell(14,4),'回到知识脉络图压缩信息，形成层级结构。',7.8)
        set_cell_text(t.cell(15,0),'作业布置\n（1 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(15,1),clean_md(f['作业布置']),8.0); set_cell_text(t.cell(15,4),'固化课堂分析成果，形成可复查的过程证据。',7.8)
        set_cell_text(t.cell(16,0),'考勤\n（1 min）',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(16,1),clean_md(f['考勤']),8.0); set_cell_text(t.cell(16,4),'保留模板考勤环节，不预填出勤事实。',7.8)
        set_cell_text(t.cell(17,0),'课后\n教学反思',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(17,1),clean_md(f['课后教学反思']),7.9); set_cell_text(t.cell(17,4),'反思必须依据真实课堂证据，课前只提供填写框架。',7.8)
        set_cell_text(t.cell(18,0),'本章节\n参考文献',9,True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.cell(18,1),clean_md(f['本章节参考文献']),7.7)
        for row in t.rows:
            for c in row.cells: set_margins(c)
    src.core_properties.title=f'{COURSE["name"]}教案'; src.core_properties.author=''; src.core_properties.last_modified_by=''; src.save(FINAL_DOCX)
    # privacy scrub
    clean=BASE/'_network_clean.docx'; subprocess.run(['python','/home/oai/skills/docx/scripts/privacy_scrub.py',str(FINAL_DOCX),'--out',str(clean)],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True); clean.replace(FINAL_DOCX)
    # build script copy + checksum manifest
    shutil.copy2(__file__,ROOT/'build_docx.py')
    manifest['sha256']={MD_NAME:sha(md_path)}
    for l in lessons: manifest['sha256'][f'images/lesson-{l["no"]:02d}-knowledge-map.png']=sha(IMGDIR/f'lesson-{l["no"]:02d}-knowledge-map.png')
    (ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    # source zip
    with zipfile.ZipFile(SOURCE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for p in ROOT.rglob('*'):
            if p.is_file(): z.write(p,p.relative_to(ROOT.parent))
    # complete package
    COMPLETE_DIR.mkdir(); shutil.copytree(ROOT,COMPLETE_DIR/'source'); shutil.copy2(FINAL_DOCX,COMPLETE_DIR/DOCX_NAME); shutil.copy2(FINAL_MD,COMPLETE_DIR/MD_NAME)
    (COMPLETE_DIR/'README.md').write_text('本包包含 canonical Markdown 源包和由其编译的学校格式 Word。内容修改请先改 source 中 Markdown，再重新编译 Word。\n',encoding='utf-8')
    with zipfile.ZipFile(COMPLETE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for p in COMPLETE_DIR.rglob('*'):
            if p.is_file(): z.write(p,p.relative_to(COMPLETE_DIR.parent))
    print(FINAL_MD); print(FINAL_DOCX); print(SOURCE_ZIP); print(COMPLETE_ZIP)

if __name__=='__main__': main()

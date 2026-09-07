from pathlib import Path
from copy import deepcopy
import re, shutil, json, hashlib, zipfile, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path('/mnt/data')
WORK_MD=BASE/'工业机器人技术_教案-working.md'
WORK_PKG=BASE/'工业机器人技术_教案_工作包'
TEMPLATE=BASE/'数据结构（C）_教案_20次课_最终版.docx'
FINAL_MD=BASE/'工业机器人技术_教案_16次课_最终版.md'
FINAL_DOCX=BASE/'工业机器人技术_教案_16次课_最终版.docx'
SRC_DIR=BASE/'工业机器人技术_教案_源码包'
FULL_DIR=BASE/'工业机器人技术_教案_完整交付包'
SRC_ZIP=BASE/'工业机器人技术_教案_16次课_源码包.zip'
FULL_ZIP=BASE/'工业机器人技术_教案_完整交付包.zip'

COURSE={
'name':'工业机器人技术','english':'Industrial Robot Technology','code':'25JD32409',
'category':'专业类','nature':'选修','language':'中文','semester':'第6学期','credits':'2',
'hours_total':'32','hours_theory':'16','hours_exp':'0','hours_lab':'16','hours_other':'0',
'major':'智能制造工程','college':'机械与动力工程学院',
'prereq':'机器人技术与应用（A）、智能测试技术、PLC技术基础或同等基础',
'followup':'先进装备系统设计与开发（ROS）、智能工厂集成技术、毕业设计（论文）',
'textbook':'张小刚、方遒.《工业机器人系统综合设计》[M]. 北京：机械工业出版社，2025. ISBN 9787111776192',
'assessment':'最终成绩100%=平时成绩20%＋课程作业20%＋期末考核60%。',
'exam':'工业机器人工作站综合项目成果＋个人答辩/核验',
'syllabus_sha':'afece6c8ea150effb4f1a6fc09c2646d557d02f6',
'skill_sha':'23bc5a10e7881e8863ad993c74dabdc9faf24a4d',
'date':'2026年9月'
}

def normalize_md(src:str)->str:
    s=src
    s=s.replace('document_state: drafting','document_state: final_compiled')
    s=s.replace('用户需要的成果：依据 GitHub 现行课程大纲与 lesson-plan-compiler Skill 编写《工业机器人技术》完整课程教案工作稿，并形成 Markdown + 知识脉络图工作包。',
                '用户需要的成果：依据 GitHub 现行课程大纲与 lesson-plan-compiler Skill 编写《工业机器人技术》完整课程教案，并形成 canonical Markdown、知识脉络图、学校格式 Word 与完整交付包。')
    s=s.replace('本次停止点：按当前 Skill 条件路由停在“工作 Markdown + 图片工作包”；主讲教师、职称、制定日期/学年等未确认，不自动编译正式 Word。',
                '本次停止点：已完成 canonical Markdown 内容校验并编译学校格式 Word；主讲教师、职称和具体校历周次未由用户提供，Word中不虚构，按学校惯例保留“按实际填写/按实际课表填写”。')
    # formal source table uses em dash for unknown; keep status in source tracking, but remove overt to_confirm wording from lesson week text.
    s=re.sub(r'第(\d+)次课（具体周次待校历确认）', r'第\1次课（Word中按实际课表填写）', s)
    s=s.replace('【教师】【教师】','【教师】').replace('【学生】【学生】','【学生】')
    # update final tracking
    s=s.replace('| P001 | 当前条件路由 | 用户确认主讲教师、职称、制定日期/学年与校历周次后，可进入 Word basic 门禁与正式编译。 | 封面、课程基本信息、各课次周次 | open |',
                '| P001 | 最终编译 | 用户本轮已要求生成Word；主讲教师、职称未提供，不虚构，Word封面保留“按实际填写”；具体周次保留“按实际课表填写”。 | 封面、各课次周次 | resolved_for_compilation |')
    return s

# create canonical final MD
raw=WORK_MD.read_text(encoding='utf-8')
final_md_text=normalize_md(raw)
FINAL_MD.write_text(final_md_text,encoding='utf-8')

# parse helpers

def md_clean(s:str)->str:
    s=s.replace('<br>','\n').replace('**','').replace('`','')
    s=s.replace('【教师】【教师】','【教师】').replace('【学生】【学生】','【学生】')
    return s.strip()

def parse_field_table(block:str):
    d={}
    for line in block.splitlines():
        if not line.startswith('|'): continue
        cells=[x.strip() for x in line.strip().strip('|').split('|')]
        if len(cells)>=4 and cells[0] not in ('字段','---') and not set(cells[0]) <= {'-'}:
            d[cells[0]]=cells[1]
    return d

def section_between(text, heading, next_heads):
    m=re.search(rf'^#### {re.escape(heading)}\s*$',text,re.M)
    if not m: return ''
    start=m.end(); ends=[]
    for h in next_heads:
        m2=re.search(rf'^#### {re.escape(h)}\s*$',text[start:],re.M)
        if m2: ends.append(start+m2.start())
    end=min(ends) if ends else len(text)
    return text[start:end].strip()

def parse_process(sec:str):
    rows=[]
    for line in sec.splitlines():
        if not line.startswith('|'): continue
        cells=[c.strip() for c in line.strip().strip('|').split('|')]
        if len(cells)!=3 or cells[0] in ('教学步骤','---') or set(cells[0]) <= {'-'}: continue
        rows.append((cells[0],md_clean(cells[1]),md_clean(cells[2])))
    return rows

# parse lessons
text=final_md_text
lesson_matches=list(re.finditer(r'^### lesson-(\d{2})\s*$',text,re.M))
lessons=[]
all_heads=['知识脉络图','教学目标','课程思政','专创融合','教学重难点','教学过程','教学反思']
for i,m in enumerate(lesson_matches):
    no=int(m.group(1)); end=lesson_matches[i+1].start() if i+1<len(lesson_matches) else text.find('# 内部追踪区',m.end())
    block=text[m.end():end]
    # metadata field table before first ####
    meta=parse_field_table(block[:block.find('####')])
    sections={}
    for hi,h in enumerate(all_heads):
        sections[h]=section_between(block,h,all_heads[hi+1:])
    proc=parse_process(sections['教学过程'])
    lessons.append({'no':no,'meta':meta,'sections':sections,'process':proc})
assert len(lessons)==16, len(lessons)

# Content QA
for L in lessons:
    for k in ['章节','授课题目','课时安排','教学方法','教学用具','教学设计','参考文献']:
        assert L['meta'].get(k), (L['no'],k)
    for h in ['教学目标','课程思政','教学重难点','教学过程','教学反思']:
        assert L['sections'].get(h), (L['no'],h)
    assert len(L['process'])==7,(L['no'],len(L['process']))
    img=WORK_PKG/L['meta']['知识脉络图']
    assert img.exists(),img
assert sum(2 for _ in lessons)==32
assert sum(1 for L in lessons if '类型：理论2' in L['meta']['课时安排'])*2==16
assert sum(1 for L in lessons if '类型：上机2' in L['meta']['课时安排'])*2==16

# DOCX functions

def set_run_font(r,size=9,bold=None,font='宋体'):
    r.font.name=font
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
    r.font.size=Pt(size)
    if bold is not None: r.bold=bold

def clear_cell(cell):
    tc=cell._tc; tcPr=tc.tcPr
    for child in list(tc):
        if child is not tcPr: tc.remove(child)
    p=OxmlElement('w:p'); tc.append(p)

def set_cell(cell,text,size=9,bold=False,align=None,bold_prefixes=()):
    clear_cell(cell)
    lines=str(text).split('\n') if text is not None else ['']
    for idx,line in enumerate(lines):
        p=cell.paragraphs[0] if idx==0 else cell.add_paragraph()
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
        if align is not None: p.alignment=align
        is_bold=bold or any(line.strip().startswith(x) for x in bold_prefixes)
        r=p.add_run(line); set_run_font(r,size,is_bold)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_cell_margins(cell,top=50,start=65,bottom=50,end=65):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+tag))
        if node is None: node=OxmlElement('w:'+tag); tcMar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')

def remove_table(table):
    el=table._element; el.getparent().remove(el)

def add_page_break_before(body, before_el):
    p=OxmlElement('w:p'); r=OxmlElement('w:r'); br=OxmlElement('w:br'); br.set(qn('w:type'),'page'); r.append(br); p.append(r); before_el.addprevious(p)

def set_row_cant_split(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)

# Build from school-format template
shutil.copy2(TEMPLATE,FINAL_DOCX)
doc=Document(FINAL_DOCX)
# preserve prototype first lesson table, expand to 20 rows by duplicating combined process row
proto=deepcopy(doc.tables[2]._element)
# remove all old lesson tables
for t in list(doc.tables[2:]): remove_table(t)
# remove old lesson separator paragraphs after notes
body=doc._element.body
for ch in list(body)[26:]:
    if ch.tag != qn('w:sectPr'): body.remove(ch)
# cover
cover=doc.tables[0]
cover_vals=[COURSE['college'],COURSE['major'],COURSE['name'],'（按实际填写）','（按实际填写）',COURSE['semester'],COURSE['major']]
for i,v in enumerate(cover_vals): set_cell(cover.cell(i,1),v,12,align=WD_ALIGN_PARAGRAPH.CENTER)
# date paragraph - locate paragraph containing 2024年9月
for p in doc.paragraphs:
    if '2024年9月' in p.text:
        p.text=COURSE['date']; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs: set_run_font(r,16,font='华文中宋')
# course info
info=doc.tables[1]
set_cell(info.cell(0,1),COURSE['name'],9)
set_cell(info.cell(1,1),COURSE['english'],9)
set_cell(info.cell(2,1),COURSE['category'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(info.cell(2,3),COURSE['nature'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(info.cell(2,5),COURSE['language'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(info.cell(3,1),COURSE['semester'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(info.cell(3,5),COURSE['credits'],9,align=WD_ALIGN_PARAGRAPH.CENTER)
for col,v in enumerate([COURSE['hours_total'],COURSE['hours_theory'],COURSE['hours_exp'],COURSE['hours_lab'],COURSE['hours_other']],1):
    set_cell(info.cell(5,col),v,9,align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(info.cell(6,1),COURSE['major'],9)
set_cell(info.cell(7,1),COURSE['textbook'],8.8)
set_cell(info.cell(8,1),COURSE['college'],9)
set_cell(info.cell(9,1),COURSE['prereq'],8.8)
set_cell(info.cell(10,1),COURSE['followup'],8.8)
set_cell(info.cell(11,1),'考试课（ ）；考查课（√）',9)
set_cell(info.cell(12,1),'闭卷（ ）；开卷（ ）；笔试（ ）；机试（ ）；口试（√，个人答辩/核验）；其它（√，工业机器人工作站综合项目成果）',8.5)
set_cell(info.cell(13,1),'平时成绩（√）；课程作业（√）；综合项目成果（√）；个人答辩/核验（√）',8.8)
set_cell(info.cell(14,1),COURSE['assessment'],9)
intro=('课程基本定位：本课程是智能制造工程专业的专业选修课。在《机器人技术与应用（A）》已经建立机器人运动学、动力学与轨迹基础的前提下，课程以工业机器人工作站为对象，围绕安全与规范、坐标与示教、程序与工艺、外围接口、工作站集成、离线仿真和调试维护组织教学。\n'
       '学时组织：总学时32学时，其中理论16学时、上机16学时；按16次课×2课时组织，每个大纲单元形成“理论讲解＋对应上机任务”的闭环。\n'
       '核心学习结果：学生能够完成工作站风险辨识、工具/工件坐标设置、基础示教程序、I/O握手与安全互锁、离线布局/路径/碰撞/节拍验证，以及报警、备份恢复和点检维护分析，并形成可复核工程文档。\n'
       '主要教学方法：短讲解—案例拆解—离线建站—程序与接口实现—故障注入—验收复盘；上机可使用RobotStudio、RoboDK、Visual Components、PQArt或等价平台，不绑定单一厂商。')
set_cell(info.cell(15,1),intro,8.5,bold_prefixes=('课程基本定位','学时组织','核心学习结果','主要教学方法'))

# rebuild lessons
for L in lessons:
    # page break paragraph then table
    p=OxmlElement('w:p'); r=OxmlElement('w:r'); br=OxmlElement('w:br'); br.set(qn('w:type'),'page'); r.append(br); p.append(r); body.append(p)
    tbl_el=deepcopy(proto); body.append(tbl_el)
    t=doc.tables[-1]
    # insert duplicate process row before old combined row to create 20 rows
    dup=deepcopy(t.rows[12]._tr); t.rows[12]._tr.addprevious(dup)
    assert len(t.rows)==20
    meta=L['meta']; sec=L['sections']; proc=L['process']
    set_cell(t.cell(0,1),meta['章节'],8.7,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(0,3),meta['授课题目'],8.7,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(1,1),'按实际课表填写',9,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(1,3),meta['课时安排'],8.8,align=WD_ALIGN_PARAGRAPH.CENTER)
    # map content cell: c1 merged through c4
    clear_cell(t.cell(2,1)); pmap=t.cell(2,1).paragraphs[0]; pmap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pmap.add_run().add_picture(str(WORK_PKG/meta['知识脉络图']),width=Cm(13.0))
    # objectives and text sections
    set_cell(t.cell(3,1),md_clean(sec['教学目标']),8.8,bold_prefixes=('知识目标','能力目标','价值目标'))
    set_cell(t.cell(4,1),md_clean(sec['课程思政']),8.6)
    set_cell(t.cell(5,1),md_clean(sec['专创融合']),8.6)
    set_cell(t.cell(6,1),md_clean(sec['教学重难点']),8.6,bold_prefixes=('教学重点','教学难点'))
    set_cell(t.cell(7,1),meta['教学方法'],8.5)
    set_cell(t.cell(8,1),meta['教学用具'],8.4)
    set_cell(t.cell(9,1),meta['教学设计'],8.4)
    # row 10 remains heading but ensure text
    set_cell(t.cell(10,1),'教学步骤及主要教学内容',9.2,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(10,4),'设计意图',9.2,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    # map process by order seven steps
    target_rows=[11,12,13,14,15,16,17]
    for row_idx,(step,content,intent) in zip(target_rows,proc):
        # normalize labels from markdown
        label=step.replace('（10 min）','\n（10 min）').replace('（共70 min）','\n（70 min）').replace('（5 min）','\n（5 min）').replace('（3 min）','\n（3 min）').replace('（2 min）','\n（2 min）')
        set_cell(t.cell(row_idx,0),label,8.5,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(t.cell(row_idx,1),content,8.15,bold_prefixes=('【教师】','【学生】','【要求】','【课堂强化】'))
        set_cell(t.cell(row_idx,4),intent,8.0)
    set_cell(t.cell(18,1),md_clean(sec['教学反思']),8.15,bold_prefixes=('【教师】',))
    set_cell(t.cell(19,1),meta['参考文献'],8.2)
    # labels of last two rows
    set_cell(t.cell(18,0),'课后\n教学反思',8.5,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(19,0),'本章节\n参考文献',8.5,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    # margins and don't split
    for row in t.rows:
        set_row_cant_split(row)
        seen=set()
        for c in row.cells:
            if id(c._tc) in seen: continue
            seen.add(id(c._tc)); set_cell_margins(c)

# Save and scrub privacy
doc.save(FINAL_DOCX)
# clean metadata
import subprocess
clean=BASE/'工业机器人技术_教案_16次课_最终版_clean.docx'
subprocess.run(['python','/home/oai/skills/docx/scripts/privacy_scrub.py',str(FINAL_DOCX),'--out',str(clean)],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True)
clean.replace(FINAL_DOCX)

# Package sources
for d in [SRC_DIR,FULL_DIR]:
    if d.exists(): shutil.rmtree(d)
    d.mkdir(parents=True)
(SRC_DIR/'images').mkdir(); (FULL_DIR/'images').mkdir()
shutil.copy2(FINAL_MD,SRC_DIR/FINAL_MD.name); shutil.copy2(FINAL_MD,FULL_DIR/FINAL_MD.name)
for p in (WORK_PKG/'images').glob('*.png'):
    shutil.copy2(p,SRC_DIR/'images'/p.name); shutil.copy2(p,FULL_DIR/'images'/p.name)
shutil.copy2(Path(__file__),SRC_DIR/'build_docx.py'); shutil.copy2(Path(__file__),FULL_DIR/'build_docx.py')
shutil.copy2(FINAL_DOCX,FULL_DIR/FINAL_DOCX.name)

def sha(p):
    h=hashlib.sha256(); h.update(p.read_bytes()); return h.hexdigest()
manifest={
 'course':COURSE,'document_state':'final_compiled','canonical_markdown':FINAL_MD.name,'compiled_docx':FINAL_DOCX.name,
 'lesson_count':16,'hours':{'total':32,'theory':16,'lab':16},'images':17,
 'source':{'syllabus_path':'course/工业机器人技术/2025/大纲/25JD32409-工业机器人技术-课程教学大纲.md','syllabus_sha':COURSE['syllabus_sha'],'skill_repo':'Teaching-Works-Lab/lesson-plan-compiler','skill_sha':COURSE['skill_sha']},
 'unresolved_handling':{'主讲教师':'Word封面保留“按实际填写”，未虚构','职称':'Word封面保留“按实际填写”，未虚构','具体周次':'Word各课次保留“按实际课表填写”，未虚构'},
 'checksums':{FINAL_MD.name:sha(FINAL_MD),FINAL_DOCX.name:sha(FINAL_DOCX)}
}
for d in [SRC_DIR,FULL_DIR]:
    (d/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
readme=('''# 《工业机器人技术》教案源码包\n\n本包按 Markdown-first / lesson-plan-compiler 工作流生成。\n\n- canonical Markdown：`工业机器人技术_教案_16次课_最终版.md`\n- `images/`：课程总图 + 16次课知识脉络图\n- `manifest.json`：大纲/Skill版本、学时、文件与未提供字段处理\n- `build_docx.py`：从 canonical Markdown 内容和学校格式模板编译 Word 的脚本\n\n正式内容修改应先回到 Markdown，再重新编译 Word。主讲教师、职称和校历周次未由用户提供，未虚构。\n''')
(SRC_DIR/'README.md').write_text(readme,encoding='utf-8')
(FULL_DIR/'README.md').write_text(readme+'\n本完整包另含学校格式 Word 成品。\n',encoding='utf-8')
# zips
for zpath,root in [(SRC_ZIP,SRC_DIR),(FULL_ZIP,FULL_DIR)]:
    if zpath.exists(): zpath.unlink()
    with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(root.rglob('*')):
            if p.is_file(): z.write(p,arcname=f'{root.name}/{p.relative_to(root).as_posix()}')
print('built',FINAL_MD,FINAL_DOCX,SRC_ZIP,FULL_ZIP,sep='\n')

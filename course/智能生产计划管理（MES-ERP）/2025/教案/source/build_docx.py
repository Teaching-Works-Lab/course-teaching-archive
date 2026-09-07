from __future__ import annotations
import re, json, shutil, zipfile, hashlib, copy, subprocess
from pathlib import Path
from typing import Dict, List
from graphviz import Digraph
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path('/mnt/data')
TEMPLATE=BASE/'程序设计基础_C_C++_I_教案_24次课_最终版.docx'
OUT_ROOT=BASE/'智能生产计划管理_MES-ERP_教案_源码包'
IMG_DIR=OUT_ROOT/'images'
MD_NAME='智能生产计划管理_MES-ERP_教案_16次课_最终版.md'
DOCX_NAME='智能生产计划管理_MES-ERP_教案_16次课_最终版.docx'
SOURCE_ZIP=BASE/'智能生产计划管理_MES-ERP_教案_16次课_源码包.zip'
COMPLETE_ZIP=BASE/'智能生产计划管理_MES-ERP_教案_完整交付包.zip'
FINAL_MD=BASE/MD_NAME
FINAL_DOCX=BASE/DOCX_NAME

COURSE={
 'name':'智能生产计划管理（MES/ERP）','english':'Intelligent Production Planning and Management (MES/ERP)','code':'25JD32416',
 'category':'专业类','nature':'选修课','language':'中文','semester':'第7学期','credits':'2','major':'智能制造工程','college':'机械与动力工程学院',
 'hours_total':32,'hours_theory':32,'hours_lab':0,
 'prereq':'智能制造导论、机械制造技术基础、数据库基础或同等基础',
 'followup':'智能制造综合实践、毕业实习、毕业设计（论文）',
 'textbook':'刘晓冰, 赵希男.《生产计划与控制（第2版）》[M]. 北京：清华大学出版社，2022',
 'assessment':'平时成绩20%＋课程作业20%＋期末考核60%',
 'syllabus_path':'course/智能生产计划管理（MES-ERP）/2025/大纲/25JD32416-智能生产计划管理（MES-ERP）-课程教学大纲.md',
 'syllabus_sha':'56c08d91171bf77f08ddbe8fcb5718eefc89fcb5','syllabus_date':'2026年9月','version':'2025版目录现行大纲'
}
REFS={
 'main':'[1] 刘晓冰, 赵希男.《生产计划与控制（第2版）》[M]. 清华大学出版社，2022。',
 'mes':'[2] 王爱民.《制造执行系统（MES）原理与应用》[M]. 北京航空航天大学出版社，2021。',
 'erp':'[3] 陈启申.《ERP——从内部集成起步（第3版）》[M]. 电子工业出版社，2022。'
}

COURSE_MAP={
 '制造运营体系':['订单到交付','计划层级','ERP / MES边界'],
 '计划与资源':['MPS / MRP','BOM / 库存 / 提前期','能力 / 排程'],
 '车间执行与集成':['工单 / WIP','质量 / 追溯','ERP-MES数据流'],
 '智能决策':['APS','KPI / 成本','持续改进']
}

UNIT_NAMES=[
 '制造运营管理与ERP/MES体系','需求管理、MPS与MRP','能力计划与生产排程基础',
 '车间作业控制与执行闭环','MES功能、数据与ERP-MES集成','智能计划、绩效与管理决策'
]


def L(no, section, topic, split, unit_hours, goals, map_, focus, diff, intro, blocks, product, evidence, assessment, checks, homework, ideology, innovation, tools, refs):
    return dict(no=no, section=section, topic=topic, kind='理论', split=split, unit_hours=unit_hours, goals=goals, map=map_, focus=focus, diff=diff,
                intro=intro, blocks=blocks, product=product, evidence=evidence, assessment=assessment, checks=checks, homework=homework,
                ideology=ideology, innovation=innovation, tools=tools, refs=refs)

lessons:List[Dict]=[
L(1,'第一单元 制造运营管理与ERP/MES体系','制造企业订单到交付业务链与生产计划层级','理论2课时',
 {'制造运营管理与ERP/MES体系':2},
 {'k':['理解制造企业从客户订单到生产交付的基本业务链，以及战略、主计划、物料/能力计划和车间执行的层级关系。','理解订单、物料、工艺路线、工作中心、工单、库存和质量记录等核心业务对象。'],
  'a':['能够从一个简化制造订单案例识别业务对象、主要流程、输入输出数据和管理目标。','能够绘制订单—计划—执行—反馈的基本闭环，并指出交期、库存、WIP和产能之间的约束关系。'],
  'v':['形成以客户需求、交付承诺和企业资源约束为基础的系统管理意识，避免局部最优替代整体协同。']},
 {'业务链':['客户订单','计划 / 采购 / 生产'],'计划层级':['长期 / 中期','MPS / MRP / 车间'],'核心对象':['物料 / BOM','工艺 / 资源 / 工单'],'管理目标':['交期','库存 / WIP / 成本']},
 '订单到交付业务链、生产计划层级与核心数据对象。','把业务事件、计划层级、数据对象和管理目标对应起来，避免把MES/ERP理解为孤立的软件菜单。',
 '给出一个“客户要求10天交付100件产品，但关键工序能力有限、部分物料不足”的场景，让学生先判断哪些部门和数据必须参与计划。',
 [('业务链拆解','从客户订单、销售承诺、物料准备、能力配置、生产执行、质量检验到完工入库逐步拆解制造闭环，明确每一步输入输出。'),
  ('计划层级与对象','用时间尺度和决策粒度解释战略/销售运营、MPS/MRP、能力计划和车间作业控制，建立订单、物料、工艺、资源、工单等对象表。'),
  ('案例复盘','学生根据简化订单案例绘制业务链，并识别若“只追求设备利用率”可能造成的WIP、交期或库存副作用。')],
 '一张订单到交付业务流程图＋核心业务对象/数据表。','流程节点、责任对象、输入输出数据、至少2个管理目标冲突及其解释。','平时成绩＋期末综合报告基础；目标1.1、2.1、3.1。',
 ['ERP/MES课程为什么要从订单到交付业务链开始？','MPS与车间派工处于同一决策层级吗？为什么？','设备利用率越高是否一定意味着企业绩效越好？'],
 '选择一个离散制造产品，列出从订单到完工至少8个业务对象及其关键字段，为后续系统边界分析准备。',
 '以交付承诺、资源约束和跨部门协同讨论企业责任，强调管理结论必须建立在真实业务数据与可执行资源条件上。',
 '将制造业务链抽象为可配置的流程与数据对象，为后续MES/ERP业务建模和数字化产品设计建立基础。',
 '投影仪、订单/BOM/库存/工艺路线简化案例表、流程图模板、白板。',REFS['main']+'\n'+REFS['erp']),

L(2,'第一单元 制造运营管理与ERP/MES体系','ERP、MRP/MRPⅡ、MES、APS体系与系统边界','理论2课时',
 {'制造运营管理与ERP/MES体系':2},
 {'k':['理解ERP、MRP、MRPⅡ、MES、APS的核心概念、主要职责和相互关系，了解ISA-95层级思想。','理解主数据、业务数据、状态数据和经营数据的基本区别，以及系统边界与数据责任。'],
  'a':['能够针对订单、BOM、库存、工艺、工单、设备状态、报工和质量数据判断其主要产生/维护/消费位置。','能够绘制ERP—MES—设备/人员的概念边界图，并识别重复维护和职责模糊风险。'],
  'v':['形成权限控制、数据主责、业务契约和跨部门沟通意识，认识系统集成首先是业务与数据边界问题。']},
 {'系统体系':['ERP / MRPⅡ','MES / APS'],'层级关系':['计划层','执行层 / 现场层'],'数据类别':['主数据','业务 / 状态数据'],'边界原则':['单一主责','接口契约 / 权限']},
 'ERP/MES/APS职责边界、计划与执行层级、数据主责。','避免“ERP什么都做”或“MES等于看板”的泛化理解，能用业务对象和数据流解释系统边界。',
 '展示同一“工单状态”在ERP和MES里都可修改的冲突案例，让学生判断谁应为主责、另一侧如何同步。',
 [('体系演进','从MRP的物料计算到MRPⅡ资源闭环、ERP企业资源集成，再到MES车间执行和APS高级计划，解释各系统的决策范围。'),
  ('数据与边界','围绕物料、BOM、工艺、计划订单、生产工单、报工、设备状态、质量记录建立“主责系统—消费系统—同步方向”表。'),
  ('冲突分析','学生分析重复主数据、双向任意修改、状态不同步等典型问题，给出数据主责和接口边界原则。')],
 'ERP/MES/APS职责矩阵＋核心数据主责表。','系统职责、对象主责、同步方向、至少2个边界冲突案例及改进规则。','平时成绩＋课程作业ERP-MES数据流前置＋期末综合报告；目标1.1、2.1、3.1。',
 ['MRP、MRPⅡ和ERP的核心扩展关系是什么？','为什么MES不能简单替代ERP？','同一主数据由多个系统随意维护会造成什么问题？'],
 '完善“对象—主责系统—消费系统—更新频率—异常处理”表，预习独立需求、相关需求与MPS。',
 '通过数据主权、权限和责任边界案例强调数字化系统中的数据诚信与岗位责任。',
 '把系统边界和接口对象视为数字化制造解决方案的架构资产，训练平台化和接口化思维。',
 '投影仪、ERP/MES概念架构图、ISA-95层级示意、数据对象卡片。',REFS['erp']+'\n'+REFS['mes']),

L(3,'第二单元 需求管理、MPS与MRP','需求管理、预测/订单与主生产计划MPS','理论2课时',
 {'需求管理、MPS与MRP':2},
 {'k':['理解独立需求与相关需求、预测与客户订单的关系，掌握MPS的对象、时段和基本编制逻辑。','理解可承诺量、计划稳定性和需求变动对生产计划的影响。'],
  'a':['能够依据简化预测、订单、期初库存和批量条件编制分时段MPS表。','能够检查MPS是否与交付承诺、库存和资源条件一致，并解释计划变更影响。'],
  'v':['形成数据准确、版本一致和计划可执行意识，认识频繁无依据变更会向采购和车间传递波动。']},
 {'需求来源':['预测','客户订单'],'需求类型':['独立需求','相关需求'],'MPS结构':['计划时段','预计库存 / 计划产量'],'计划控制':['稳定区','变更 / 承诺']},
 '需求管理与MPS时间分段计算。','区分预测、订单和实际需求口径，并在库存、交期与计划稳定性之间保持一致。',
 '给出预测与实际订单在第3周发生偏差的案例，让学生判断能否直接把预测和订单相加，以及MPS应如何响应。',
 [('需求口径','解释独立/相关需求、预测、客户订单和订单承诺，讨论不同阶段需求数据的来源与可信度。'),
  ('MPS推演','用4—6个计划时段的简化数据计算预计可用库存、计划产量和可能的可承诺量，学生分组核对。'),
  ('变更分析','改变一个时段订单量或库存值，比较MPS前后变化，讨论计划冻结、沟通与版本控制。')],
 '一份6时段MPS计算表＋需求版本说明。','需求来源、计算公式/口径、时段结果、变更前后对比和一条可执行性说明。','课程作业MPS/MRP计算（30%内部权重）阶段1＋期末考核；目标1.2、2.2、3.2。',
 ['独立需求与相关需求的差别是什么？','MPS为什么不是简单的销售订单列表？','计划版本变化后哪些下游数据可能受影响？'],
 '完成给定案例MPS表，并写出一个“订单增加20%”时的影响清单。',
 '通过计划承诺与版本变更强调守信、协同和数据责任，避免用随意改数“满足”交期。',
 '把MPS表设计成可复用的计划计算模板，训练将业务规则转化为可验证的数据产品。',
 '投影仪、MPS时段案例表、计算器或电子表格（仅辅助计算）、白板。',REFS['main']),

L(4,'第二单元 需求管理、MPS与MRP','BOM、库存状态、提前期与批量规则','理论2课时',
 {'需求管理、MPS与MRP':2},
 {'k':['掌握BOM层级、父子件用量、库存现有量/已分配量/预计到货、提前期和批量规则等MRP基础数据。','理解低层码、提前期偏置和基础数据准确性对MRP结果的影响。'],
  'a':['能够从产品结构图建立简化BOM并按层级展开需求。','能够构造库存状态记录并判断错误BOM、库存或提前期如何导致计划偏差。'],
  'v':['形成基础数据“一处错误、全链扩散”的质量意识，重视BOM、库存和提前期版本一致性。']},
 {'产品结构':['父件 / 子件','BOM层级 / 用量'],'库存状态':['现有量','预计到货 / 已分配'],'时间参数':['采购提前期','制造提前期'],'批量规则':['逐批','固定批量 / 安全库存']},
 'BOM展开、库存状态与提前期偏置。','把数量逻辑与时间逻辑同时放入MRP基础数据，理解错误主数据的连锁效应。',
 '给出一个BOM中紧固件用量误写为1而实际需要4的案例，让学生追踪该错误在多批计划中的放大效应。',
 [('BOM结构','从产品—组件—零件建立多层BOM，解释父子关系、用量、层级和版本。'),
  ('库存与时间','建立现有量、预计到货、已分配和提前期数据表，演示需求数量与计划时间的两条计算线。'),
  ('数据质量','注入一个BOM用量、库存或提前期错误，学生比较MRP前置数据和可能的生产后果。')],
 '多层BOM图＋库存/提前期基础数据表。','BOM层级、单位用量、库存状态、提前期、批量规则和一条数据错误影响链。','课程作业MPS/MRP计算（30%内部权重）阶段2＋期末考核；目标1.2、2.2、3.2。',
 ['BOM中的“单位用量”为什么会沿层级放大？','库存现有量和预计到货量能否直接等同？','提前期错误主要影响MRP的数量还是时间？'],
 '完善案例BOM和库存状态记录，为下一课完整MRP净需求展开准备。',
 '以BOM和库存数据错误造成缺料/积压为例强调基础数据治理、版本责任和工程诚信。',
 '将BOM与库存状态视为计划算法的核心数据模型，训练数据结构化与规则可配置思维。',
 '投影仪、产品结构卡片、BOM/库存状态模板、电子表格（可选）。',REFS['main']+'\n'+REFS['erp']),

L(5,'第二单元 需求管理、MPS与MRP','MRP净需求、计划订单与安全库存校核','理论2课时',
 {'需求管理、MPS与MRP':2},
 {'k':['掌握总需求、预计可用量、净需求、计划订单接收与计划订单下达的基本计算逻辑。','理解安全库存、批量规则、提前期和例外信息对MRP结果的影响。'],
  'a':['能够依据MPS、BOM、库存和提前期完成一个两层产品的MRP展开。','能够通过数量守恒、时段偏置和边界条件检查MRP结果，并解释计划订单为何变化。'],
  'v':['形成计划结果必须可追溯到数据和规则的证据意识，不把软件输出直接等同于正确结论。']},
 {'需求展开':['总需求','低层传递'],'净需求':['预计可用','净需求'],'计划订单':['接收','下达 / 提前期偏置'],'参数影响':['安全库存','批量 / 例外信息']},
 'MRP净需求计算、计划订单接收/下达及校核。','同时处理数量、时段、批量和安全库存，能从结果反查数据与规则。',
 '给出“系统显示第4周缺料，但采购提前期为2周”的MRP表，让学生判断计划订单应在哪一周下达。',
 [('MRP计算链','按“总需求→预计可用→净需求→计划接收→计划下达”逐行推演，关联上层计划与BOM用量。'),
  ('参数实验','改变安全库存、固定批量或提前期，比较净需求和计划下达的变化，讨论库存与交付权衡。'),
  ('结果校核','学生使用数量守恒、时段偏置和零/边界库存检查MRP表，指出一个潜在例外并提出处理建议。')],
 '完整MPS/MRP计算表＋结果解释短报告。','计算过程、参数、净需求、计划订单、至少2个参数对照和校核结论。','课程作业MPS/MRP计算与结果解释（30%内部权重）完成＋期末考核；目标1.2、2.2、3.2。',
 ['计划订单接收与计划订单下达为什么不一定在同一时段？','安全库存增加通常会怎样影响计划订单？','如何验证MRP结果而不是只相信表格输出？'],
 '整理MPS/MRP阶段作业；预习工艺路线、工作中心、标准工时和能力负荷。',
 '强调计算结果必须可复核、参数变更必须有依据，避免通过随意调整库存/提前期“美化”计划。',
 '把MRP推演表视作可审核的计划算法原型，为后续ERP计划模块与APS优化理解打基础。',
 '投影仪、MRP案例数据、计算器/电子表格（仅辅助）、校核清单。',REFS['main']+'\n'+REFS['erp']),

L(6,'第三单元 能力计划与生产排程基础','工艺路线、工作中心、标准工时与能力数据','理论2课时',
 {'能力计划与生产排程基础':2},
 {'k':['理解工艺路线、工序顺序、工作中心、设备/人员能力、准备时间、加工时间和标准工时等能力计划基础数据。','理解产能日历、可用时间、效率/利用率等参数的基本作用。'],
  'a':['能够把一个产品工艺路线转换为工作中心负荷需求，并建立简化能力数据表。','能够识别工艺路线或标准工时错误对能力计划和交期判断的影响。'],
  'v':['形成资源节约、标准数据有据和工艺—计划协同意识，认识能力数据既影响效率也影响员工负荷与安全。']},
 {'工艺路线':['工序顺序','工作中心'],'时间数据':['准备时间','加工时间 / 标准工时'],'能力资源':['设备 / 人员','日历 / 可用时间'],'负荷基础':['订单数量','工时×数量']},
 '工艺路线、工作中心和标准工时到负荷需求的转换。','区分工艺时间、日历能力和有效能力，并避免把理论能力直接当可用能力。',
 '给出两条工艺路线，其中关键工序都经过同一工作中心，让学生判断为什么物料齐套仍可能无法按期交付。',
 [('能力数据模型','解释工艺路线、工序、工作中心、准备/加工时间和资源日历，建立“产品—工序—中心—工时”关系。'),
  ('负荷计算','根据订单数量和标准工时计算各工作中心工时负荷，区分理论可用时间与考虑效率后的有效能力。'),
  ('数据审查','比较标准工时低估、工艺路线漏工序和日历未更新三类错误，学生说明对交期和资源配置的影响。')],
 '工艺路线表＋工作中心能力表＋初步负荷计算表。','工序顺序、标准工时、工作中心、日历能力、负荷公式和数据风险说明。','课程作业能力负荷与排程分析（25%内部权重）阶段1＋期末考核；目标1.3、2.3、3.3。',
 ['工艺路线与BOM分别回答什么问题？','理论能力与有效能力为什么不同？','标准工时偏小会怎样影响交期承诺？'],
 '完成给定订单在3个工作中心上的工时负荷表，标注能力数据来源/假设。',
 '从工时标准、人员负荷和设备能力讨论效率目标与安全/合理劳动负荷之间的责任边界。',
 '把工艺路线和能力表结构化，为数字化产能分析、排程和瓶颈可视化准备数据模型。',
 '投影仪、工艺路线/工作中心案例表、能力日历模板、计算器。',REFS['main']),

L(7,'第三单元 能力计划与生产排程基础','能力需求、负荷平衡与瓶颈调整','理论2课时',
 {'能力计划与生产排程基础':2},
 {'k':['理解粗能力计划/能力需求计划、负荷、产能、瓶颈和能力平衡的基本关系。','理解加班、外协、转移工序、调整批量/交期等常见能力调整手段及其约束。'],
  'a':['能够依据工作中心负荷与能力数据识别超负荷时段和瓶颈。','能够提出至少两种能力调整方案，并比较交期、成本、质量和人员影响。'],
  'v':['形成瓶颈管理、资源节约和协同决策意识，避免只通过无限加班解决计划问题。']},
 {'负荷需求':['工时负荷','时段分布'],'能力供给':['可用工时','效率 / 日历'],'瓶颈判断':['负荷率','超负荷时段'],'调整手段':['交期 / 批量','加班 / 外协 / 转移']},
 '负荷—能力对比、瓶颈识别和调整方案评价。','在交期、成本、质量和人员影响之间做有约束的方案权衡。',
 '给出关键工作中心负荷率130%的周计划，让学生先回答“加班30%”是否是唯一方案，并列出潜在副作用。',
 [('负荷平衡','将上一课负荷表按时段汇总，与能力日历比较，计算负荷率并识别瓶颈。'),
  ('方案设计','讨论加班、外协、转移、批量调整、交期协商等方案，要求说明前提与代价。'),
  ('方案评审','学生用交期、成本、质量、员工负荷、风险五个维度比较两种方案，形成简短管理建议。')],
 '工作中心负荷—能力图＋两方案比较表。','负荷率、瓶颈时段、调整量、约束条件、方案比较和推荐理由。','课程作业能力负荷与排程分析（25%内部权重）阶段2＋期末考核；目标1.3、2.3、3.3。',
 ['负荷率超过100%意味着什么？','为什么“无限能力计划”仍有分析价值？','加班方案评价至少还应考虑哪些非交期因素？'],
 '完成瓶颈能力调整方案，预习正排、倒排、优先规则和生产工单下达。',
 '强调产能决策对员工、安全、质量和成本的综合影响，避免把效率指标作为唯一价值尺度。',
 '将负荷—能力比较转化为排程决策看板，为APS与智能排程的约束建模建立直觉。',
 '投影仪、能力负荷案例、方案评价矩阵、白板。',REFS['main']),

L(8,'第三→第四单元过渡','正排/倒排、优先规则与生产工单下达衔接','第三单元1课时＋第四单元1课时',
 {'能力计划与生产排程基础':1,'车间作业控制与执行闭环':1},
 {'k':['理解有限/无限能力、正排/倒排及常见优先规则的基本思想。','理解计划订单转生产工单、下达条件和计划层向执行层移交的业务含义。'],
  'a':['能够对同一组工单按交期、最短加工时间等规则形成简化排序并比较结果。','能够制定工单下达前检查清单，连接物料、能力、工艺和版本状态。'],
  'v':['形成计划可执行、下达有条件和跨层级责任清晰的意识，避免把“排出来”误当“能执行”。']},
 {'排程方向':['正排','倒排'],'能力约束':['有限能力','无限能力'],'优先规则':['交期优先','加工时间 / 紧迫度'],'工单下达':['齐套 / 能力','版本 / 状态']},
 '排程规则与计划到车间工单的移交条件。','理解优先规则的局限，并把排程结果与工单下达的物料/能力/工艺约束连接起来。',
 '给出3张工单：一张交期最紧、一张加工时间最短、一张客户等级最高，让学生说明不同优先规则为何给出不同顺序。',
 [('排程基础','用简化甘特表解释正排、倒排、有限/无限能力和常见优先规则，比较完工时间、延期和在制品倾向。'),
  ('工单下达','从计划订单到生产工单说明齐套、工艺版本、工作中心能力、质量/图纸版本和授权状态等下达条件。'),
  ('衔接评审','学生把排序结果放入工单下达检查表，判断哪些工单“顺序靠前但暂不能下达”，解释原因。')],
 '简化工单排序表＋生产工单下达检查清单。','规则、排程顺序、延期/等待结果、物料/能力/版本下达条件及判断依据。','课程作业能力负荷与排程分析（25%内部权重）完成，并衔接工单执行作业；目标1.3/2.3/3.3与1.4/2.4/3.4。',
 ['正排和倒排各适合回答什么问题？','不同优先规则为什么可能相互冲突？','排程顺序第一的工单为什么仍可能不能下达？'],
 '整理能力与排程阶段作业；选一张工单列出从“计划下达”到“完工报工”的预计状态。',
 '通过工单下达条件强调计划人员对可执行性、版本和资源事实负责，避免把压力简单转嫁给车间。',
 '把排程规则与下达检查清单转化为可配置决策规则，为APS/工作流系统设计提供原型。',
 '投影仪、工单/交期/工时案例、甘特图模板、下达检查表。',REFS['main']),

L(9,'第四单元 车间作业控制与执行闭环','工单状态、派工、报工与WIP控制','理论2课时',
 {'车间作业控制与执行闭环':2},
 {'k':['掌握生产工单从创建、下达、派工、开工、报工到完工/关闭的基本状态流。','理解WIP、工序进度、投入/产出记录和车间反馈对计划调整的作用。'],
  'a':['能够绘制工单状态机并判断不允许的状态跳转。','能够根据简化报工数据计算在制数量、完工数量和偏差，并提出计划反馈信息。'],
  'v':['形成生产记录真实、状态可追溯和跨岗位协同意识，不用事后补录掩盖真实执行偏差。']},
 {'工单生命周期':['创建 / 下达','派工 / 开工'],'执行记录':['报工','投入 / 产出 / 工时'],'WIP':['工序在制','排队 / 加工'],'反馈闭环':['进度偏差','计划调整']},
 '工单状态流、报工数据和WIP反馈。','把业务状态与实际生产事实对应，防止非法跳转、漏报/补报造成执行数据失真。',
 '给出“工单尚未下达却出现完工报工”的记录，要求学生判断这是界面问题、流程问题还是数据治理问题。',
 [('状态流','建立工单创建—下达—派工—开工—暂停/异常—报工—完工—关闭的状态机，明确每次转换的业务条件。'),
  ('报工与WIP','用投入、良品、不良品、在制、工时数据分析工序进度，解释WIP过高/过低的管理含义。'),
  ('反馈分析','学生根据一组实际进度低于计划的数据，形成“事实—偏差—影响—建议”反馈卡，区分记录事实与管理判断。')],
 '工单状态图＋WIP/报工分析表。','状态转换条件、投入产出数据、WIP计算、计划偏差及一条调整建议。','课程作业工单执行与异常闭环（20%内部权重）阶段1＋期末考核；目标1.4、2.4、3.4。',
 ['工单“下达”和“开工”有什么区别？','WIP为什么不是越低越好或越高越好？','为什么报工数据必须区分事实记录和计划判断？'],
 '完善工单状态机，为下一课质量、不合格、返工和异常流程增加分支。',
 '通过真实报工、不可篡改事实与责任追溯强调生产数据诚信和质量责任。',
 '把工单状态机和WIP数据设计成MES执行工作流核心模型，训练事件驱动和状态化产品思维。',
 '投影仪、工单状态卡、报工/WIP案例数据、状态图模板。',REFS['main']+'\n'+REFS['mes']),

L(10,'第四单元 车间作业控制与执行闭环','物料配送、质量异常、返工与计划调整闭环','理论2课时',
 {'车间作业控制与执行闭环':2},
 {'k':['理解车间物料配送、领退补料、质量检验、不合格、返工/报废和异常停机的基本业务关系。','理解异常事实如何触发工单、库存、质量和计划层的协同调整。'],
  'a':['能够为“不合格返工”或“关键物料短缺”绘制跨部门异常闭环。','能够判断异常发生后哪些状态/数量必须同步更新，哪些决策需人工授权。'],
  'v':['形成质量追溯、异常闭环和不得掩盖不合格/停机事实的职业责任意识。']},
 {'物料执行':['领料 / 配送','退料 / 补料'],'质量流程':['检验','不合格 / 返工 / 报废'],'异常管理':['缺料','停机 / 质量异常'],'计划反馈':['数量变化','交期 / 能力再平衡']},
 '物料—质量—异常—计划调整的闭环。','跨系统/部门保持数量、状态和追溯一致，同时保留人工审批与责任边界。',
 '给出“工序完成100件，其中8件不合格、5件返工、3件报废”的数据，让学生判断完工、WIP、库存和订单交付量应如何更新。',
 [('物料与质量','串联领料、退料、补料、检验、合格、不合格、返工和报废，说明批次/序列追溯的重要性。'),
  ('异常闭环','构建缺料、设备停机和质量异常三类事件的发现—隔离—处置—恢复—复测—关闭流程。'),
  ('计划调整','学生根据不合格/返工造成的可交付量变化，判断是否需要补产、延期或能力调整，并说明批准/沟通对象。')],
 '异常闭环流程图＋质量/数量状态调整表。','不合格数量、返工/报废状态、物料/工单/计划影响、责任与审批点、追溯字段。','课程作业工单执行与异常闭环（20%内部权重）阶段2＋期末考核；目标1.4、2.4、3.4。',
 ['返工品在业务上为什么不能简单当作完工品？','异常关闭至少需要哪些证据？','质量异常为什么可能需要重新运行MRP/能力评估？'],
 '完成一个异常闭环案例，预习MES功能模型与车间事实数据。',
 '以不合格品和停机记录说明“如实记录比数据好看更重要”，强化质量追溯和管理责任。',
 '将异常闭环设计为可配置工作流，为MES质量、维护和计划联动的产品功能设计建立框架。',
 '投影仪、质量/物料异常案例、流程图模板、数量状态表。',REFS['main']+'\n'+REFS['mes']),

L(11,'第四→第五单元过渡','车间执行反馈、MES功能入口与执行闭环','第四单元1课时＋第五单元1课时',
 {'车间作业控制与执行闭环':1,'MES功能、数据与ERP-MES集成':1},
 {'k':['理解车间执行事实（工单、报工、WIP、质量、设备状态）如何形成MES的核心信息入口。','理解MES“管理执行闭环”而非简单数据展示，其功能围绕人员、设备、物料、质量、工单和文档组织。'],
  'a':['能够把上一单元工单/异常流程映射到MES功能模块和数据对象。','能够区分计划要求、执行事实、分析结果和管理指令的方向。'],
  'v':['形成事实先于看板、业务闭环先于界面美观的数字化建设意识。']},
 {'执行事实':['工单 / 报工','WIP / 质量 / 设备'],'MES对象':['人员 / 设备','物料 / 工单 / 文档'],'数据方向':['计划下达','执行反馈'],'闭环价值':['状态一致','异常可追溯']},
 '执行事实到MES功能模型的映射。','避免以界面/看板代替真实执行闭环，明确计划下行、事实上行和异常处理的方向。',
 '展示一个“看板显示100%完成，但报工/质量记录不完整”的案例，让学生判断MES是否真的完成了执行闭环。',
 [('执行收口','汇总工单、WIP、质量、物料和异常反馈，建立车间事实清单与状态来源。'),
  ('MES功能映射','把事实清单映射到工单管理、资源管理、物料管理、质量管理、追溯、文档和数据采集等功能。'),
  ('方向检查','学生在数据流图上标注“计划/指令下行、执行事实/异常上行、分析/决策回路”，查找方向反转或职责混乱。')],
 '车间事实—MES功能映射表＋上下行数据流草图。','事实来源、MES对象、计划/状态方向、异常闭环和至少一条界面与事实不一致风险。','课程作业工单执行与异常闭环（20%内部权重）完成，并衔接ERP-MES数据流作业；目标1.4/2.4/3.4与1.5/2.5/3.5。',
 ['MES的核心价值为什么不等于“可视化大屏”？','计划数据与执行事实通常各从哪个方向流动？','看板状态和实际报工不一致时应先相信哪一个？为什么？'],
 '整理车间执行阶段作业；为下一课列出MES需要管理的至少6类对象和其关键状态。',
 '强调数字化系统必须忠实反映生产事实，避免“数据装饰”替代真实质量和执行管理。',
 '将车间业务对象映射到MES模块，训练从流程痛点到数字化产品功能的需求分析能力。',
 '投影仪、工单/质量/设备状态案例、MES功能卡片、数据流模板。',REFS['mes']+'\n'+REFS['main']),

L(12,'第五单元 MES功能、数据与ERP-MES集成','MES核心功能、生产数据采集与资源状态','理论2课时',
 {'MES功能、数据与ERP-MES集成':2},
 {'k':['理解MES在生产工单、物料、人员、设备、质量、追溯和文档等方面的核心功能。','理解生产数据采集的事件、时间戳、数据来源、状态语义及条码/RFID等识别技术的概念作用。'],
  'a':['能够为一个制造单元设计基本MES对象清单、事件清单和数据采集点。','能够识别“采集到了数据但语义不清/时间错乱/来源不可追溯”的数据质量问题。'],
  'v':['形成数据来源可追溯、采集最小必要和员工/设备数据使用合规意识。']},
 {'MES对象':['工单 / 物料','人员 / 设备 / 质量'],'采集事件':['开工 / 完工','停机 / 报警 / 检验'],'标识技术':['条码','RFID概念'],'数据质量':['时间戳','来源 / 状态语义']},
 'MES功能对象、事件采集和数据质量。','从业务事件定义采集点与状态，而不是先收集大量无语义数据。',
 '给出“设备每秒上报一条状态，但没有工单号、人员和事件语义”的数据流，让学生判断其对生产追溯到底有多大价值。',
 [('功能模型','围绕工单、资源、物料、质量、文档、追溯建立MES功能矩阵，说明功能之间通过共同业务对象连接。'),
  ('事件采集','从开工、完工、停机、换型、报警、检验等事件定义时间戳、对象标识、状态和值，讨论条码/RFID概念用途。'),
  ('数据质量','学生检查一组缺失工单号、重复时间戳或单位不一致的数据，说明为什么“采集量大”不代表“可用”。')],
 'MES对象/事件/采集点矩阵＋数据质量问题清单。','对象ID、事件、时间戳、状态、来源、单位、关联工单/物料和数据质量检查项。','课程作业ERP-MES数据流与追溯（15%内部权重）阶段1＋期末考核；目标1.5、2.5、3.5。',
 ['MES数据采集为什么必须有业务语义？','条码/RFID解决的是识别问题还是完整业务流程问题？','生产数据的时间戳和来源为什么必须可追溯？'],
 '完善MES采集矩阵，并标注哪些数据应由ERP主责、哪些由MES/现场产生。',
 '讨论人员、质量和设备数据使用中的隐私、权限与最小必要原则，强化规范采集与审计意识。',
 '把MES对象和事件抽象为可扩展数据模型，为接口API、事件总线和工业数据平台设计建立基础。',
 '投影仪、MES功能模型示意、生产事件卡片、采集点模板。',REFS['mes']),

L(13,'第五单元 MES功能、数据与ERP-MES集成','ERP-MES主数据/业务数据接口、状态同步与追溯','理论2课时',
 {'MES功能、数据与ERP-MES集成':2},
 {'k':['掌握ERP-MES常见主数据和业务数据对象、下行/上行方向及基本同步逻辑。','理解接口幂等、版本、时间戳、状态映射、失败重试和一致性风险的基本概念。'],
  'a':['能够绘制物料/BOM/工艺/工单下行与报工/质量/消耗/完工上行的数据流。','能够设计一个产品批次从订单到原料/工序/质量/成品的追溯链，并指出断链风险。'],
  'v':['形成权限边界、接口契约、数据一致性和跨系统审计意识。']},
 {'下行对象':['物料 / BOM / 工艺','计划 / 生产工单'],'上行对象':['报工 / 消耗','质量 / 完工 / 状态'],'接口一致性':['版本 / 幂等','状态映射 / 重试'],'追溯链':['批次 / 序列','订单—工序—质量']},
 'ERP-MES数据流、状态同步和端到端追溯。','处理跨系统对象主责、状态映射和异常重试，避免重复/丢失/乱序导致业务不一致。',
 '给出“ERP重试发送同一工单，MES因此创建两张工单”的案例，让学生判断需要何种接口标识和幂等规则。',
 [('接口对象','建立ERP→MES主数据/工单下行和MES→ERP报工/质量/消耗/完工上行表，明确主责、触发时机和关键字段。'),
  ('一致性控制','通过重复消息、乱序状态、接口超时案例讨论唯一标识、版本、状态映射、重试和人工补偿。'),
  ('追溯设计','学生用订单号、批次/序列号、工单、工序、物料批次、检验记录建立追溯链，并检查一处断链。')],
 'ERP-MES数据流图＋接口对象表＋追溯链。','对象主责、方向、关键字段、状态映射、失败处理和端到端追溯路径。','课程作业ERP-MES数据流与追溯（15%内部权重）阶段2＋期末考核；目标1.5、2.5、3.5。',
 ['为什么接口重试可能导致重复工单？','主数据和业务状态同步的频率是否应该完全一样？','产品追溯链至少需要哪些关联标识？'],
 '完成ERP-MES接口分析作业；预习APS/KPI需要哪些来自计划和执行层的数据。',
 '通过接口权限、审计和追溯链强调跨系统协作中的契约意识和数据责任。',
 '把接口对象、状态映射和追溯标识设计为可版本化契约，训练数字化平台集成产品思维。',
 '投影仪、ERP-MES数据流案例、接口对象模板、追溯链模板。',REFS['mes']+'\n'+REFS['erp']),

L(14,'第五→第六单元过渡','ERP-MES集成数据到APS/KPI决策输入','第五单元1课时＋第六单元1课时',
 {'MES功能、数据与ERP-MES集成':1,'智能计划、绩效与管理决策':1},
 {'k':['理解计划、工单、资源、质量和执行事实如何成为APS排程和KPI计算输入。','理解数据时效性、口径、缺失/异常和权限对智能计划与绩效判断的影响。'],
  'a':['能够建立“业务问题—所需数据—来源系统—口径/刷新频率—质量检查”表。','能够识别用错误时间窗、错误口径或缺失状态计算KPI/排程造成的决策风险。'],
  'v':['形成“算法之前先治理数据、指标之前先统一口径”的决策责任意识。']},
 {'决策输入':['订单 / 计划','资源 / WIP / 质量'],'来源系统':['ERP','MES / 现场'],'数据治理':['时间戳 / 口径','缺失 / 异常 / 版本'],'决策输出':['APS排程','KPI / 管理建议']},
 'ERP/MES集成数据作为APS和KPI决策输入。','把数据口径、时效和质量前置到算法/指标讨论，避免“精确计算错误数据”。',
 '展示“ERP库存为昨日日终、MES WIP为实时、产能日历未更新”却直接用于APS排程的案例，让学生指出至少三类数据时效风险。',
 [('集成收口','回顾ERP-MES接口对象和追溯链，区分计划数据、执行事实、主数据和状态数据。'),
  ('决策数据准备','面向APS与KPI建立数据需求表，讨论时间窗、刷新频率、单位口径、缺失/异常和版本。'),
  ('风险评审','学生对一组不同时间戳/口径的数据做“可用/需修正/不可用于决策”判定，并说明理由。')],
 'APS/KPI数据需求与质量检查表。','数据项、来源系统、时间戳/频率、口径、质量检查和是否可用于决策的判断。','课程作业ERP-MES数据流与追溯（15%内部权重）完成，并衔接KPI判断作业；目标1.5/2.5/3.5与1.6/2.6/3.6。',
 ['为什么实时MES数据与日结ERP数据不能不加说明地直接拼接？','KPI口径不一致会造成什么管理问题？','APS优化前最少应做哪些数据质量检查？'],
 '整理ERP-MES数据流作业；准备一个计划方案，列出需要评价的交期、库存、WIP、利用率和质量指标。',
 '强调智能决策不能把数据问题隐藏在算法之后，技术人员必须说明数据时效、口径和不确定性。',
 '把数据质量门控设计成APS/KPI服务的前置校验机制，训练智能制造数据产品的可靠性设计。',
 '投影仪、跨系统数据时间戳案例、数据质量检查表、白板。',REFS['mes']+'\n'+REFS['erp']),

L(15,'第六单元 智能计划、绩效与管理决策','APS多约束排程与制造KPI评价','理论2课时',
 {'智能计划、绩效与管理决策':2},
 {'k':['理解APS、多约束计划和优化排程的基本思想，认识约束、目标函数和启发式/优化结果的边界。','掌握OEE、准时交付率、库存周转、WIP、一次合格率等常见指标的基本含义和相互影响。'],
  'a':['能够比较两个简化排程方案在交期、库存/WIP、利用率和质量风险上的差异。','能够选择与业务问题匹配的KPI，并识别单指标优化带来的副作用。'],
  'v':['形成成本、效率、质量、交付和员工影响综合平衡的管理意识，不以单一KPI驱动不合理行为。']},
 {'APS':['多约束','目标 / 优先规则'],'交付指标':['准时交付率','延期 / 周期'],'资源指标':['OEE / 利用率','瓶颈 / 负荷'],'库存质量':['库存周转 / WIP','一次合格率']},
 'APS约束与制造KPI的多目标评价。','避免把OEE/利用率等指标等同于最终经营目标，能解释指标间冲突和局限。',
 '给出方案A“高利用率但WIP高、交期波动”和方案B“利用率略低但准时交付高”的数据，让学生先判断哪一个更优以及还缺什么信息。',
 [('APS思想','从人工优先规则扩展到多约束排程，解释订单、物料、能力、换型、设备状态、交期等约束与目标。'),
  ('KPI计算','用简化数据计算/解释准时交付率、OEE、库存周转、WIP和一次合格率，强调指标口径和时间窗。'),
  ('方案比较','学生用多指标评分/权衡表比较两个排程方案，说明推荐方案、适用条件和可能副作用。')],
 '双方案APS/KPI比较表＋管理建议。','约束清单、关键KPI、计算口径、方案差异、推荐理由和至少一条指标副作用。','课程作业KPI与管理判断（10%内部权重）阶段1＋期末考核；目标1.6、2.6、3.6。',
 ['APS与简单优先规则的核心区别是什么？','OEE提高是否一定提高准时交付率？','为什么KPI必须同时说明口径和时间窗？'],
 '完善方案比较表，并补充成本/经济性和员工/合规影响，为期末综合业务分析报告准备。',
 '通过单指标激励失真案例强调负责任绩效管理和对员工、质量、交付的综合责任。',
 '将APS约束、KPI和方案评价组成可解释决策支持原型，训练智能计划产品的指标设计能力。',
 '投影仪、APS简化排程案例、KPI数据表、方案评价矩阵。',REFS['main']+'\n'+REFS['mes']),

L(16,'第六单元 智能计划、绩效与管理决策','成本经济决策、可视化看板与MES/ERP综合方案评审','理论2课时',
 {'智能计划、绩效与管理决策':2},
 {'k':['理解库存、加班、外协、换型、停机、质量和延期等成本因素，以及可视化看板与持续改进的基本逻辑。','理解期末MES/ERP综合业务分析报告应覆盖业务流程、计划与能力、MES执行追溯、ERP-MES集成、KPI/经济性和个人核验。'],
  'a':['能够对综合制造案例形成端到端MES/ERP方案并用交付、库存、WIP、资源、质量和成本证据评价。','能够独立解释方案假设、数据来源、系统边界和改进建议，并针对指定数据变更现场修正部分计算/流程。'],
  'v':['形成经济决策、合规、数据诚信、协作责任和持续改进意识，对采用的数字化方案和管理建议承担解释责任。']},
 {'成本决策':['库存 / 加班 / 外协','质量 / 延期 / 换型'],'可视化':['业务看板','指标口径 / 异常钻取'],'持续改进':['问题—原因','措施—验证'],'综合方案':['计划—执行—集成','KPI / 经济性 / 答辩']},
 '成本/KPI综合评价、可视化与端到端MES/ERP方案评审。','在不同系统、数据和指标间保持一致的业务逻辑，并明确假设、成本、人员与合规边界。',
 '展示一个“通过大量加班实现100%准时交付”的方案，让学生判断如果只看OTD为什么可能得出错误管理结论。',
 [('经济性与看板','梳理库存、加班、外协、换型、质量、停机和延期成本，讨论看板应从异常和决策问题出发而非堆叠图表。'),
  ('综合方案串联','把订单/MPS/MRP、能力/排程、工单/WIP/质量、MES、ERP-MES接口、追溯、APS/KPI串成端到端方案。'),
  ('方案评审与个人核验','学生用评审清单交叉检查业务边界、计算、数据流、KPI和经济性，并模拟回答“若订单/能力/库存改变如何调整方案”。')],
 'MES/ERP综合业务分析报告框架＋个人答辩核验清单。','端到端业务图、MPS/MRP与能力摘要、MES执行/追溯、ERP-MES数据流、KPI/成本比较、假设与改进记录。','课程作业KPI与管理判断（10%内部权重）完成＋期末考核60%核心准备；全部课程目标。',
 ['为什么准时交付率不能脱离成本和质量单独评价？','一个合格的MES/ERP综合方案必须说明哪些系统边界？','个人核验为什么需要现场修改/指定计算，而不是只做团队汇报？'],
 '按期末要求完成综合业务分析报告和个人准备；所有数据、流程、计算与引用应可追溯并标明假设。',
 '通过成本、员工影响、权限、数据诚信和可追溯性强化负责任数字化决策，反对只追求“漂亮指标”。',
 '将综合方案视作制造数字化咨询/产品方案原型，训练需求分析、数据设计、流程编排、指标与经济性论证能力。',
 '投影仪、综合制造案例包、业务流程/数据流模板、KPI/成本评审表。',REFS['main']+'\n'+REFS['mes']+'\n'+REFS['erp'])
]


def render_map(path:Path,title:str,branches:Dict[str,List[str]]):
    dot=Digraph('G',format='png')
    dot.attr(rankdir='TB',bgcolor='white',margin='0.02',pad='0.05',nodesep='0.20',ranksep='0.36',dpi='190')
    dot.attr('node',shape='box',style='rounded,filled',fontname='Noto Sans CJK SC',fontsize='11',color='#6B7280',penwidth='1.0',fillcolor='#F8FAFC',margin='0.10,0.06')
    dot.attr('edge',color='#94A3B8',arrowsize='0.55',penwidth='0.9')
    dot.node('root',title,fillcolor='#E6F1F8',color='#2F6F8F',fontsize='14',penwidth='1.4')
    for bi,(b,leaves) in enumerate(branches.items(),1):
        bid=f'b{bi}'; dot.node(bid,b,fillcolor='#EEF6F1',color='#5D846E',fontsize='12'); dot.edge('root',bid)
        for li,leaf in enumerate(leaves,1):
            nid=f'{bid}_{li}'; dot.node(nid,leaf,fillcolor='#FFFFFF',color='#A5B4C2',fontsize='10'); dot.edge(bid,nid)
    outbase=str(path.with_suffix('')); dot.render(outbase,cleanup=True)
    gen=Path(outbase+'.png')
    if gen!=path: shutil.move(gen,path)

def obj_md(g):
    def B(label,arr): return f'**{label}：**\n'+'\n'.join(f'{i+1}. {x}' for i,x in enumerate(arr))
    return '\n\n'.join([B('知识目标',g['k']),B('能力目标',g['a']),B('价值目标',g['v'])])

def process_md(l):
    times=[('一、业务概念与数据模型',30),('二、学生计算/流程/数据推演',25),('三、方案比较、校核与错误复盘',20)]
    out=[]
    for (label,mins),(title,detail) in zip(times,l['blocks']):
        out.append(f'**{label}（约{mins} min）—{title}**\n\n【教师】{detail}\n\n【学生】先独立完成流程、表格、计算或数据流推演并写出判断依据，再与同伴互核；需要电子表格/Python时仅作为计算辅助，不把软件操作本身作为课程目标。')
    out += [f'**独立学习产物**：{l["product"]}',f'**学习证据**：{l["evidence"]}',f'**评价映射**：{l["assessment"]}']
    return '\n\n'.join(out)

def preclass(l):
    return f'''【教师】发布“{l['topic']}”对应大纲/教材范围和一个最小制造业务案例，不提前给出完整结论。\n\n1. 阅读对应大纲与教材内容；\n2. 圈出3个关键术语并记录至少1个疑问；\n3. 预读订单/BOM/库存/工艺/能力/工单/接口/KPI中的相关数据表，写出一个预期业务关系或计算结果。\n\n【学生】完成准备并带着“业务关系预测/疑问/已有计算或流程证据”进入课堂。'''

def reflection(l):
    return f'''【课后填写，不预填事实】记录本次实际用时、“{l['diff']}”的真实掌握/错误，以及流程图、计算表、数据流/接口表、方案比较等评价证据完整性；仅依据真实课堂记录确定后续补救或拓展，不补写未发生事实。'''

def clean_md(s):
    s=re.sub(r'!\[[^\]]*\]\([^\)]*\)','',s); s=s.replace('**','').replace('`',''); s=re.sub(r'^- ','',s,flags=re.M); s=re.sub(r'\n{3,}','\n\n',s); return s.strip()

def extract_info(s):
    d={}
    for line in s.splitlines():
        m=re.match(r'- \*\*(.+?)\*\*：(.+)',line.strip())
        if m: d[m.group(1)]=m.group(2).strip()
    return d

def parse_md(path:Path):
    text=path.read_text(encoding='utf-8'); parts=re.split(r'(?=^## 第\d+次课\s)',text,flags=re.M); parts=[p for p in parts if re.match(r'^## 第\d+次课\s',p)]
    heads=['课次信息','教学目标','课程思政','专创融合','教学重难点','教学方法与用具','教学设计','课前任务','互动导入','传授新知与课堂训练','过关检测','课堂小结','作业布置','考勤','课后教学反思','本章节参考文献']
    parsed=[]
    for p in parts:
        h=re.match(r'^## 第(\d+)次课\s+(.+)$',p,re.M); fields={}
        for i,hd in enumerate(heads):
            st=re.search(rf'^### {re.escape(hd)}\s*$',p,re.M)
            if not st: continue
            pos=st.end(); nxt=[]
            for hd2 in heads[i+1:]:
                m2=re.search(rf'^### {re.escape(hd2)}\s*$',p[pos:],re.M)
                if m2: nxt.append(pos+m2.start())
            fields[hd]=p[pos:min(nxt) if nxt else len(p)].strip()
        img=re.search(r'!\[知识脉络图\]\((images/lesson-\d+-knowledge-map\.png)\)',p); fields['image']=img.group(1) if img else None
        parsed.append({'no':int(h.group(1)),'topic':h.group(2).strip(),'fields':fields})
    return parsed,text

def remove_table(table):
    e=table._element; e.getparent().remove(e)

def set_cell_text(cell,text,font_size=10.5,bold_prefixes=None,align=None,line_spacing=1.05):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=line_spacing
    if align is not None: p.alignment=align
    for i,line in enumerate(str(text).split('\n')):
        if i>0:p.add_run().add_break()
        r=p.add_run(line); r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(font_size)
        if bold_prefixes and any(line.strip().startswith(x) for x in bold_prefixes): r.bold=True

def set_cell_margins(cell,top=60,start=70,bottom=60,end=70):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        n=tcMar.find(qn(f'w:{m}'))
        if n is None:n=OxmlElement(f'w:{m}'); tcMar.append(n)
        n.set(qn('w:w'),str(v)); n.set(qn('w:type'),'dxa')

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); h=OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); trPr.append(h)

def sha256(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for c in iter(lambda:f.read(1<<20),b''): h.update(c)
    return h.hexdigest()

def main():
    if OUT_ROOT.exists(): shutil.rmtree(OUT_ROOT)
    IMG_DIR.mkdir(parents=True)
    render_map(IMG_DIR/'course-knowledge-map.png','智能生产计划管理（MES/ERP）知识主线',COURSE_MAP)
    for l in lessons: render_map(IMG_DIR/f'lesson-{l["no"]:02d}-knowledge-map.png',l['topic'],l['map'])

    md=['# 《智能生产计划管理（MES/ERP）》教学设计（16次课）','',
        '> **Canonical source / 内容权威源**：本 Markdown 与 `images/`、`manifest.json` 共同构成教案语义主源。Word 仅由本源包编译，负责学校格式呈现。','',
        '## 课程基本信息',
        f'- **课程名称**：{COURSE["name"]}',f'- **课程英文名称**：{COURSE["english"]}',f'- **课程代码**：{COURSE["code"]}',f'- **课程类别**：{COURSE["category"]}',f'- **课程性质**：{COURSE["nature"]}',f'- **授课语言**：{COURSE["language"]}',f'- **授课学期**：{COURSE["semester"]}',f'- **学分**：{COURSE["credits"]}',
        '- **课程学时及分配**：总学时32；理论32；实践0；实验0；上机0；项目式0',f'- **适用专业**：{COURSE["major"]}',f'- **授课学院**：{COURSE["college"]}',f'- **先修课程**：{COURSE["prereq"]}',f'- **后续课程**：{COURSE["followup"]}',f'- **选用教材**：{COURSE["textbook"]}',f'- **课程评价**：{COURSE["assessment"]}',f'- **内容依据**：{COURSE["version"]}（大纲更新时间：{COURSE["syllabus_date"]}）','',
        '## 课程总体设计',
        '课程严格按现行培养方案与大纲执行32学时全理论教学，不虚构实验、上机或项目式学时。16次课×2课时=32学时。课程采用“业务场景—数据模型—计划计算—执行闭环—案例评审”的理论教学方式，允许使用Excel、Python或演示型MES/ERP界面辅助课堂计算和说明，但软件操作本身不是课程目标。六个单元学时严格闭合为4+6+5+6+6+5=32学时，其中第8、11、14次课各承担相邻单元1+1课时的过渡衔接。课程学习证据以业务流程图、MPS/MRP计算表、能力负荷表、工单状态/异常闭环、ERP-MES数据流与追溯链、KPI/成本方案比较和综合业务分析报告为主。','',
        '![课程知识脉络](images/course-knowledge-map.png)','',
        '### 课次总览','| 次数 | 课型 | 单元学时分配 | 单元/过渡 | 授课题目 | 主要评价 |','|---:|---|---|---|---|---|']
    for l in lessons: md.append(f'| {l["no"]} | 理论 | {l["split"]} | {l["section"]} | {l["topic"]} | {l["assessment"].split("；")[0]} |')
    md.append('')
    for l in lessons:
        md += [f'## 第{l["no"]}次课 {l["topic"]}','','### 课次信息',f'- **章节/单元**：{l["section"]}','- **周次**：按实际课表填写','- **课时安排**：2课时（100 min）','- **课型**：理论',f'- **单元学时分配**：{l["split"]}','','![知识脉络图](images/lesson-%02d-knowledge-map.png)'%l['no'],'','### 教学目标',obj_md(l['goals']),'','### 课程思政',l['ideology'],'','### 专创融合',l['innovation'],'','### 教学重难点',f'**教学重点：**{l["focus"]}\n\n**教学难点：**{l["diff"]}','','### 教学方法与用具',f'**教学方法：**业务场景法、案例分析法、表格推演法、问题驱动法、数据流/流程建模法、方案比较法、错误复盘法\n\n**教学用具：**{l["tools"]}','','### 教学设计','课前任务 → 互动导入（10 min）→ 传授新知与课堂训练（75 min）→ 过关检测（10 min）→ 课堂小结（3 min）→ 作业布置（1 min）→ 考勤（1 min）','','### 课前任务',preclass(l),'','### 互动导入',f'【教师】{l["intro"]}\n\n【学生】先独立判断业务关系、流程或计算结果，再与同伴交换依据；教师收集典型分歧后进入本课。','','### 传授新知与课堂训练',process_md(l),'','### 过关检测','【教师】组织当堂检测：\n\n'+'\n'.join(f'{i+1}. {q}' for i,q in enumerate(l['checks']))+'\n\n【学生】独立作答/计算/绘图；【教师】按错误类型讲解，要求说明业务、数据、计算或流程依据。','','### 课堂小结',f'【教师】回到本课知识脉络图，用“业务场景—对象/数据—计划/执行逻辑—校核/异常—管理判断”总结“{l["topic"]}”，再次强调教学重点：{l["focus"]}\n\n【学生】写下“本课一个关键业务规则 + 一个最容易造成错误决策的数据/边界”。','','### 作业布置',f'【教师】布置课后任务：{l["homework"]}\n\n【学生】按课程文件命名与证据要求整理流程图、计算表、数据流或方案说明；使用电子工具时需保留输入、规则和人工核验记录。','','### 考勤','【教师】利用学校/课程实际使用的平台进行签到。\n\n【学生】按课程要求完成签到。','','### 课后教学反思',reflection(l),'','### 本章节参考文献',l['refs'],'']
    md_path=OUT_ROOT/MD_NAME; md_path.write_text('\n'.join(md),encoding='utf-8')

    manifest={'course':{'name':COURSE['name'],'code':COURSE['code'],'hours':32,'theory_hours':32,'lab_hours':0},'canonical_markdown':MD_NAME,'lesson_count':16,
              'syllabus':{'path':COURSE['syllabus_path'],'sha':COURSE['syllabus_sha'],'updated':COURSE['syllabus_date'],'version':COURSE['version']},
              'template_identifier':'程序设计基础（C&C++）Ⅰ最终校版学校教学设计布局基线（仅作Word格式模板）','course_map':'images/course-knowledge-map.png',
              'lessons':[{'no':l['no'],'type':'理论','hour_split':l['split'],'unit_hours':l['unit_hours'],'section':l['section'],'title':l['topic'],'image':f'images/lesson-{l["no"]:02d}-knowledge-map.png','assessment':l['assessment']} for l in lessons]}
    (OUT_ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT_ROOT/'README.md').write_text(f'''# 智能生产计划管理（MES/ERP）教案源码包\n\n本包采用 Markdown-first 工作流。\n\n- `{MD_NAME}`：唯一课程内容主源。\n- `images/`：课程总图与16次课知识脉络图。\n- `manifest.json`：课程、课次、图片、大纲版本和六单元学时映射。\n- `build_docx.py`：从内容主源编译学校格式Word的脚本。\n\n## 内容/格式边界\nMarkdown负责课程内容正确性；Word负责学校格式、表格、字体、分页和图片尺寸。Word视觉检查发现语义问题时，必须回到Markdown修改后重新编译。\n\n## 权威依据\n`{COURSE['syllabus_path']}`，SHA `{COURSE['syllabus_sha']}`，大纲更新时间 `{COURSE['syllabus_date']}`。\n\n## 学时硬边界\n本课程32学时全部为理论学时（理论32、上机0、实验0、项目式0）。Excel、Python或演示型MES/ERP界面仅可作为课堂辅助工具，不据此声明上机或实验学时。\n''',encoding='utf-8')

    # CONTENT QA GATE
    parsed,text=parse_md(md_path)
    refs=re.findall(r'!\[[^\]]*\]\((images/[^\)]+)\)',text); missing=[r for r in refs if not (OUT_ROOT/r).exists()]
    assert len(parsed)==16 and not missing
    assert '总学时32；理论32；实践0；实验0；上机0；项目式0' in text
    assert COURSE['assessment'] in text
    assert len(set(p['fields']['image'] for p in parsed))==16
    assert all('学习证据' in p['fields']['传授新知与课堂训练'] and '评价映射' in p['fields']['传授新知与课堂训练'] for p in parsed)
    assert all('【课后填写，不预填事实】' in p['fields']['课后教学反思'] for p in parsed)
    sums={u:sum(l['unit_hours'].get(u,0) for l in lessons) for u in UNIT_NAMES}
    assert sums==dict(zip(UNIT_NAMES,[4,6,5,6,6,5])), sums
    shutil.copy2(md_path,FINAL_MD)

    # WORD COMPILATION
    doc=Document(str(TEMPLATE)); proto=copy.deepcopy(doc.tables[2]._element)
    for t in list(doc.tables[2:]): remove_table(t)
    for par in list(doc.paragraphs[24:]): par._element.getparent().remove(par._element)
    cover=doc.tables[0]
    for i,val in enumerate([COURSE['college'],COURSE['major'],COURSE['name'],'（按实际填写）','（按实际填写）',COURSE['semester'],COURSE['major']]): set_cell_text(cover.cell(i,1),val,13.0)
    info=doc.tables[1]
    set_cell_text(info.cell(0,1),COURSE['name'],10.1); set_cell_text(info.cell(1,1),COURSE['english'],8.9); set_cell_text(info.cell(2,1),COURSE['category']); set_cell_text(info.cell(2,3),COURSE['nature']); set_cell_text(info.cell(2,5),COURSE['language'])
    set_cell_text(info.cell(3,1),COURSE['semester']); set_cell_text(info.cell(3,5),COURSE['credits'])
    for c,v in zip(range(1,6),[32,32,0,0,0]): set_cell_text(info.cell(5,c),str(v))
    set_cell_text(info.cell(6,1),COURSE['major'],9.2); set_cell_text(info.cell(7,1),COURSE['textbook'],9.0); set_cell_text(info.cell(8,1),COURSE['college']); set_cell_text(info.cell(9,1),COURSE['prereq'],9.0); set_cell_text(info.cell(10,1),COURSE['followup'],9.0)
    set_cell_text(info.cell(11,1),'考试课（ ）；考查课（√）',10.4)
    set_cell_text(info.cell(12,1),'闭卷（ ）；开卷（ ）；笔试（ ）；机试（ ）；口试（ ）；其它（√）：综合业务分析报告＋个人答辩/核验',9.2)
    set_cell_text(info.cell(13,1),'平时成绩（√）；课程作业（√）；期末考核（√）',9.8)
    set_cell_text(info.cell(14,1),COURSE['assessment'],10)
    intro='课程基本定位：面向制造企业“订单—计划—物料—能力—工单—车间执行—质量—追溯—成本”业务链，建立ERP/MRP/MRPⅡ、MPS/MRP、能力排程、车间执行、MES与ERP-MES集成、APS/KPI/成本的连续知识链。\n核心学习结果：能够完成MPS/MRP基本计算、能力负荷分析、工单状态与异常闭环、ERP-MES数据流与追溯设计，并比较不同计划方案的交期、库存、WIP、利用率、质量与成本影响。\n主要教学方法：业务场景—数据模型—计划计算—执行闭环—案例评审；32学时全部为理论，电子表格/Python/MES-ERP演示界面仅作辅助，不虚构上机或实验学时。'
    set_cell_text(info.cell(15,1),intro,8.9,bold_prefixes=['课程基本定位','核心学习结果','主要教学方法'])

    for p in parsed:
        sep=doc.add_paragraph(); sep.paragraph_format.page_break_before=True; sep.paragraph_format.space_before=Pt(0); sep.paragraph_format.space_after=Pt(0); new_tbl=copy.deepcopy(proto); sep._p.addnext(new_tbl); t=doc.tables[-1]
        f=p['fields']; li=extract_info(f['课次信息'])
        set_cell_text(t.cell(0,1),li.get('章节/单元',''),9.2); set_cell_text(t.cell(0,3),p['topic'],9.2); set_cell_text(t.cell(1,1),li.get('周次','按实际课表填写')); set_cell_text(t.cell(1,3),li.get('课时安排','2课时（100 min）')+'；'+li.get('单元学时分配',''))
        c=t.cell(2,1); c.text=''; pp=c.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.add_run().add_picture(str(OUT_ROOT/f['image']),width=Cm(11.7))
        set_cell_text(t.cell(3,1),clean_md(f['教学目标']),9.25,bold_prefixes=['知识目标','能力目标','价值目标']); set_cell_text(t.cell(4,1),clean_md(f['课程思政']),9.0); set_cell_text(t.cell(5,1),clean_md(f['专创融合']),9.0); set_cell_text(t.cell(6,1),clean_md(f['教学重难点']),9.0,bold_prefixes=['教学重点','教学难点'])
        meth=clean_md(f['教学方法与用具']); mm=re.search(r'教学方法：(.*?)(?:\n\n|\n)教学用具：(.*)',meth,re.S)
        if mm: set_cell_text(t.cell(7,1),mm.group(1).strip(),8.7); set_cell_text(t.cell(8,1),mm.group(2).strip(),8.7)
        else: set_cell_text(t.cell(7,1),meth,8.8); set_cell_text(t.cell(8,1),'制造业务案例表、流程图模板。',8.8)
        set_cell_text(t.cell(9,1),clean_md(f['教学设计']),8.8); set_cell_text(t.cell(10,1),'教学步骤及主要教学内容',10.1); set_cell_text(t.cell(10,4),'设计意图',10.1)
        set_cell_text(t.cell(11,1),clean_md(f['课前任务']),8.45,bold_prefixes=['【教师】','【学生】'],line_spacing=1.0); set_cell_text(t.cell(11,4),'通过业务数据预读、关系预测和问题清单建立先备认知，避免课堂只记概念不理解数据与流程。',8.2)
        combo='【互动导入 10 min】\n'+clean_md(f['互动导入'])+'\n\n【传授新知与课堂训练 75 min】\n'+clean_md(f['传授新知与课堂训练'])
        set_cell_text(t.cell(12,0),'互动导入\n（10 min）\n\n传授新知与课堂训练\n（75 min）',8.6,line_spacing=1.0); set_cell_text(t.cell(12,1),combo,7.95,bold_prefixes=['【互动导入','【传授','一、','二、','三、','独立学习产物','学习证据','评价映射'],line_spacing=0.98); set_cell_text(t.cell(12,4),'采用“业务场景→对象/数据→计划/执行逻辑→计算/流程推演→方案比较→证据与管理判断”的闭环，所有结论保留数据口径和假设。',8.05)
        set_cell_text(t.cell(13,1),clean_md(f['过关检测']),8.1,line_spacing=1.0); set_cell_text(t.cell(13,4),'检测学生是否能从业务、数据、计算与流程依据解释结论，而不是只复述系统名词。',8.2)
        set_cell_text(t.cell(14,1),clean_md(f['课堂小结']),8.1,line_spacing=1.0); set_cell_text(t.cell(14,4),'回到知识脉络图，固化“业务—数据—计划—执行—反馈—决策”的可迁移结构。',8.2)
        set_cell_text(t.cell(15,1),clean_md(f['作业布置']),8.1,line_spacing=1.0); set_cell_text(t.cell(15,4),'固化流程图、计算表、接口表和方案比较证据，递进形成期末综合业务分析报告。',8.2)
        set_cell_text(t.cell(16,1),clean_md(f['考勤']),8.1,line_spacing=1.0); set_cell_text(t.cell(16,4),'保留学校模板考勤环节，不预填实际出勤事实。',8.2)
        set_cell_text(t.cell(17,1),clean_md(f['课后教学反思']),7.85,line_spacing=1.0); set_cell_text(t.cell(17,4),'课后反思必须基于真实课堂计算、流程和讨论证据；课前只保留填写框架。',8.05)
        set_cell_text(t.cell(18,0),'本章节参考文献',7.8,align=WD_ALIGN_PARAGRAPH.CENTER,line_spacing=1.0); set_cell_text(t.cell(18,1),clean_md(f['本章节参考文献']),7.7,line_spacing=1.0)
        for ri,row in enumerate(t.rows):
            for cc in row.cells:
                cc.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri==12: set_cell_margins(cc,top=30,start=65,bottom=30,end=65)
                elif ri in (17,18): set_cell_margins(cc,top=18,start=65,bottom=18,end=65)
                else: set_cell_margins(cc,top=40,start=65,bottom=40,end=65)
        set_repeat_header(t.rows[0])

    cp=doc.core_properties
    cp.title='智能生产计划管理（MES/ERP） 教学设计（16次课）'
    cp.subject='32学时全理论；智能制造工程'
    cp.keywords='MES,ERP,MPS,MRP,能力计划,工单,WIP,追溯,APS,KPI,教案,教学设计'
    cp.comments='依据现行智能生产计划管理（MES/ERP）课程教学大纲编制；Markdown为内容主源，Word为学校格式编译产物。'
    doc.save(FINAL_DOCX)
    scrub=Path('/home/oai/skills/docx/scripts/privacy_scrub.py'); clean=BASE/'mes_erp_clean.docx'
    subprocess.run(['python',str(scrub),str(FINAL_DOCX),'--out',str(clean)],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True); shutil.move(clean,FINAL_DOCX)
    shutil.copy2(Path(__file__),OUT_ROOT/'build_docx.py')
    manifest['files']={MD_NAME:sha256(md_path),'docx':sha256(FINAL_DOCX),'images_count':len(list(IMG_DIR.glob('*.png')))}; (OUT_ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    for z in [SOURCE_ZIP,COMPLETE_ZIP]:
        if z.exists(): z.unlink()
    with zipfile.ZipFile(SOURCE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for fp in OUT_ROOT.rglob('*'):
            if fp.is_file(): z.write(fp,fp.relative_to(OUT_ROOT.parent))
    comp=BASE/'智能生产计划管理_MES-ERP_教案_完整交付包'; shutil.rmtree(comp,ignore_errors=True); shutil.copytree(OUT_ROOT,comp/'source'); shutil.copy2(FINAL_DOCX,comp/DOCX_NAME); shutil.copy2(FINAL_MD,comp/MD_NAME); (comp/'README.md').write_text('本包包含canonical Markdown源包与由其编译的学校格式Word。内容修改请先修改source中的Markdown和语义图片，再重新编译Word。\n',encoding='utf-8')
    with zipfile.ZipFile(COMPLETE_ZIP,'w',zipfile.ZIP_DEFLATED) as z:
        for fp in comp.rglob('*'):
            if fp.is_file(): z.write(fp,fp.relative_to(comp.parent))
    print(json.dumps({'md':str(FINAL_MD),'docx':str(FINAL_DOCX),'source_zip':str(SOURCE_ZIP),'complete_zip':str(COMPLETE_ZIP),'source_dir':str(OUT_ROOT)},ensure_ascii=False,indent=2))

if __name__=='__main__': main()
